from flask import Flask, render_template, request, redirect, send_from_directory, send_file, session, url_for, Response
from flask import make_response
import base64
import qrcode
import os
import io
import uuid
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
import json
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from database import (
    get_connection,
    save_uploaded_file,
    get_uploaded_file,
    delete_uploaded_file,
    json_delete
)


app = Flask(__name__)

app.secret_key = os.environ.get("SECRET_KEY", "development-secret-key")

ADMIN_FILE = "admin.json"
def get_admin_data():

    if not os.path.exists(ADMIN_FILE):

        admin_data = {
            "admin_id": "ADMIN001",
            "password": generate_password_hash("admin123")
        }

        with open(ADMIN_FILE, "w") as file:
            json.dump(admin_data, file, indent=4)

        return admin_data

    try:
        with open(ADMIN_FILE, "r") as file:
            admin_data = json.load(file)

        return admin_data

    except Exception:
        admin_data = {
            "admin_id": "ADMIN001",
            "password": generate_password_hash("admin123")
        }

        with open(ADMIN_FILE, "w") as file:
            json.dump(admin_data, file, indent=4)

        return admin_data


# =========================================================
# SERVER-SIDE ROLE / ROUTE SECURITY
# =========================================================

@app.before_request
def protect_private_routes():

    path = request.path

    # Public/login pages
    public_paths = {
        "/",
        "/admin-login",
        "/faculty-login",
        "/student-login",
        "/admission-login",
    }

    if path in public_paths:
        return None

    # ADMIN-ONLY ROUTES
    admin_paths = (
        "/admin-",
        "/edit-student/",
        "/delete-student/",
        "/reset-student-password/",
        "/admin-delete-student/",
        "/admin-reset-student-password/",
        "/edit-faculty/",
        "/delete-faculty/",
        "/reset-faculty-password/",
        "/edit-course/",
        "/delete-course/",
        "/delete-subject/",
        "/approve-application/",
        "/reject-application/",
    )

    if path.startswith(admin_paths):
        if not session.get("admin_id"):
            return redirect(url_for("admin_login"))
        return None

    # FACULTY-ONLY ROUTES
    faculty_paths = (
        "/manage-attendance",
        "/upload-notes",
        "/delete-note/",
        "/faculty-syllabus",
        "/delete-syllabus/",
        "/send-notice",
        "/delete-notice/",
        "/holiday-information",
        "/faculty-attendance-record",
        "/faculty-attendance-record-excel",
    )

    if path.startswith(faculty_paths):
        if not session.get("faculty_id"):
            return redirect(url_for("faculty_login"))
        return None

    return None

# =========================================================
# ADMIN LOGIN SECURITY
# =========================================================

def admin_required():

    admin_id = session.get("admin_id")
    department_id = session.get("admin_department_id")

    if not admin_id or not department_id:
        session.pop("admin_id", None)
        session.pop("admin_department", None)
        session.pop("admin_department_id", None)

        return redirect(
            url_for("admin_login")
        )

    conn = get_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                SELECT
                    admin_id,
                    department_id,
                    department_name,
                    is_closed
                FROM department_admins
                WHERE admin_id = %s
                  AND department_id = %s
            """, (
                admin_id,
                department_id
            ))

            admin_data = cur.fetchone()

    finally:
        conn.close()

    # Admin account exist nahi karta
    if not admin_data:

        session.pop("admin_id", None)
        session.pop("admin_department", None)
        session.pop("admin_department_id", None)

        return redirect(
            url_for("admin_login")
        )

    # Department temporarily closed hai
    if admin_data["is_closed"]:

        session.pop("admin_id", None)
        session.pop("admin_department", None)
        session.pop("admin_department_id", None)

        return redirect(
            url_for("admin_login")
        )

    return None

# =========================================================
# GLOBAL BROWSER HISTORY + CACHE CONTROL
# =========================================================

@app.after_request
def global_navigation_control(response):

    # Browser ko purane pages cache karne se roko
    response.headers["Cache-Control"] = (
        "no-store, no-cache, must-revalidate, max-age=0"
    )
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"

    # Sirf HTML pages par JavaScript add karo
    content_type = response.headers.get("Content-Type", "")

    if "text/html" in content_type:

        html = response.get_data(as_text=True)

        global_navigation_script = """
<script>
(function () {

    if (window.__globalNavigationInstalled) {
        return;
    }

    window.__globalNavigationInstalled = true;


    // =====================================================
    // POST FORMS
    // =====================================================

    document.addEventListener("submit", function (event) {

        const form = event.target;

        if (!(form instanceof HTMLFormElement)) {
            return;
        }

        const method = (
            form.getAttribute("method") || "GET"
        ).toLowerCase();

        if (method !== "post") {
            return;
        }

        if (
            form.target &&
            form.target !== "_self"
        ) {
            return;
        }

        const action = (
            form.getAttribute("action") ||
            window.location.href
        );

        const url = new URL(
            action,
            window.location.href
        );

        if (
            url.origin !== window.location.origin
        ) {
            return;
        }

        event.preventDefault();

        const submitButton = form.querySelector(
            "button[type='submit'], input[type='submit']"
        );

        if (submitButton) {
            submitButton.disabled = true;
        }

        fetch(
            url.href,
            {
                method: "POST",
                body: new FormData(form),
                credentials: "same-origin",
                redirect: "follow"
            }
        )
        .then(function (response) {

            if (response.redirected) {

                // History entry create nahi hogi
                window.location.replace(
                    response.url
                );

                return;
            }

            return response.text();

        })
        .then(function (html) {

            if (!html) {
                return;
            }

            document.open();
            document.write(html);
            document.close();

        })
        .catch(function () {

            form.submit();

        });

    }, true);



    // =====================================================
    // DELETE / RESET GET LINKS
    // =====================================================

    document.addEventListener("click", function (event) {

        const link = event.target.closest("a");

        if (!link) {
            return;
        }

        const href = link.getAttribute("href");

        if (!href) {
            return;
        }

        const url = new URL(
            href,
            window.location.href
        );


        // Sirf same website ke links
        if (
            url.origin !== window.location.origin
        ) {
            return;
        }


        const path = url.pathname;


        // =================================================
        // DATA DELETE / RESET ACTIONS
        // =================================================

        const isActionLink =
            path.startsWith("/delete-") ||
            path.startsWith("/reset-");


        if (!isActionLink) {
            return;
        }


        // Normal browser navigation rok do
        event.preventDefault();


        // Action request
        fetch(
            url.href,
            {
                method: "GET",
                credentials: "same-origin",
                redirect: "follow"
            }
        )
        .then(function (response) {

            /*
             * replace() use karne se
             * Delete/Reset ke liye
             * extra history entry nahi banegi.
             */

            window.location.replace(
                response.url
            );

        })
        .catch(function () {

            // Agar fetch fail ho to normal link
            window.location.href = url.href;

        });

    }, true);

})();
</script>
"""


        # </body> se pehle script insert karo
        lower_html = html.lower()

        body_position = lower_html.rfind(
            "</body>"
        )

        if body_position != -1:

            html = (
                html[:body_position]
                + global_navigation_script
                + html[body_position:]
            )

            response.set_data(html)


    return response

# =========================
# FOLDERS
# =========================

UPLOAD_FOLDER = "uploads"
SYLLABUS_FOLDER = "uploads/syllabus"

app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
app.config["SYLLABUS_FOLDER"] = SYLLABUS_FOLDER

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(SYLLABUS_FOLDER, exist_ok=True)

# =========================
# ALLOWED FILES
# =========================

ALLOWED_EXTENSIONS = {"pdf", "jpg", "jpeg", "png"}


def allowed_file(filename):
    return (
        "." in filename
        and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS
    )


NOTES_FILE = "notes.json"
SYLLABUS_FILE = "syllabus.json"
SUBJECTS_FILE = "subjects.json"
DEPARTMENTS_FILE = "departments.json"
COURSES_FILE = "courses.json"

def get_admin_courses():

    courses = []

    if os.path.exists(COURSES_FILE):

        try:
            with open(COURSES_FILE, "r") as file:
                courses = json.load(file)

        except:
            courses = []

    return courses



@app.route("/admin-subjects", methods=["GET", "POST"])
def admin_subjects():

    # =========================
    # ADMIN LOGIN CHECK
    # =========================

    security_check = admin_required()

    if security_check:
        return security_check

    subjects = []
    courses = []

    # =========================
    # SUBJECTS LOAD
    # =========================

    if os.path.exists(SUBJECTS_FILE):

        try:
            with open(SUBJECTS_FILE, "r") as file:
                subjects = json.load(file)

            if not isinstance(subjects, list):
                subjects = []

        except:
            subjects = []


    # =========================
    # COURSES LOAD
    # Course Management se
    # =========================

    courses = get_admin_courses()


    # =========================
    # ADD SUBJECT
    # =========================

    if request.method == "POST":

        subject_name = request.form.get(
            "subject_name", ""
        ).strip()

        course = request.form.get(
            "course", ""
        ).strip()

        semester = request.form.get(
            "semester", ""
        ).strip()


        # =========================
        # VALIDATION
        # =========================

        if not subject_name or not course or not semester:

            return "Please fill all fields."


        if semester not in [
            "1", "2", "3", "4",
            "5", "6", "7", "8", "9", "10"
        ]:

            return "Invalid Semester."


        # =========================
        # CHECK COURSE
        # Admin Course Management
        # se hi course allowed hoga
        # =========================

        valid_course = None

        for item in courses:

            if isinstance(item, dict):

                course_name = str(
                    item.get("name")
                    or item.get("course_name")
                    or item.get("course")
                    or ""
                ).strip()

            else:

                course_name = str(item).strip()


            if course_name.lower() == course.lower():

                valid_course = course_name
                break


        if not valid_course:

            return "Invalid Course."


        # =========================
        # DUPLICATE CHECK
        # =========================

        for subject in subjects:

            existing_course = str(
                subject.get("course", "")
            ).strip()

            # Purane records ke liye
            # department ko ignore kiya jayega

            if (
                str(
                    subject.get(
                        "subject_name", ""
                    )
                ).strip().lower()
                == subject_name.lower()

                and existing_course.lower()
                == valid_course.lower()

                and str(
                    subject.get(
                        "semester", ""
                    )
                ).strip()
                == semester
            ):

                return (
                    "This Subject already exists "
                    "for this Course and Semester."
                )


        # =========================
        # SUBJECT ID
        # =========================

        existing_ids = []

        for subject in subjects:

            existing_ids.append(
                str(
                    subject.get(
                        "subject_id", ""
                    )
                ).strip()
            )


        number = len(subjects) + 1

        subject_id = (
            "SUB" + str(number).zfill(3)
        )


        while subject_id in existing_ids:

            number += 1

            subject_id = (
                "SUB" + str(number).zfill(3)
            )


        # =========================
        # SUBJECT DATA
        # =========================

        subject = {

            "subject_id":
                subject_id,

            "subject_name":
                subject_name,

            "course":
                valid_course,

            "semester":
                semester
        }


        subjects.append(subject)


        # =========================
        # SAVE
        # =========================

        with open(
            SUBJECTS_FILE,
            "w"
        ) as file:

            json.dump(
                subjects,
                file,
                indent=4
            )


        return redirect(
            "/admin-subjects"
        )


    # =========================
    # SHOW PAGE
    # =========================

    return render_template(
        "admin_subjects.html",
        subjects=subjects,
        courses=courses
    )

@app.route("/delete-subject/<subject_id>")
def delete_subject(subject_id):

    # =========================
    # ADMIN LOGIN CHECK
    # =========================

    admin_data = get_admin_data()

    if session.get("admin_id") != admin_data.get("admin_id"):
        session.pop("admin_id", None)
        return redirect(url_for("admin_login"))


    subjects = []

    # Purane subjects read karo
    if os.path.exists(SUBJECTS_FILE):

        try:
            with open(SUBJECTS_FILE, "r") as file:
                subjects = json.load(file)

        except:
            subjects = []

    # Selected subject ko remove karo
    updated_subjects = []

    for subject in subjects:

        if str(subject.get("subject_id")).strip() != str(subject_id).strip():
            updated_subjects.append(subject)

    # Updated list save karo
    with open(SUBJECTS_FILE, "w") as file:

        json.dump(
            updated_subjects,
            file,
            indent=4
        )

    return redirect("/admin-subjects")

# =========================
# MAIN PAGE
# =========================

@app.route("/")
def home():

    return render_template(
        "index.html"
    )


# =========================
# STUDENT REGISTRATION QR CODE
# =========================
# QR ko direct PNG response ke through serve karte hain.
# Isse browser me broken data-image issue nahi hoga.
@app.route("/student-registration-qr")
def student_registration_qr():

    registration_url = (
        "https://mangalayatan-university-portal-2.onrender.com"
        "/student-register"
    )

    qr = qrcode.make(registration_url)

    qr_buffer = io.BytesIO()

    qr.save(
        qr_buffer,
        format="PNG"
    )

    qr_buffer.seek(0)

    return send_file(
        qr_buffer,
        mimetype="image/png",
        download_name="student-registration-qr.png"
    )


@app.route("/student-register", methods=["GET", "POST"])
def student_register():

    courses = get_admin_courses()

    departments = []

    if os.path.exists(DEPARTMENTS_FILE):

        try:

            with open(
                DEPARTMENTS_FILE,
                "r"
            ) as file:

                departments = json.load(file)

            if not isinstance(
                departments,
                list
            ):

                departments = []

        except:

            departments = []

    students = []

    # -----------------------------
    # EXISTING STUDENTS LOAD
    # -----------------------------

    if os.path.exists(STUDENTS_FILE):

        try:

            with open(
                STUDENTS_FILE,
                "r"
            ) as file:

                students = json.load(file)

            if not isinstance(
                students,
                list
            ):

                students = []

        except:

            students = []

    # -----------------------------
    # REGISTRATION
    # -----------------------------

    if request.method == "POST":

        student_name = request.form.get(
            "student_name",
            ""
        ).strip()

        enrollment = request.form.get(
            "enrollment",
            ""
        ).strip()

        department = request.form.get(
            "department",
            ""
        ).strip()

        course = request.form.get(
            "course",
            ""
        ).strip()

        semester = request.form.get(
            "semester",
            ""
        ).strip()

        branch = request.form.get(
            "branch",
            ""
        ).strip().upper()

        section = request.form.get(
            "section",
            ""
        ).strip().upper()

        group = request.form.get(
            "group",
            ""
        ).strip().upper()

        # -------------------------
        # DCEA SEMESTER 1/2 ADMISSION FLOW
        # -------------------------

        admission_type = request.form.get(
            "admission_type",
            ""
        ).strip()

        admission_date = request.form.get(
            "admission_date",
            ""
        ).strip()

        application_no = request.form.get(
            "application_no",
            ""
        ).strip()

        admission_card = request.files.get(
            "admission_card"
        )

        is_dcea_sem_1_2 = (
            department.strip().lower() == "dcea"
            and semester in ["1", "2"]
        )

        # -------------------------
        # REQUIRED FIELDS
        # -------------------------
        # Section and Group are optional.

        if (
            not student_name
            or not enrollment
            or not department
            or not course
            or not semester
        ):

            return "Please fill all required fields."

        # -------------------------
        # DEPARTMENT CHECK
        # -------------------------

        valid_department = None

        for item in departments:

            if isinstance(
                item,
                dict
            ):

                department_name = str(
                    item.get(
                        "name",
                        ""
                    )
                ).strip()

                if (
                    department_name.lower()
                    == department.lower()
                ):

                    valid_department = (
                        department_name
                    )

                    break

        if not valid_department:

            return "Invalid Department."

        department = valid_department

        # -------------------------
        # SEMESTER CHECK
        # -------------------------

        if semester not in [
            "1", "2", "3", "4",
            "5", "6", "7", "8",
            "9", "10"
        ]:

            return "Invalid Semester."

        # -------------------------
        # COURSE CHECK
        # -------------------------

        valid_course = None

        for item in courses:

            if isinstance(
                item,
                dict
            ):

                course_name = str(
                    item.get(
                        "name",
                        ""
                    )
                ).strip()

                if (
                    course_name.lower()
                    == course.lower()
                ):

                    valid_course = course_name

                    break

        if not valid_course:

            return "Invalid Course."

        course = valid_course

        # -------------------------
        # BRANCH VALIDATION
        # -------------------------
        # Branch is ONLY required for
        # B.TECH ALL + Semester 1 or 2.

        branches = [

            "CSE",

            "MECHANICAL",

            "MECHANICAL WITH AI & ML",

            "ELECTRICAL",

            "CIVIL",

            "ECE",

            "CSE WITH AI & ML",

            "CSE WITH CYBER SECURITY",

            "CSE WITH DATA SCIENCE"

        ]

        is_btech_all = (
            course.strip().lower()
            == "b.tech all"
        )

        if (
            is_btech_all
            and semester in ["1", "2"]
        ):

            if not branch:

                return "Please select Branch."

            if branch not in branches:

                return "Invalid Branch."

        else:

            # Other courses and semesters
            # do not use branch.

            branch = ""

        # -------------------------
        # SECTION CHECK
        # -------------------------
        # Section is optional.

        if section and section not in [

            "A",
            "B",
            "C",
            "D"

        ]:

            return "Invalid Section."

        # -------------------------
        # GROUP CHECK
        # -------------------------
        # Group is optional.

        if group and group not in [

            "G1",
            "G2",
            "G3",
            "G4"

        ]:

            return "Invalid Group."

        # -------------------------
        # DUPLICATE ENROLLMENT
        # -------------------------

        for student in students:

            old_enrollment = str(
                student.get(
                    "enrollment",
                    ""
                )
            ).strip()

            if (
                old_enrollment.lower()
                == enrollment.lower()
            ):

                return (
                    "This Enrollment Number "
                    "already exists."
                )


            # -------------------------
        # OPEN ADMISSION PROGRAM PAGE
        # -------------------------

        if (
            is_dcea_sem_1_2
            and not admission_type
        ):

            return render_template(
                "student_admission_type.html",

                form_action="/student-register",

                student_name=student_name,

                enrollment=enrollment,

                department=department,

                course=course,

                semester=semester,

                branch=branch,

                section=section,

                group=group
            )

        # -------------------------
        # DCEA ADMISSION VALIDATION
        # -------------------------

        if is_dcea_sem_1_2:

            if admission_type not in [
                "orientation",
                "after_orientation"
            ]:

                return (
                    "Invalid admission type.",
                    400
                )

            if not admission_date:

                return (
                    "Please fill Date of Admission.",
                    400
                )

            try:

                from datetime import date

                admission_date_obj = (
                    date.fromisoformat(
                        admission_date
                    )
                )

            except ValueError:

                return (
                    "Invalid admission date.",
                    400
                )

            orientation_start = date(
                2026,
                6,
                1
            )

            orientation_end = date(
                2026,
                8,
                23
            )

            after_orientation_start = date(
                2026,
                8,
                24
            )

            # Before 1 June 2026

            if (
                admission_date_obj
                < orientation_start
            ):

                return (
                    "Admission date cannot be before 1 June 2026.",
                    400
                )

            # Orientation Program

            if admission_type == "orientation":

                if (
                    admission_date_obj
                    >= after_orientation_start
                ):

                    return (
                        "Please go to After Orientation Program side.",
                        400
                    )

            # After Orientation Program

            if admission_type == "after_orientation":

                if (
                    admission_date_obj
                    <= orientation_end
                ):

                    return (
                        "Please use Orientation Program side for this date.",
                        400
                    )

                if not application_no:

                    return (
                        "Please enter Application No.",
                        400
                    )

                if not admission_card:

                    return (
                        "Please upload Admission Card PDF.",
                        400
                    )

                if not admission_card.filename:

                    return (
                        "Please upload Admission Card PDF.",
                        400
                    )

                if not admission_card.filename.lower().endswith(
                    ".pdf"
                ):

                    return (
                        "Only PDF Admission Card is allowed.",
                        400
                    )

        else:

            # Existing QR registrations:
            # No admission flow.

            admission_type = ""

            admission_date = ""

            application_no = ""

            admission_card = None

        # -------------------------
        # GENERATE STUDENT ID
        # -------------------------

        numbers = []

        for student in students:

            old_id = str(
                student.get(
                    "student_id",
                    ""
                )
            ).strip().upper()

            if old_id.startswith("STU"):

                try:

                    number = int(
                        old_id[3:]
                    )

                    numbers.append(number)

                except:

                    pass

        if numbers:

            next_number = max(
                numbers
            ) + 1

        else:

            next_number = 1

        student_id = (
            "STU"
            + str(
                next_number
            ).zfill(3)
        )

        # -------------------------
        # CREATE STUDENT
        # -------------------------

        student = {

            "student_id": student_id,

            "name": student_name,

            "enrollment": enrollment,

            "department": department,

            "course": course,

            "semester": semester,

            "branch": branch,

            "section": section,

            "group": group,

            "admission_type": admission_type,

            "admission_date": admission_date,

            "application_no": application_no,

            "admission_card": ""

        }

        # -------------------------
        # SAVE ADMISSION CARD
        # -------------------------

        if (
            is_dcea_sem_1_2
            and admission_type == "after_orientation"
        ):

            try:

                card_filename = (
                    "admission_card_"
                    + student_id
                    + ".pdf"
                )

                save_uploaded_file(
                    admission_card,
                    card_filename
                )

                student["admission_card"] = (
                    card_filename
                )

            except Exception as e:

                return (
                    f"Admission Card could not be saved: {e}",
                    500
                )

        students.append(student)

        # -------------------------
        # SAVE
        # -------------------------

        with open(
            STUDENTS_FILE,
            "w"
        ) as file:

            json.dump(
                students,
                file,
                indent=4
            )

        # -------------------------
        # SUCCESS
        # -------------------------

        return (
            "<h2>Registration Successful!</h2>"
            "<p>Your Student ID is: "
            + student_id
            + "</p>"
            "<p>Username: Enrollment Number</p>"
            "<p>Password: Student ID</p>"
            "<p><a href='/student-login'>"
            "Go to Student Login"
            "</a></p>"
        )

    # -----------------------------
    # REGISTRATION PAGE
    # -----------------------------

    return render_template(
        "student_register.html",
        courses=courses,
        departments=departments,
        semesters=[
            "1", "2", "3", "4",
            "5", "6", "7", "8",
            "9", "10"
        ],
        branches=[
            "CSE",
            "MECHANICAL",
            "MECHANICAL WITH AI & ML",
            "ELECTRICAL",
            "CIVIL",
            "ECE",
            "CSE WITH AI & ML",
            "CSE WITH CYBER SECURITY",
            "CSE WITH DATA SCIENCE"
        ],
        sections=[
            "A", "B", "C", "D"
        ],
        groups=[
            "G1", "G2", "G3", "G4"
        ]
    )


# =========================================================
# STUDENT LOGIN
# =========================================================

@app.route("/student-login", methods=["GET", "POST"])
def student_login():

    if request.method == "POST":

        enrollment = request.form.get(
            "student_id",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        )

        students = []

        # =================================================
        # LOAD STUDENTS
        # =================================================

        if os.path.exists(STUDENTS_FILE):

            try:

                with open(
                    STUDENTS_FILE,
                    "r",
                    encoding="utf-8"
                ) as file:

                    students = json.load(file)

                if not isinstance(
                    students,
                    list
                ):
                    students = []

            except Exception:

                students = []

        # =================================================
        # FIND STUDENT BY ENROLLMENT
        # =================================================

        student = None

        for item in students:

            if not isinstance(
                item,
                dict
            ):
                continue

            old_enrollment = str(
                item.get(
                    "enrollment",
                    ""
                )
            ).strip()

            if (
                old_enrollment.lower()
                ==
                enrollment.lower()
            ):

                student = item
                break

        # =================================================
        # INVALID STUDENT
        # =================================================

        if student is None:

            return (
                "Invalid Enrollment Number "
                "or Password"
            )

        # =================================================
        # STUDENT ID
        # =================================================

        student_id = str(
            student.get(
                "student_id",
                ""
            )
        ).strip()

        if not student_id:

            return (
                "Invalid Enrollment Number "
                "or Password"
            )

        # =================================================
        # PASSWORD CHECK
        #
        # IMPORTANT:
        #
        # If password_hash already exists:
        # ONLY the saved password is accepted.
        #
        # Student ID will NOT work anymore.
        #
        # If password_hash does NOT exist:
        # Student ID is used as the FIRST/default password.
        # =================================================

        stored_hash = str(
            student.get(
                "password_hash",
                ""
            )
        ).strip()

        login_success = False

        # -------------------------------------------------
        # EXISTING PASSWORD
        # -------------------------------------------------

        if stored_hash:

            try:

                if check_password_hash(
                    stored_hash,
                    password
                ):

                    login_success = True

            except Exception:

                login_success = False

        # -------------------------------------------------
        # FIRST-TIME LOGIN
        #
        # Only if there is NO password_hash.
        # -------------------------------------------------

        else:

            if password == student_id:

                student[
                    "password_hash"
                ] = generate_password_hash(
                    student_id
                )

                try:

                    with open(
                        STUDENTS_FILE,
                        "w",
                        encoding="utf-8"
                    ) as file:

                        json.dump(
                            students,
                            file,
                            indent=4
                        )

                    login_success = True

                except Exception:

                    login_success = False

        # =================================================
        # FINAL PASSWORD CHECK
        # =================================================

        if not login_success:

            return (
                "Invalid Enrollment Number "
                "or Password"
            )

        # =================================================
        # DEPARTMENT LOCK CHECK
        # =================================================

        student_department = str(
            student.get(
                "department",
                ""
            )
        ).strip()

        conn = get_connection()

        try:

            with conn.cursor() as cur:

                cur.execute("""
                    SELECT is_closed
                    FROM department_admins
                    WHERE LOWER(department_name) = LOWER(%s)
                """, (
                    student_department,
                ))

                department_data = cur.fetchone()

        finally:

            conn.close()

        # Department Owner ke system me nahi hai
        # OR Department temporarily closed hai

        if (
            not department_data
            or department_data["is_closed"]
        ):

            return (
                "Your service is locked. "
                "First payment by department then use this service."
            )

        # =================================================
        # ROLE SESSION ISOLATION
        # =================================================

        session.pop(
            "admin_id",
            None
        )

        session.pop(
            "faculty_id",
            None
        )

        session.pop(
            "faculty_name",
            None
        )

        # =================================================
        # STUDENT SESSION
        # =================================================

        session[
            "student_id"
        ] = student_id

        session[
            "student_enrollment"
        ] = str(
            student.get(
                "enrollment",
                ""
            )
        ).strip()

        session[
            "student_course"
        ] = str(
            student.get(
                "course",
                ""
            )
        ).strip()

        session[
            "student_semester"
        ] = str(
            student.get(
                "semester",
                ""
            )
        ).strip()

        session[
            "student_section"
        ] = str(
            student.get(
                "section",
                ""
            )
        ).strip().upper()

        # =================================================
        # OPEN STUDENT DASHBOARD
        # =================================================

        return redirect(
            "/student-dashboard"
        )

    # =====================================================
    # GET - LOGIN PAGE
    # =====================================================

    return render_template(
        "student_login.html"
    )


# =========================
# STUDENT DASHBOARD
# =========================

@app.route("/student-dashboard")
def student_dashboard():

    # Student login check
    if "student_enrollment" not in session:
        return redirect(url_for("student_login"))

    student_enrollment = str(
        session.get("student_enrollment", "")
    ).strip()

    # Student data load
    students = []

    if os.path.exists(STUDENTS_FILE):
        try:
            with open(STUDENTS_FILE, "r") as file:
                students = json.load(file)
        except:
            students = []

    # Logged-in student find
    student = None

    for item in students:
        if str(
            item.get("enrollment", "")
        ).strip() == student_enrollment:

            student = item
            break

    if student is None:
        return "Student not found.", 404

    return render_template(
        "student_dashboard.html",
        student=student
    )

    # -------------------------
    # LOGIN CHECK
    # -------------------------

    if not session.get("student_id"):

        return redirect(
            "/student-login"
        )

    # -------------------------
    # STUDENT DATA LOAD
    # -------------------------

    students = []

    if os.path.exists(STUDENTS_FILE):

        try:

            with open(
                STUDENTS_FILE,
                "r"
            ) as file:

                students = json.load(file)

        except:

            students = []

    # -------------------------
    # CURRENT STUDENT
    # -------------------------

    student = None

    current_student_id = str(
        session.get(
            "student_id",
            ""
        )
    ).strip()

    for item in students:

        if str(
            item.get(
                "student_id",
                ""
            )
        ).strip() == current_student_id:

            student = item
            break

    # -------------------------
    # STUDENT NOT FOUND
    # -------------------------

    if student is None:

        session.pop(
            "student_id",
            None
        )

        session.pop(
            "student_enrollment",
            None
        )

        session.pop(
            "student_course",
            None
        )

        session.pop(
            "student_semester",
            None
        )

        session.pop(
            "student_section",
            None
        )

        return redirect(
            "/student-login"
        )

    # -------------------------
    # DASHBOARD
    # -------------------------

    return render_template(
        "student_dashboard.html",
        student=student
    )


   


# =========================
# STUDENT SYLLABUS
# =========================

@app.route("/syllabus")
def syllabus():

    # Student login check
    if "student_enrollment" not in session:
        return redirect(url_for("student_login"))

    student_enrollment = str(
        session.get("student_enrollment", "")
    ).strip()

    # =========================
    # STUDENTS LOAD
    # =========================

    students = []

    if os.path.exists(STUDENTS_FILE):
        try:
            with open(STUDENTS_FILE, "r") as file:
                students = json.load(file)
        except:
            students = []

    # =========================
    # CURRENT STUDENT FIND
    # =========================

    student = None

    for item in students:
        if (
            str(item.get("enrollment", "")).strip()
            == student_enrollment
        ):
            student = item
            break

    if not student:
        return "Student not found", 404

    # =========================
    # STUDENT DETAILS
    # =========================

    student_course = str(
        student.get("course", "")
    ).strip().lower()

    student_semester = str(
        student.get("semester", "")
    ).strip()

    student_section = str(
        student.get("section", "")
    ).strip().upper()

    # =========================
    # SYLLABUS LOAD
    # =========================

    all_syllabuses = []

    if os.path.exists(SYLLABUS_FILE):
        try:
            with open(SYLLABUS_FILE, "r") as file:
                all_syllabuses = json.load(file)
        except:
            all_syllabuses = []

    # =========================
    # FILTER SYLLABUS
    # Course + Semester + Section
    # =========================

    syllabuses = []

    for syllabus_item in all_syllabuses:

        syllabus_course = str(
            syllabus_item.get("course", "")
        ).strip().lower()

        syllabus_semester = str(
            syllabus_item.get("semester", "")
        ).strip()

        syllabus_section = str(
            syllabus_item.get("section", "")
        ).strip().upper()

        if (
            syllabus_course == student_course
            and syllabus_semester == student_semester
            and syllabus_section == student_section
        ):
            syllabuses.append(syllabus_item)

    # =========================
    # SHOW STUDENT SYLLABUS
    # =========================

    return render_template(
        "syllabus.html",
        syllabuses=syllabuses,
        student=student
    )


# =========================
# OPEN SYLLABUS FILE
# =========================

@app.route("/syllabus-file/<filename>")
def syllabus_file(filename):

    stored = get_uploaded_file(filename)

    if not stored:
        return "File not found", 404

    return send_file(
    io.BytesIO(bytes(stored["content"])),
    mimetype=(
        "image/jpeg"
        if filename.lower().endswith((".jpg", ".jpeg"))
        else "image/png"
        if filename.lower().endswith(".png")
        else stored["content_type"]
    ),
    as_attachment=False,
    download_name=filename
)


# =========================
# STUDENT ATTENDANCE
# =========================

@app.route("/attendance")
def attendance():

    # Student login check
    if "student_enrollment" not in session:
        return redirect(url_for("student_login"))

    student_enrollment = str(
        session.get("student_enrollment", "")
    ).strip()

    # =========================
    # STUDENTS LOAD
    # =========================

    students = []

    if os.path.exists(STUDENTS_FILE):
        try:
            with open(STUDENTS_FILE, "r") as file:
                students = json.load(file)
        except:
            students = []

    # =========================
    # CURRENT STUDENT FIND
    # =========================

    student = None

    for item in students:
        if (
            str(item.get("enrollment", "")).strip()
            == student_enrollment
        ):
            student = item
            break

    if not student:
        return "Student not found", 404

    # =========================
    # STUDENT DETAILS
    # =========================

    student_course = str(
        student.get("course", "")
    ).strip().lower()

    student_semester = str(
        student.get("semester", "")
    ).strip()

    student_section = str(
        student.get("section", "")
    ).strip().upper()

    # =========================
    # ATTENDANCE LOAD
    # =========================

    all_attendance = []

    if os.path.exists("attendance.json"):
        try:
            with open("attendance.json", "r") as file:
                all_attendance = json.load(file)
        except:
            all_attendance = []

    # =========================
    # CURRENT STUDENT ATTENDANCE
    # =========================

    student_attendance = []

    for record in all_attendance:

        record_course = str(
            record.get("course", "")
        ).strip().lower()

        record_semester = str(
            record.get("semester", "")
        ).strip()

        record_section = str(
            record.get("section", "")
        ).strip().upper()

        record_student = str(
            record.get("student_id", "")
        ).strip()

        if (
            record_student == student_enrollment
            and record_course == student_course
            and record_semester == student_semester
            and record_section == student_section
        ):
            student_attendance.append(record)

    # =========================
    # SUBJECT WISE ATTENDANCE
    # =========================

    subject_attendance = {}

    total_classes = 0
    total_present = 0

    for record in student_attendance:

        subject = str(
            record.get("subject", "Unknown")
        ).strip()

        status = str(
            record.get("status", "")
        ).strip()

        if subject not in subject_attendance:
            subject_attendance[subject] = {
                "total": 0,
                "present": 0,
                "absent": 0,
                "percentage": 0
            }

        subject_attendance[subject]["total"] += 1
        total_classes += 1

        if status == "Present":
            subject_attendance[subject]["present"] += 1
            total_present += 1
        else:
            subject_attendance[subject]["absent"] += 1

    # =========================
    # PERCENTAGE
    # =========================

    for subject in subject_attendance:

        data = subject_attendance[subject]

        if data["total"] > 0:
            data["percentage"] = round(
                (data["present"] / data["total"]) * 100,
                2
            )

    if total_classes > 0:
        attendance_percentage = round(
            (total_present / total_classes) * 100,
            2
        )
    else:
        attendance_percentage = 0

    # =========================
    # SHOW ATTENDANCE
    # =========================

    return render_template(
        "attendance.html",
        student=student,
        subject_attendance=subject_attendance,
        total_classes=total_classes,
        total_present=total_present,
        attendance_percentage=attendance_percentage,
        daily_history=student_attendance
    )

# =========================
# MANAGE ATTENDANCE
# =========================

@app.route("/manage-attendance", methods=["GET", "POST"])
def manage_attendance():

    # =========================
    # FACULTY LOGIN SECURITY
    # =========================
    faculty_id = str(
        session.get("faculty_id", "")
    ).strip()

    if not faculty_id:
        return redirect("/faculty-login")

    faculty_name = str(
        session.get("faculty_name", "")
    ).strip()


    # =========================
    # STUDENTS LOAD
    # =========================
    all_students = []

    if os.path.exists(STUDENTS_FILE):
        try:
            with open(STUDENTS_FILE, "r") as file:
                all_students = json.load(file)
        except:
            all_students = []


    # =========================
    # SUBJECTS LOAD
    # =========================
    subjects = []

    if os.path.exists(SUBJECTS_FILE):
        try:
            with open(SUBJECTS_FILE, "r") as file:
                subjects = json.load(file)
        except:
            subjects = []


    # =========================
    # COURSES FROM ADMIN
    # =========================
    courses = get_admin_courses()


    # =========================
    # POST - SAVE ATTENDANCE
    # =========================
    if request.method == "POST":

        course = request.form.get(
            "course", ""
        ).strip()

        subject = request.form.get(
            "subject", ""
        ).strip()

        semester = request.form.get(
            "semester", ""
        ).strip()

        section = request.form.get(
            "section", ""
        ).strip()

        date = request.form.get(
            "date", ""
        ).strip()

        present_students = request.form.getlist(
            "present"
        )


        # =========================
        # BASIC VALIDATION
        # =========================
        if (
            not course
            or not subject
            or not semester
            or not section
            or not date
        ):
            return "Please fill all fields."


        # =========================
        # SEMESTER VALIDATION
        # =========================
        if semester not in [
            "1", "2", "3", "4",
            "5", "6", "7", "8", "9", "10"
        ]:
            return "Invalid Semester."


        # =========================
        # SECTION VALIDATION
        # =========================
        if section not in [
            "A", "B", "C", "D"
        ]:
            return "Invalid Section."


        # =========================
        # COURSE VALIDATION
        # =========================
        valid_course = False

        for item in courses:

            if (
                str(item.get("name", "")).strip().lower()
                == course.lower()
            ):
                valid_course = True
                break

        if not valid_course:
            return "Invalid Course."


        # =========================
        # SUBJECT VALIDATION
        # COURSE + SEMESTER
        # =========================
        valid_subject = False

        for item in subjects:

            if (
                str(
                    item.get("subject_name", "")
                ).strip().lower()
                == subject.lower()

                and

                str(
                    item.get("course", "")
                ).strip().lower()
                == course.lower()

                and

                str(
                    item.get("semester", "")
                ).strip()
                == semester
            ):
                valid_subject = True
                break


        if not valid_subject:
            return "Invalid Subject for selected Course and Semester."


        # =========================
        # FIND STUDENTS
        # COURSE + SEMESTER + SECTION
        # =========================
        selected_students = []

        for student in all_students:

            student_course = str(
                student.get("course", "")
            ).strip()

            student_semester = str(
                student.get("semester", "")
            ).strip()

            student_section = str(
                student.get("section", "")
            ).strip().upper()


            if (
                student_course.lower()
                == course.lower()

                and

                student_semester
                == semester

                and

                student_section
                == section.upper()
            ):
                selected_students.append(student)


        # =========================
        # ATTENDANCE FILE LOAD
        # =========================
        attendance_data = []

        if os.path.exists("attendance.json"):

            try:
                with open(
                    "attendance.json",
                    "r"
                ) as file:

                    attendance_data = json.load(file)

            except:
                attendance_data = []

                # =========================
        # DUPLICATE ATTENDANCE CHECK
        # =========================
        for old_record in attendance_data:

            if (
                str(old_record.get("date", "")).strip()
                == date

                and

                str(old_record.get("course", "")).strip().lower()
                == course.lower()

                and

                str(old_record.get("subject", "")).strip().lower()
                == subject.lower()

                and

                str(old_record.get("semester", "")).strip()
                == semester

                and

                str(old_record.get("section", "")).strip().upper()
                == section.upper()
            ):

                return (
                    "You have done present on this day "
                    "for this course, semester, section and subject."
                )

        # =========================
        # SAVE EACH STUDENT
        # =========================
        for student in selected_students:

            enrollment = str(
                student.get("enrollment", "")
            ).strip()

            student_name = str(
                student.get("name", "")
            ).strip()


            if enrollment in present_students:
                status = "Present"
            else:
                status = "Absent"


            record = {
                "faculty_id": faculty_id,
                "faculty_name": faculty_name,

                "date": date,

                "course": course,
                "subject": subject,

                "semester": semester,
                "section": section,

                "student_id": enrollment,
                "student_name": student_name,

                "status": status
            }


            attendance_data.append(record)


        # =========================
        # SAVE ATTENDANCE.JSON
        # =========================
        with open(
            "attendance.json",
            "w"
        ) as file:

            json.dump(
                attendance_data,
                file,
                indent=4
            )


        # =========================
        # AFTER SAVE
        # =========================
        return redirect(
            "/manage-attendance"
        )


    # =========================
    # SHOW ATTENDANCE PAGE
    # =========================
    return render_template(
        "manage_attendance.html",

        students=all_students,

        courses=courses,

        subjects=subjects,

        faculty_name=faculty_name
    )

# =========================
# FACULTY ATTENDANCE HISTORY
# =========================

@app.route("/faculty-attendance-history", methods=["GET", "POST"])
def faculty_attendance_history():

    # =========================
    # CURRENT FACULTY
    # =========================

    faculty_id = str(
        session.get("faculty_id", "")
    ).strip()

    faculty_name = str(
        session.get("faculty_name", "")
    ).strip()

    if not faculty_id:
        return redirect("/faculty-login")


    # =========================
    # ATTENDANCE LOAD
    # =========================

    attendance_data = []

    if os.path.exists("attendance.json"):
        try:
            with open("attendance.json", "r") as file:
                attendance_data = json.load(file)
        except:
            attendance_data = []


    # =========================
    # SELECTED FILTERS
    # =========================

    selected_date = request.values.get(
        "date", ""
    ).strip()

    selected_section = request.values.get(
        "section", ""
    ).strip()

    selected_subject = request.values.get(
        "subject", ""
    ).strip()


    # =========================
    # FACULTY SUBJECTS
    # =========================

    faculty_subjects = []

    for record in attendance_data:

        record_faculty = str(
            record.get("faculty_id", "")
        ).strip()

        if record_faculty != faculty_id:
            continue

        subject = str(
            record.get("subject", "")
        ).strip()

        if subject and subject not in faculty_subjects:
            faculty_subjects.append(subject)

    faculty_subjects.sort()


    # =========================
    # FACULTY SECTIONS
    # =========================

    faculty_sections = []

    for record in attendance_data:

        record_faculty = str(
            record.get("faculty_id", "")
        ).strip()

        if record_faculty != faculty_id:
            continue

        section = str(
            record.get("section", "")
        ).strip()

        if section and section not in faculty_sections:
            faculty_sections.append(section)

    faculty_sections.sort()


    # =========================
    # FILTER ATTENDANCE
    # =========================

    attendance_records = []

    if (
        selected_date
        and selected_section
        and selected_subject
    ):

        for record in attendance_data:

            record_faculty = str(
                record.get("faculty_id", "")
            ).strip()

            record_date = str(
                record.get("date", "")
            ).strip()

            record_section = str(
                record.get("section", "")
            ).strip()

            record_subject = str(
                record.get("subject", "")
            ).strip()


            if (
                record_faculty == faculty_id
                and record_date == selected_date
                and record_section.upper()
                == selected_section.upper()
                and record_subject.lower()
                == selected_subject.lower()
            ):

                attendance_records.append(record)


    # =========================
    # PAGE
    # =========================

    return render_template(
        "faculty_attendance_history.html",
        faculty_name=faculty_name,
        faculty_subjects=faculty_subjects,
        faculty_sections=faculty_sections,
        selected_date=selected_date,
        selected_section=selected_section,
        selected_subject=selected_subject,
        attendance_records=attendance_records
    )

# =========================
# FACULTY ATTENDANCE RECORD
# =========================

def load_attendance_data():

    if not os.path.exists("attendance.json"):
        return []

    try:
        with open("attendance.json", "r") as file:
            data = json.load(file)

        return data if isinstance(data, list) else []

    except Exception:
        return []


def get_faculty_attendance_combinations(attendance_data, faculty_id):

    combinations = []
    seen = set()

    for record in attendance_data:

        if str(record.get("faculty_id", "")).strip() != faculty_id:
            continue

        course = str(record.get("course", "")).strip()
        semester = str(record.get("semester", "")).strip()
        section = str(record.get("section", "")).strip().upper()
        subject = str(record.get("subject", "")).strip()

        if not course or not semester or not section or not subject:
            continue

        key = (
            course.lower(),
            semester,
            section,
            subject.lower()
        )

        if key in seen:
            continue

        seen.add(key)

        combinations.append({
            "course": course,
            "semester": semester,
            "section": section,
            "subject": subject
        })

    combinations.sort(
        key=lambda item: (
            item["course"].lower(),
            int(item["semester"]) if item["semester"].isdigit() else 99,
            item["section"],
            item["subject"].lower()
        )
    )

    return combinations


def build_faculty_attendance_record(
    attendance_data,
    faculty_id,
    course,
    semester,
    section,
    subject,
    end_date
):

    """
    Build cumulative attendance for the CURRENT student roster of the
    selected Course + Semester + Section.

    Student order follows students.json exactly. This keeps the Attendance
    Record/Excel order the same as the Admin Manage Students roster instead
    of sorting students alphabetically.

    Attendance itself is restricted to the logged-in faculty and the exact
    Course + Semester + Section + Subject combination.
    """

    # ---------------------------------------------------------
    # 1. Get the current roster for the selected section.
    #    This is the source of truth for which students belong in
    #    the report and also preserves their existing order.
    # ---------------------------------------------------------
    current_roster = []

    if os.path.exists(STUDENTS_FILE):
        try:
            with open(STUDENTS_FILE, "r") as file:
                loaded_students = json.load(file)
        except Exception:
            loaded_students = []

        if isinstance(loaded_students, list):
            seen_enrollments = set()

            for student in loaded_students:
                student_course = str(student.get("course", "")).strip()
                student_semester = str(student.get("semester", "")).strip()
                student_section = str(student.get("section", "")).strip().upper()
                enrollment = str(student.get("enrollment", "")).strip()
                name = str(student.get("name", "")).strip()

                if (
                    student_course.lower() == course.lower()
                    and student_semester == semester
                    and student_section == section.upper()
                    and enrollment
                    and enrollment.lower() not in seen_enrollments
                ):
                    current_roster.append({
                        "enrollment": enrollment,
                        "name": name
                    })
                    seen_enrollments.add(enrollment.lower())

    # ---------------------------------------------------------
    # 2. Find only attendance taken by this faculty for the exact
    #    selected combination and up to the selected end date.
    # ---------------------------------------------------------
    matching = []

    for record in attendance_data:

        if str(record.get("faculty_id", "")).strip() != faculty_id:
            continue

        if str(record.get("course", "")).strip().lower() != course.lower():
            continue

        if str(record.get("semester", "")).strip() != semester:
            continue

        if str(record.get("section", "")).strip().upper() != section.upper():
            continue

        if str(record.get("subject", "")).strip().lower() != subject.lower():
            continue

        record_date = str(record.get("date", "")).strip()

        if not record_date:
            continue

        if record_date > end_date:
            continue

        matching.append(record)

    if not matching:
        return None

    dates = sorted({
        str(record.get("date", "")).strip()
        for record in matching
        if str(record.get("date", "")).strip()
    })

    if not dates:
        return None

    start_date = dates[0]

    # ---------------------------------------------------------
    # 3. For duplicate student/date entries, latest saved status wins.
    # ---------------------------------------------------------
    latest = {}

    for record in matching:
        enrollment = str(record.get("student_id", "")).strip()
        record_date = str(record.get("date", "")).strip()

        if not enrollment or not record_date:
            continue

        latest[(enrollment.lower(), record_date)] = record

    # ---------------------------------------------------------
    # 4. Build rows ONLY from the current selected-section roster.
    #    Historical records for students who are no longer in this
    #    section will not incorrectly appear in the report.
    # ---------------------------------------------------------
    total_lectures = len(dates)
    rows = []

    for student in current_roster:

        enrollment = student["enrollment"]
        name = student["name"]

        present = 0
        absent = 0

        for lecture_date in dates:

            record = latest.get((enrollment.lower(), lecture_date))

            if record is None:
                # Attendance was not saved for this student on this
                # lecture date, so count it as absent.
                absent += 1
                continue

            status = str(record.get("status", "")).strip().lower()

            if status == "present":
                present += 1
            else:
                absent += 1

        percentage = (
            round((present / total_lectures) * 100, 2)
            if total_lectures
            else 0
        )

        rows.append({
            "enrollment": enrollment,
            "name": name,
            "total_lectures": total_lectures,
            "present": present,
            "absent": absent,
            "percentage": percentage
        })

    return {
        "start_date": start_date,
        "end_date": end_date,
        "course": course,
        "semester": semester,
        "section": section.upper(),
        "subject": subject,
        "rows": rows,
        "total_lectures": total_lectures
    }


@app.route("/faculty-attendance-record", methods=["GET", "POST"])
def faculty_attendance_record():

    faculty_id = str(session.get("faculty_id", "")).strip()
    faculty_name = str(session.get("faculty_name", "")).strip()

    if not faculty_id:
        return redirect("/faculty-login")

    attendance_data = load_attendance_data()

    combinations = get_faculty_attendance_combinations(
        attendance_data,
        faculty_id
    )

    selected_course = request.values.get("course", "").strip()
    selected_semester = request.values.get("semester", "").strip()
    selected_section = request.values.get("section", "").strip().upper()
    selected_subject = request.values.get("subject", "").strip()
    selected_end_date = request.values.get("end_date", "").strip()

    if not selected_end_date:
        from datetime import date
        selected_end_date = date.today().isoformat()

    record = None
    error = ""

    if selected_course or selected_semester or selected_section or selected_subject:

        if not all([
            selected_course,
            selected_semester,
            selected_section,
            selected_subject
        ]):
            error = "Please select Course, Semester, Section and Subject."

        else:

            valid_combination = False

            for item in combinations:

                if (
                    item["course"].lower() == selected_course.lower()
                    and item["semester"] == selected_semester
                    and item["section"].upper() == selected_section.upper()
                    and item["subject"].lower() == selected_subject.lower()
                ):
                    valid_combination = True
                    break

            if not valid_combination:
                error = (
                    "Attendance Record is not available for this selection. "
                    "You can only view attendance taken by your own faculty account."
                )

            else:

                # Future dates are not allowed.
                from datetime import date
                today = date.today().isoformat()

                if selected_end_date > today:
                    selected_end_date = today

                record = build_faculty_attendance_record(
                    attendance_data,
                    faculty_id,
                    selected_course,
                    selected_semester,
                    selected_section,
                    selected_subject,
                    selected_end_date
                )

                if record is None:
                    error = "No attendance data found for this selection."

    return render_template(
        "faculty_attendance_record.html",
        faculty_name=faculty_name,
        combinations=combinations,
        selected_course=selected_course,
        selected_semester=selected_semester,
        selected_section=selected_section,
        selected_subject=selected_subject,
        selected_end_date=selected_end_date,
        record=record,
        error=error
    )


@app.route("/faculty-attendance-record-excel")
def faculty_attendance_record_excel():

    faculty_id = str(session.get("faculty_id", "")).strip()
    faculty_name = str(session.get("faculty_name", "")).strip()

    if not faculty_id:
        return redirect("/faculty-login")

    course = request.args.get("course", "").strip()
    semester = request.args.get("semester", "").strip()
    section = request.args.get("section", "").strip().upper()
    subject = request.args.get("subject", "").strip()
    end_date = request.args.get("end_date", "").strip()

    if not all([course, semester, section, subject, end_date]):
        return "Invalid attendance record request.", 400

    attendance_data = load_attendance_data()

    # Security: exact combination must belong to logged-in faculty.
    combinations = get_faculty_attendance_combinations(
        attendance_data,
        faculty_id
    )

    valid_combination = any(
        item["course"].lower() == course.lower()
        and item["semester"] == semester
        and item["section"].upper() == section.upper()
        and item["subject"].lower() == subject.lower()
        for item in combinations
    )

    if not valid_combination:
        return "Unauthorized attendance record request.", 403

    from datetime import date
    today = date.today().isoformat()

    if end_date > today:
        end_date = today

    record = build_faculty_attendance_record(
        attendance_data,
        faculty_id,
        course,
        semester,
        section,
        subject,
        end_date
    )

    if record is None:
        return "No attendance data found for this selection.", 404

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Attendance Record"

    worksheet["A1"] = "Attendance Record"
    worksheet["A1"].font = Font(bold=True, size=16)

    worksheet["A3"] = "Faculty"
    worksheet["B3"] = faculty_name
    worksheet["A4"] = "Course"
    worksheet["B4"] = record["course"]
    worksheet["A5"] = "Semester"
    worksheet["B5"] = record["semester"]
    worksheet["A6"] = "Section"
    worksheet["B6"] = record["section"]
    worksheet["A7"] = "Subject"
    worksheet["B7"] = record["subject"]
    worksheet["A8"] = "Attendance Period"
    worksheet["B8"] = f'{record["start_date"]} to {record["end_date"]}'
    worksheet["A9"] = "Total Lectures"
    worksheet["B9"] = record["total_lectures"]

    for row in range(3, 10):
        worksheet[f"A{row}"].font = Font(bold=True)

    header_row = 11
    headers = [
        "Student Name",
        "Enrollment",
        "Total Lecture",
        "Present",
        "Absent",
        "Percentage"
    ]

    for col, header in enumerate(headers, start=1):
        cell = worksheet.cell(
            row=header_row,
            column=col,
            value=header
        )
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row_index, student in enumerate(record["rows"], start=header_row + 1):

        worksheet.cell(row=row_index, column=1, value=student["name"])
        worksheet.cell(row=row_index, column=2, value=student["enrollment"])
        worksheet.cell(row=row_index, column=3, value=student["total_lectures"])
        worksheet.cell(row=row_index, column=4, value=student["present"])
        worksheet.cell(row=row_index, column=5, value=student["absent"])
        worksheet.cell(row=row_index, column=6, value=student["percentage"] / 100)
        worksheet.cell(row=row_index, column=6).number_format = "0.00%"

    widths = [25, 20, 16, 12, 12, 15]

    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    worksheet.freeze_panes = "A12"
    worksheet.auto_filter.ref = (
        f"A{header_row}:F{header_row + len(record['rows'])}"
    )

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    safe_course = "".join(
        char if char.isalnum() or char in "-_" else "_"
        for char in course
    )
    safe_subject = "".join(
        char if char.isalnum() or char in "-_" else "_"
        for char in subject
    )

    filename = (
        f"Attendance_Record_{safe_course}_Sem{semester}_"
        f"Section_{section}_{safe_subject}.xlsx"
    )

    return send_file(
        output,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        as_attachment=True,
        download_name=filename
    )


# =========================
# STUDENT NOTES
# =========================

@app.route("/notes")
def student_notes():

    # Student login check
    if "student_enrollment" not in session:
        return redirect(url_for("student_login"))

    # Login ke time save hua enrollment number
    student_enrollment = str(
        session.get("student_enrollment", "")
    ).strip()

    students = []

    if os.path.exists(STUDENTS_FILE):
        try:
            with open(STUDENTS_FILE, "r") as f:
                students = json.load(f)
        except:
            students = []

    # Enrollment number se student find karo
    student = None

    for s in students:

        if str(
            s.get("enrollment", "")
        ).strip() == student_enrollment:

            student = s
            break

    if not student:
        return "Student not found", 404

    # Student ki Course + Semester + Section
    student_course = str(
        student.get("course", "")
    ).strip().lower()

    student_semester = str(
        student.get("semester", "")
    ).strip()

    student_section = str(
        student.get("section", "")
    ).strip().upper()

    notes = []

    if os.path.exists(NOTES_FILE):

        try:
            with open(NOTES_FILE, "r") as f:
                all_notes = json.load(f)
        except:
            all_notes = []

        # Sirf matching Course + Semester + Section
        for note in all_notes:

            note_course = str(
                note.get("course", "")
            ).strip().lower()

            note_semester = str(
                note.get("semester", "")
            ).strip()

            note_section = str(
                note.get("section", "")
            ).strip().upper()

            if (
                note_course == student_course
                and note_semester == student_semester
                and note_section == student_section
            ):
                notes.append(note)

    return render_template(
        "notes.html",
        notes=notes,
        student=student
    )


@app.route("/student-note/<filename>")
def student_note(filename):

    # Student login check
    if "student_enrollment" not in session:
        return redirect(url_for("student_login"))

    return send_from_directory(
        UPLOAD_FOLDER,
        filename
    )

# =========================
# OPEN NOTES / UPLOADED FILE
# =========================

@app.route("/uploads/<filename>")
def uploaded_file(filename):

    # -----------------------------
    # STUDENT ACCESS
    # -----------------------------
    if "student_enrollment" in session:

        student_id = str(
            session.get("student_enrollment", "")
        ).strip()

        students = []

        if os.path.exists(STUDENTS_FILE):

            try:
                with open(STUDENTS_FILE, "r") as f:
                    students = json.load(f)
            except:
                students = []

        student = None

        for s in students:

            if str(
                s.get("enrollment", "")
            ).strip() == student_id:

                student = s
                break

        if not student:
            return "Unauthorized", 403

        student_course = str(
            student.get("course", "")
        ).strip().lower()

        student_semester = str(
            student.get("semester", "")
        ).strip()

        student_section = str(
            student.get("section", "")
        ).strip().upper()

        notes = []

        if os.path.exists(NOTES_FILE):

            try:
                with open(NOTES_FILE, "r") as f:
                    notes = json.load(f)
            except:
                notes = []

        # File ko matching note me dhundo
        for note in notes:

            if note.get("filename") != filename:
                continue

            note_course = str(
                note.get("course", "")
            ).strip().lower()

            note_semester = str(
                note.get("semester", "")
            ).strip()

            note_section = str(
                note.get("section", "")
            ).strip().upper()

            # Student ke Course + Semester + Section ka exact match
            if (
                note_course == student_course
                and note_semester == student_semester
                and note_section == student_section
            ):

                stored = get_uploaded_file(filename)

                if not stored:
                    return "File not found", 404

                return send_file(
                    io.BytesIO(bytes(stored["content"])),
                    mimetype=stored["content_type"],
                    download_name=filename
                )

        return "Unauthorized", 403

    # -----------------------------
    # FACULTY ACCESS
    # -----------------------------
    if "faculty_id" in session:

        faculty_id = str(
            session.get("faculty_id", "")
        ).strip()

        notes = []

        if os.path.exists(NOTES_FILE):

            try:
                with open(NOTES_FILE, "r") as f:
                    notes = json.load(f)
            except:
                notes = []

        for note in notes:

            if (
                note.get("filename") == filename
                and str(
                    note.get("faculty_id", "")
                ).strip() == faculty_id
            ):

                stored = get_uploaded_file(filename)

                if not stored:
                    return "File not found", 404

                return send_file(
                    io.BytesIO(bytes(stored["content"])),
                    mimetype=stored["content_type"],
                    download_name=filename
                )

        return "Unauthorized", 403

    return redirect(url_for("student_login"))


# =========================
# ADMISSION LOGIN
# =========================

@app.route("/admission-login", methods=["GET", "POST"])
def admission_login():

    if request.method == "POST":

        admission_id = request.form.get(
            "admission_id",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        ).strip()

        if (
            admission_id.upper() == "ADM001"
            and password == "admission123"
        ):

            return redirect("/admission-dashboard")

        return "Invalid Admission ID or Password"

    return render_template("admission_login.html")


# =========================
# ADMISSION DASHBOARD
# =========================

@app.route("/admission-dashboard")
def admission_dashboard():

    applications = []

    try:

        with open("applications.json", "r") as file:
            applications = json.load(file)

    except:

        applications = []

    total_applications = len(applications)

    pending = 0
    approved = 0
    rejected = 0

    for application in applications:

        status = application.get(
            "status",
            "Pending"
        )

        if status == "Pending":
            pending += 1

        elif status == "Approved":
            approved += 1

        elif status == "Rejected":
            rejected += 1

    return render_template(
        "admission_dashboard.html",
        total_applications=total_applications,
        pending=pending,
        approved=approved,
        rejected=rejected
    )


# =========================
# NEW ADMISSION
# =========================

@app.route("/new-admission", methods=["GET", "POST"])
def new_admission():

    if request.method == "POST":

        applications = []

        try:

            with open("applications.json", "r") as file:
                applications = json.load(file)

        except:

            applications = []

        application = {

            "application_id":
                "APP" + str(
                    len(applications) + 1
                ).zfill(3),

            "student_name":
                request.form.get("student_name"),

            "father_name":
                request.form.get("father_name"),

            "dob":
                request.form.get("dob"),

            "gender":
                request.form.get("gender"),

            "course":
                request.form.get("course"),

            "mobile":
                request.form.get("mobile"),

            "email":
                request.form.get("email"),

            "address":
                request.form.get("address"),

            "tenth":
                request.form.get("tenth"),

            "twelfth":
                request.form.get("twelfth"),

            "status":
                "Pending"
        }

        applications.append(application)

        with open("applications.json", "w") as file:

            json.dump(
                applications,
                file,
                indent=4
            )

        return redirect("/applications")

    return render_template(
        "new_admission.html"
    )


# =========================
# APPLICATIONS
# =========================

@app.route("/applications")
def applications():

    # =========================
    # ADMIN LOGIN CHECK
    # =========================

    admin_data = get_admin_data()

    if session.get("admin_id") != admin_data.get("admin_id"):
        session.pop("admin_id", None)
        return redirect(url_for("admin_login"))


    application_data = []

    try:

        with open("applications.json", "r") as file:
            application_data = json.load(file)

    except:

        application_data = []

    return render_template(
        "applications.html",
        applications=application_data
    )


# =========================
# APPROVE APPLICATION
# =========================

@app.route("/approve-application/<application_id>")
def approve_application(application_id):

    # =========================
    # ADMIN LOGIN CHECK
    # =========================

    admin_data = get_admin_data()

    if session.get("admin_id") != admin_data.get("admin_id"):
        session.pop("admin_id", None)
        return redirect(url_for("admin_login"))


    applications = []

    try:

        with open("applications.json", "r") as file:
            applications = json.load(file)

    except:

        applications = []

    for application in applications:

        if application["application_id"] == application_id:

            application["status"] = "Approved"

            break

    with open("applications.json", "w") as file:

        json.dump(
            applications,
            file,
            indent=4
        )

    return redirect("/applications")


# =========================
# REJECT APPLICATION
# =========================

@app.route("/reject-application/<application_id>")
def reject_application(application_id):

    # =========================
    # ADMIN LOGIN CHECK
    # =========================

    admin_data = get_admin_data()

    if session.get("admin_id") != admin_data.get("admin_id"):
        session.pop("admin_id", None)
        return redirect(url_for("admin_login"))


    applications = []

    try:

        with open("applications.json", "r") as file:
            applications = json.load(file)

    except:

        applications = []

    for application in applications:

        if application["application_id"] == application_id:

            application["status"] = "Rejected"

            break

    with open("applications.json", "w") as file:

        json.dump(
            applications,
            file,
            indent=4
        )

    return redirect("/applications")


# =========================
# ADMISSION STATUS
# =========================

@app.route("/admission-status", methods=["GET", "POST"])
def admission_status():

    application = None
    not_found = False

    if request.method == "POST":

        application_id = request.form.get(
            "application_id",
            ""
        ).strip().upper()

        try:

            with open("applications.json", "r") as file:
                applications_data = json.load(file)

        except:

            applications_data = []

        for item in applications_data:

            if (
                item.get(
                    "application_id",
                    ""
                ).upper()
                == application_id
            ):

                application = item

                break

        if application is None:

            not_found = True

    return render_template(
        "admission_status.html",
        application=application,
        not_found=not_found
    )


# =========================
# ADMISSION DOCUMENTS
# =========================

@app.route(
    "/admission-documents",
    methods=["GET", "POST"]
)
def admission_documents():

    if request.method == "POST":

        application_id = request.form.get(
            "application_id",
            ""
        ).strip().upper()

        files = [

            (
                "tenth_marksheet",
                request.files.get(
                    "tenth_marksheet"
                )
            ),

            (
                "twelfth_marksheet",
                request.files.get(
                    "twelfth_marksheet"
                )
            ),

            (
                "identity_document",
                request.files.get(
                    "identity_document"
                )
            )
        ]

        for document_name, document_file in files:

            if (
                document_file
                and document_file.filename != ""
            ):

                filename = (
                    application_id
                    + "_"
                    + document_name
                    + "_"
                    + secure_filename(
                        document_file.filename
                    )
                )

                save_uploaded_file(document_file, filename)

        return "Documents Uploaded Successfully!"

    return render_template(
        "admission_documents.html"
    )


# =========================
# VIEW DOCUMENTS
# =========================

@app.route("/view-documents")
def view_documents():

    # =========================
    # ADMIN LOGIN CHECK
    # =========================

    admin_data = get_admin_data()

    if session.get("admin_id") != admin_data.get("admin_id"):
        session.pop("admin_id", None)
        return redirect(url_for("admin_login"))


    documents = []

    if os.path.exists("uploads"):

        for filename in os.listdir("uploads"):

            parts = filename.split("_", 2)

            if len(parts) == 3:

                application_id = parts[0]

                document_type = parts[1]

                documents.append({

                    "application_id":
                        application_id,

                    "document_type":
                        document_type,

                    "filename":
                        filename
                })

    return render_template(
        "view_documents.html",
        documents=documents
    )


# =========================
# FACULTY LOGIN
# =========================

@app.route(
    "/faculty-login",
    methods=["GET", "POST"]
)
def faculty_login():

    if request.method == "POST":

        faculty_id = request.form.get(
            "faculty_id", ""
        ).strip()

        password = request.form.get(
            "password", ""
        )

        faculty_list = []

        if os.path.exists(FACULTY_FILE):

            try:

                with open(
                    FACULTY_FILE,
                    "r"
                ) as file:

                    faculty_list = json.load(file)

            except:

                faculty_list = []

        faculty = None

        # Faculty ID find karo
        for item in faculty_list:

            if str(
                item.get("faculty_id", "")
            ).strip() == faculty_id:

                faculty = item
                break

        if faculty is None:

            return "Invalid Faculty ID or Password"

        # Faculty ID
        actual_faculty_id = str(
            faculty.get("faculty_id", "")
        ).strip()

        # Password check
        if faculty.get("password_hash"):

            # Changed password
            if not check_password_hash(
                faculty["password_hash"],
                password
            ):

                return "Invalid Faculty ID or Password"

        else:

            # First login: password = Faculty ID
            if password != actual_faculty_id:

                return "Invalid Faculty ID or Password"

            # Password ko hash karke save karo
            faculty["password_hash"] = generate_password_hash(
                actual_faculty_id
            )

            # Purana plaintext password remove karo
            faculty.pop("password", None)

            with open(
                FACULTY_FILE,
                "w"
            ) as file:

                json.dump(
                    faculty_list,
                    file,
                    indent=4
                )

        # =================================================
        # DEPARTMENT LOCK CHECK
        # =================================================

        faculty_department = str(
            faculty.get(
                "department",
                ""
            )
        ).strip()

        conn = get_connection()

        try:

            with conn.cursor() as cur:

                cur.execute("""
                    SELECT is_closed
                    FROM department_admins
                    WHERE LOWER(department_name) = LOWER(%s)
                """, (
                    faculty_department,
                ))

                department_data = cur.fetchone()

        finally:

            conn.close()

        # Department Owner ke system me nahi hai
        # OR Department temporarily closed hai

        if (
            not department_data
            or department_data["is_closed"]
        ):

            return (
                "Your service is locked. "
                "First payment by department then use this service."
            )

        # -------------------------
        # ROLE SESSION ISOLATION
        # -------------------------
        # Faculty login ke time purane Admin/Student session keys hatao.
        session.pop("admin_id", None)
        session.pop("student_id", None)
        session.pop("student_enrollment", None)
        session.pop("student_course", None)
        session.pop("student_semester", None)
        session.pop("student_section", None)

        # Faculty session

        session["faculty_id"] = actual_faculty_id
        session["faculty_name"] = str(
            faculty.get("name", "")
        ).strip()

        return redirect(
            "/faculty-dashboard"
        )

    return render_template(
        "faculty_login.html"
    )

@app.route("/faculty-change-password", methods=["GET", "POST"])
def faculty_change_password():

    # =========================
    # FACULTY LOGIN SECURITY
    # =========================

    faculty_id = str(
        session.get("faculty_id", "")
    ).strip()

    if not faculty_id:
        return redirect(
            url_for("faculty_login")
        )


    # =========================
    # LOAD FACULTY DATA
    # =========================

    faculty_list = []

    if os.path.exists(FACULTY_FILE):

        try:

            with open(
                FACULTY_FILE,
                "r",
                encoding="utf-8"
            ) as file:

                faculty_list = json.load(file)

            if not isinstance(
                faculty_list,
                list
            ):
                faculty_list = []

        except Exception:

            faculty_list = []


    # =========================
    # FIND LOGGED-IN FACULTY
    # =========================

    faculty = None

    for item in faculty_list:

        if not isinstance(
            item,
            dict
        ):
            continue

        stored_faculty_id = str(
            item.get(
                "faculty_id",
                ""
            )
        ).strip()

        if stored_faculty_id == faculty_id:

            faculty = item

            break


    # =========================
    # FACULTY NOT FOUND
    # =========================

    if faculty is None:

        session.pop(
            "faculty_id",
            None
        )

        session.pop(
            "faculty_name",
            None
        )

        return redirect(
            url_for("faculty_login")
        )


    # =========================
    # CHANGE PASSWORD
    # =========================

    if request.method == "POST":

        current_password = request.form.get(
            "current_password",
            ""
        )

        new_password = request.form.get(
            "new_password",
            ""
        )

        confirm_password = request.form.get(
            "confirm_password",
            ""
        )


        # =========================
        # REQUIRED FIELDS
        # =========================

        if (
            not current_password
            or not new_password
            or not confirm_password
        ):

            return "Please fill all fields."


        # =========================
        # CURRENT PASSWORD CHECK
        # =========================

        stored_password_hash = str(
            faculty.get(
                "password_hash",
                ""
            )
        )

        if not stored_password_hash:

            return (
                "Faculty password is not "
                "properly configured.",
                500
            )


        if not check_password_hash(
            stored_password_hash,
            current_password
        ):

            return (
                "Current password is incorrect."
            )


        # =========================
        # CONFIRM NEW PASSWORD
        # =========================

        if (
            new_password
            != confirm_password
        ):

            return (
                "New passwords do not match."
            )


        # =========================
        # SAME PASSWORD CHECK
        # =========================

        if (
            current_password
            == new_password
        ):

            return (
                "New password must be "
                "different from current password."
            )


        # =========================
        # NEW PASSWORD LENGTH
        # =========================

        if len(new_password) < 6:

            return (
                "New password must be at least "
                "6 characters."
            )


        # =========================
        # HASH NEW PASSWORD
        # =========================

        faculty["password_hash"] = (
            generate_password_hash(
                new_password
            )
        )


        # =========================
        # REMOVE OLD PLAINTEXT PASSWORD
        # =========================

        faculty.pop(
            "password",
            None
        )


        # =========================
        # SAVE FACULTY DATA
        # =========================

        with open(
            FACULTY_FILE,
            "w",
            encoding="utf-8"
        ) as file:

            json.dump(
                faculty_list,
                file,
                indent=4,
                ensure_ascii=False
            )


        # =========================
        # RETURN TO DASHBOARD
        # =========================

        return redirect(
            url_for("faculty_dashboard")
        )


    # =========================
    # CHANGE PASSWORD PAGE
    # =========================

    return render_template(
        "faculty_change_password.html"
    )

@app.route("/reset-faculty-password/<faculty_id>")
def reset_faculty_password(faculty_id):

    # =========================
    # ADMIN LOGIN CHECK
    # =========================

    admin_data = get_admin_data()

    if session.get("admin_id") != admin_data.get("admin_id"):
        session.pop("admin_id", None)
        return redirect(url_for("admin_login"))


    faculty_list = []

    if os.path.exists(FACULTY_FILE):

        try:
            with open(FACULTY_FILE, "r") as file:
                faculty_list = json.load(file)

        except:
            faculty_list = []

    faculty_found = False

    for faculty in faculty_list:

        if str(
            faculty.get("faculty_id", "")
        ).strip() == str(faculty_id).strip():

            # Password ko Faculty ID par reset karo
            faculty["password_hash"] = generate_password_hash(
                faculty_id
            )

            # Agar purana plaintext password hai
            # to remove kar do
            faculty.pop("password", None)

            faculty_found = True
            break

    if not faculty_found:
        return "Faculty not found."

    with open(FACULTY_FILE, "w") as file:

        json.dump(
            faculty_list,
            file,
            indent=4
        )

    return redirect("/admin-faculty")

@app.route("/edit-faculty/<faculty_id>", methods=["GET", "POST"])
def edit_faculty(faculty_id):

    # =========================
    # ADMIN LOGIN CHECK
    # =========================

    admin_data = get_admin_data()

    if session.get("admin_id") != admin_data.get("admin_id"):
        session.pop("admin_id", None)
        return redirect(url_for("admin_login"))


    faculty_list = []

    if os.path.exists(FACULTY_FILE):

        try:
            with open(FACULTY_FILE, "r") as file:
                faculty_list = json.load(file)

            if not isinstance(faculty_list, list):
                faculty_list = []

        except:
            faculty_list = []


    # ==============================
    # FIND FACULTY
    # ==============================

    faculty = None

    for item in faculty_list:

        if str(
            item.get("faculty_id", "")
        ).strip().lower() == str(
            faculty_id
        ).strip().lower():

            faculty = item
            break


    if faculty is None:
        return "Faculty not found."


    # ==============================
    # WHERE DID EDIT OPEN FROM?
    # ==============================

    from_page = request.args.get(
        "from",
        ""
    )

    search = request.args.get(
        "q",
        ""
    ).strip()


    # ==============================
    # UPDATE FACULTY
    # ==============================

    if request.method == "POST":

        name = request.form.get(
            "name",
            ""
        ).strip()

        department = request.form.get(
            "department",
            ""
        ).strip()


        if not name or not department:
            return "Please fill all fields."


        # ==============================
        # VALIDATE DEPARTMENT
        # ==============================

        departments = []

        if os.path.exists(DEPARTMENTS_FILE):

            try:
                with open(
                    DEPARTMENTS_FILE,
                    "r"
                ) as file:

                    departments = json.load(file)

                if not isinstance(
                    departments,
                    list
                ):
                    departments = []

            except:
                departments = []


        valid_department = None

        for dept in departments:

            dept_name = str(
                dept.get("name", "")
            ).strip()

            if dept_name.lower() == department.lower():

                valid_department = dept_name
                break


        if not valid_department:
            return "Invalid Department."


        faculty["name"] = name
        faculty["department"] = valid_department


        # ==============================
        # SAVE
        # ==============================

        with open(
            FACULTY_FILE,
            "w"
        ) as file:

            json.dump(
                faculty_list,
                file,
                indent=4
            )


        # ==============================
        # RETURN TO SEARCH
        # ==============================

        if from_page == "search":

            return redirect(
                "/admin-search-faculty?q="
                + search
            )


        # ==============================
        # RETURN TO FACULTY MANAGEMENT
        # ==============================

        return redirect(
            "/admin-faculty"
        )


    # ==============================
    # EDIT PAGE
    # ==============================

    return render_template(
        "edit_faculty.html",
        faculty=faculty,
        from_page=from_page,
        search=search
    )


# =========================
# FACULTY DASHBOARD
# =========================

@app.route("/faculty-dashboard")
def faculty_dashboard():

    # Faculty login check
    if "faculty_id" not in session:
        return redirect(url_for("faculty_login"))

    faculty_id = str(
        session.get("faculty_id", "")
    ).strip()

    # Faculty data load
    faculty_list = []

    if os.path.exists(FACULTY_FILE):
        try:
            with open(FACULTY_FILE, "r") as file:
                faculty_list = json.load(file)
        except:
            faculty_list = []

    # Logged-in faculty find
    faculty = None

    for item in faculty_list:
        if str(
            item.get("faculty_id", "")
        ).strip() == faculty_id:
            faculty = item
            break

    if faculty is None:
        return "Faculty not found.", 404

    return render_template(
        "faculty_dashboard.html",
        faculty=faculty
    )


# =========================
# FACULTY PROFILE
# =========================

@app.route(
    "/faculty-profile",
    methods=["GET", "POST"]
)
def faculty_profile():

    # =========================
    # FACULTY LOGIN SECURITY
    # =========================

    faculty_id = str(
        session.get("faculty_id", "")
    ).strip()

    if not faculty_id:

        return redirect(
            url_for("faculty_login")
        )


    # =========================
    # PROFILE FILE
    # =========================

    profile_file = "faculty_profile.json"


    # =========================
    # DEFAULT PROFILE
    # =========================

    profile = {

        "faculty_id":
            faculty_id,

        "name":
            "Faculty Member",

        "department":
            "Computer Science & Engineering",

        "email":
            "faculty@university.com",

        "designation":
            "Assistant Professor"
    }


    # =========================
    # LOAD PROFILE
    # =========================

    if os.path.exists(profile_file):

        try:

            with open(
                profile_file,
                "r",
                encoding="utf-8"
            ) as file:

                loaded_profile = json.load(file)

            if isinstance(
                loaded_profile,
                dict
            ):

                profile = loaded_profile

        except Exception:

            pass


    # =========================
    # PROFILE OWNERSHIP CHECK
    # =========================

    stored_profile_id = str(
        profile.get(
            "faculty_id",
            ""
        )
    ).strip()

    if (
        stored_profile_id
        and stored_profile_id != faculty_id
    ):

        return (
            "Faculty profile not found.",
            404
        )


    # Always keep the logged-in
    # Faculty ID.

    profile["faculty_id"] = faculty_id


    # =========================
    # UPDATE PROFILE
    # =========================

    if request.method == "POST":

        name = request.form.get(
            "name",
            ""
        ).strip()

        department = request.form.get(
            "department",
            ""
        ).strip()

        email = request.form.get(
            "email",
            ""
        ).strip()

        designation = request.form.get(
            "designation",
            ""
        ).strip()


        profile["faculty_id"] = faculty_id

        profile["name"] = name

        profile["department"] = department

        profile["email"] = email

        profile["designation"] = designation


        # =========================
        # SAVE
        # =========================

        with open(
            profile_file,
            "w",
            encoding="utf-8"
        ) as file:

            json.dump(
                profile,
                file,
                indent=4,
                ensure_ascii=False
            )


        return redirect(
            url_for("faculty_profile")
        )


    # =========================
    # SHOW PROFILE
    # =========================

    return render_template(
        "faculty_profile.html",
        profile=profile
    )


# =========================
# ATTENDANCE HISTORY
# =========================

@app.route("/attendance-history")
def attendance_history():

    # =========================
    # STUDENT LOGIN CHECK
    # =========================

    if "student_enrollment" not in session:
        return redirect(url_for("student_login"))

    student_enrollment = str(
        session.get("student_enrollment", "")
    ).strip()

    if not student_enrollment:
        return redirect(url_for("student_login"))


    attendance_records = []

    # Attendance data पढ़ना
    if os.path.exists("attendance.json"):

        try:
            with open("attendance.json", "r") as file:
                all_attendance = json.load(file)

        except:
            all_attendance = []

    else:
        all_attendance = []


    # =========================
    # LOGGED-IN STUDENT
    # =========================

    students = []

    if os.path.exists(STUDENTS_FILE):

        try:
            with open(STUDENTS_FILE, "r") as file:
                students = json.load(file)

        except:
            students = []


    # Current student ढूँढना
    student = None

    for item in students:

        if str(
            item.get("enrollment", "")
        ).strip() == student_enrollment:

            student = item
            break


    if not student:
        return "Student not found.", 404


    # Student की attendance
    student_course = str(
        student.get("course", "")
    ).strip()

    student_semester = str(
        student.get("semester", "")
    ).strip()

    student_section = str(
        student.get("section", "")
    ).strip()

    student_enrollment = str(
        student.get("enrollment", "")
    ).strip()


    # सिर्फ current student की attendance
    for record in all_attendance:

        if (
            str(record.get("course", "")).strip()
            == student_course

            and

            str(record.get("semester", "")).strip()
            == student_semester

            and

            str(record.get("section", "")).strip()
            == student_section

            and

            str(record.get("student_id", "")).strip()
            == student_enrollment
        ):

            attendance_records.append(record)


    # Attendance percentage
    total_classes = len(attendance_records)

    present_classes = 0

    for record in attendance_records:

        if record.get("status") == "Present":
            present_classes += 1


    if total_classes > 0:

        attendance_percentage = round(
            (present_classes / total_classes) * 100,
            2
        )

    else:

        attendance_percentage = 0


    return render_template(
        "attendance_history.html",
        attendance_records=attendance_records,
        total_classes=total_classes,
        present_classes=present_classes,
        attendance_percentage=attendance_percentage
    )


# =========================
# UPLOAD NOTES - FACULTY + COURSE + SECTION
# =========================

@app.route("/upload-notes", methods=["GET", "POST"])
def upload_notes():

    # Faculty login check
    if "faculty_id" not in session:
        return redirect(url_for("faculty_login"))

    faculty_id = str(session.get("faculty_id", "")).strip()
    faculty_name = str(session.get("faculty_name", "")).strip()

    # Faculty name session me nahi hai to faculty file se lo
    if not faculty_name:

        faculties = []

        if os.path.exists(FACULTY_FILE):
            try:
                with open(FACULTY_FILE, "r") as file:
                    faculties = json.load(file)
            except:
                faculties = []

        for faculty in faculties:

            if str(faculty.get("faculty_id", "")).strip() == faculty_id:
                faculty_name = str(
                    faculty.get("name", "")
                ).strip()
                break

    # Admin courses
    courses = get_admin_courses()

    # Admin subjects
    subjects = []

    if os.path.exists(SUBJECTS_FILE):
        try:
            with open(SUBJECTS_FILE, "r") as file:
                subjects = json.load(file)
        except:
            subjects = []

    if request.method == "POST":

        course = str(
            request.form.get("course", "")
        ).strip()

        semester = str(
            request.form.get("semester", "")
        ).strip()

        section = str(
            request.form.get("section", "")
        ).strip().upper()

        subject = str(
            request.form.get("subject", "")
        ).strip()

        title = str(
            request.form.get("title", "")
        ).strip()

        file = request.files.get("file")

        # Validation
        valid_courses = []

        for c in courses:

            if isinstance(c, dict):
                name = c.get("course_name") or c.get("name")
            else:
                name = c

            if name:
                valid_courses.append(str(name).strip())

        valid_course = None

        for c in valid_courses:

            if c.lower() == course.lower():
                valid_course = c
                break

        if not valid_course:
            return "Invalid course"

        if semester not in [str(i) for i in range(1, 11)]:
            return "Invalid semester"

        if section not in ["A", "B", "C", "D"]:
            return "Invalid section"

        # Subject must belong to selected course + semester
        valid_subject = None

        for s in subjects:

            s_course = str(
                s.get("course", "")
            ).strip()

            s_semester = str(
                s.get("semester", "")
            ).strip()

            s_name = str(
                s.get("subject_name", "")
            ).strip()

            if (
                s_course.lower() == valid_course.lower()
                and s_semester == semester
                and s_name.lower() == subject.lower()
            ):
                valid_subject = s_name
                break

        if not valid_subject:
            return "Invalid subject"

        if not file or not file.filename:
            return "Please select a file"

        # Allowed file types
        if not allowed_file(file.filename):
            return "Only PDF, JPG, JPEG and PNG files are allowed"

        # Secure filename
        filename = secure_filename(file.filename)

        # Unique filename
        import uuid

        filename = str(uuid.uuid4()) + "_" + filename

        os.makedirs(UPLOAD_FOLDER, exist_ok=True)

        save_uploaded_file(file, filename)

        notes = []

        if os.path.exists(NOTES_FILE):
            try:
                with open(NOTES_FILE, "r") as f:
                    notes = json.load(f)
            except:
                notes = []

        # Save note
        notes.append({
            "faculty_id": faculty_id,
            "faculty_name": faculty_name,
            "course": valid_course,
            "semester": semester,
            "section": section,
            "subject": valid_subject,
            "title": title,
            "filename": filename
        })

        with open(NOTES_FILE, "w") as f:
            json.dump(notes, f, indent=4)

        return redirect(url_for("upload_notes"))

    # Only this faculty's notes
    my_notes = []

    if os.path.exists(NOTES_FILE):

        try:
            with open(NOTES_FILE, "r") as f:
                notes = json.load(f)
        except:
            notes = []

        for index, note in enumerate(notes):

            if str(
                note.get("faculty_id", "")
            ).strip() == faculty_id:

                note["_index"] = index
                my_notes.append(note)

    return render_template(
        "upload_notes.html",
        faculty_name=faculty_name,
        courses=courses,
        subjects=subjects,
        notes=my_notes
    )


@app.route("/delete-note/<int:note_index>", methods=["POST", "GET"])
def delete_note(note_index):

    # Faculty login check
    if "faculty_id" not in session:
        return redirect(url_for("faculty_login"))

    faculty_id = str(
        session.get("faculty_id", "")
    ).strip()

    notes = []

    if os.path.exists(NOTES_FILE):

        try:
            with open(NOTES_FILE, "r") as f:
                notes = json.load(f)
        except:
            notes = []

    if note_index < 0 or note_index >= len(notes):
        return redirect(url_for("upload_notes"))

    note = notes[note_index]

    # केवल अपना note delete कर सकता है
    if str(
        note.get("faculty_id", "")
    ).strip() != faculty_id:
        return "Unauthorized", 403

    filename = note.get("filename", "")

    # Delete physical file
    if filename:
        delete_uploaded_file(filename)

    notes.pop(note_index)

    with open(NOTES_FILE, "w") as f:
        json.dump(notes, f, indent=4)

    return redirect(url_for("upload_notes"))

# =========================
# FACULTY SYLLABUS UPLOAD
# =========================

@app.route("/faculty-syllabus", methods=["GET", "POST"])
def faculty_syllabus():

    if "faculty_id" not in session:
        return redirect(url_for("faculty_login"))

    faculty_id = str(session.get("faculty_id", "")).strip()
    faculty_name = str(session.get("faculty_name", "")).strip()

    # Faculty name session में न हो तो file से लो
    if not faculty_name:

        faculties = []

        if os.path.exists(FACULTY_FILE):
            try:
                with open(FACULTY_FILE, "r") as file:
                    faculties = json.load(file)
            except:
                faculties = []

        for faculty in faculties:
            if str(
                faculty.get("faculty_id", "")
            ).strip() == faculty_id:

                faculty_name = str(
                    faculty.get("name", "")
                ).strip()

                break

    # Admin Course Management से courses
    courses = get_admin_courses()

    # Admin Subject Management से subjects
    subjects = []

    if os.path.exists(SUBJECTS_FILE):
        try:
            with open(SUBJECTS_FILE, "r") as file:
                subjects = json.load(file)
        except:
            subjects = []

    if request.method == "POST":

        course = str(
            request.form.get("course", "")
        ).strip()

        semester = str(
            request.form.get("semester", "")
        ).strip()

        section = str(
            request.form.get("section", "")
        ).strip().upper()

        subject = str(
            request.form.get("subject", "")
        ).strip()

        file = request.files.get("syllabus_file")

        # -----------------------------
        # COURSE VALIDATION
        # -----------------------------

        valid_courses = []

        for c in courses:

            if isinstance(c, dict):
                course_name = (
                    c.get("course_name")
                    or c.get("name")
                )
            else:
                course_name = c

            if course_name:
                valid_courses.append(
                    str(course_name).strip()
                )

        valid_course = None

        for c in valid_courses:

            if c.lower() == course.lower():
                valid_course = c
                break

        if not valid_course:
            return "Invalid course"

        # -----------------------------
        # SEMESTER VALIDATION
        # -----------------------------

        if semester not in [
            str(i) for i in range(1, 11)
        ]:
            return "Invalid semester"

        # -----------------------------
        # SECTION VALIDATION
        # -----------------------------

        if section not in ["A", "B", "C", "D"]:
            return "Invalid section"

        # -----------------------------
        # SUBJECT VALIDATION
        # Course + Semester के according
        # -----------------------------

        valid_subject = None

        for s in subjects:

            s_course = str(
                s.get("course", "")
            ).strip()

            s_semester = str(
                s.get("semester", "")
            ).strip()

            s_name = str(
                s.get("subject_name", "")
            ).strip()

            if (
                s_course.lower() == valid_course.lower()
                and s_semester == semester
                and s_name.lower() == subject.lower()
            ):
                valid_subject = s_name
                break

        if not valid_subject:
            return "Invalid subject"

        # -----------------------------
        # FILE VALIDATION
        # -----------------------------

        if not file or not file.filename:
            return "Please select a syllabus file"

        if not allowed_file(file.filename):
            return "Only PDF, JPG, JPEG and PNG files are allowed"

        filename = secure_filename(file.filename)

        import uuid

        filename = str(uuid.uuid4()) + "_" + filename

        os.makedirs(
            SYLLABUS_FOLDER,
            exist_ok=True
        )

        save_uploaded_file(file, filename)

        # -----------------------------
        # OLD SYLLABUS DATA
        # -----------------------------

        syllabuses = []

        if os.path.exists(SYLLABUS_FILE):

            try:
                with open(
                    SYLLABUS_FILE,
                    "r"
                ) as f:

                    syllabuses = json.load(f)

            except:
                syllabuses = []

        # -----------------------------
        # SAVE SYLLABUS
        # -----------------------------

        syllabuses.append({

            "faculty_id": faculty_id,

            "faculty_name": faculty_name,

            "course": valid_course,

            "semester": semester,

            "section": section,

            "subject": valid_subject,

            "filename": filename

        })

        with open(
            SYLLABUS_FILE,
            "w"
        ) as f:

            json.dump(
                syllabuses,
                f,
                indent=4
            )

        return redirect(
            url_for("faculty_syllabus")
        )

    # -----------------------------
    # ONLY OWN SYLLABUS
    # -----------------------------

    my_syllabuses = []

    if os.path.exists(SYLLABUS_FILE):

        try:
            with open(
                SYLLABUS_FILE,
                "r"
            ) as f:

                all_syllabuses = json.load(f)

        except:
            all_syllabuses = []

        for index, syllabus in enumerate(
            all_syllabuses
        ):

            if str(
                syllabus.get("faculty_id", "")
            ).strip() == faculty_id:

                syllabus["_index"] = index

                my_syllabuses.append(
                    syllabus
                )

    return render_template(
        "faculty_syllabus.html",
        faculty_name=faculty_name,
        courses=courses,
        subjects=subjects,
        syllabuses=my_syllabuses
    )


@app.route(
    "/delete-syllabus/<int:syllabus_index>",
    methods=["POST", "GET"]
)
def delete_syllabus(syllabus_index):

    if "faculty_id" not in session:
        return redirect(url_for("faculty_login"))

    faculty_id = str(
        session.get("faculty_id", "")
    ).strip()

    syllabuses = []

    if os.path.exists(SYLLABUS_FILE):

        try:
            with open(
                SYLLABUS_FILE,
                "r"
            ) as f:

                syllabuses = json.load(f)

        except:
            syllabuses = []

    if (
        syllabus_index < 0
        or syllabus_index >= len(syllabuses)
    ):
        return redirect(
            url_for("faculty_syllabus")
        )

    syllabus = syllabuses[syllabus_index]

    # केवल अपना syllabus delete कर सकता है
    if str(
        syllabus.get("faculty_id", "")
    ).strip() != faculty_id:

        return "Unauthorized", 403

    filename = syllabus.get(
        "filename",
        ""
    )

    if filename:
        delete_uploaded_file(filename)

    syllabuses.pop(syllabus_index)

    with open(
        SYLLABUS_FILE,
        "w"
    ) as f:

        json.dump(
            syllabuses,
            f,
            indent=4
        )

    return redirect(
        url_for("faculty_syllabus")
    )

# =========================
# FACULTY SEND NOTICE
# =========================

NOTICES_FILE = "notices.json"


# =========================
# FACULTY SEND NOTICE
# =========================

NOTICES_FILE = "notices.json"


@app.route("/send-notice", methods=["GET", "POST"])
def send_notice():

    # -------------------------
    # Faculty login check
    # -------------------------
    faculty_id = session.get("faculty_id")

    if not faculty_id:
        return redirect("/faculty-login")

    faculty_name = str(
        session.get("faculty_name", "")
    ).strip()

    # -------------------------
    # Faculty name fallback
    # -------------------------
    if not faculty_name:

        faculty_list = []

        if os.path.exists(FACULTY_FILE):
            try:
                with open(FACULTY_FILE, "r") as file:
                    faculty_list = json.load(file)
            except:
                faculty_list = []

        for faculty in faculty_list:

            if str(
                faculty.get("faculty_id", "")
            ).strip() == str(faculty_id).strip():

                faculty_name = str(
                    faculty.get("name", "")
                ).strip()

                session["faculty_name"] = faculty_name
                break

    # -------------------------
    # Admin Course Management
    # -------------------------
    courses = get_admin_courses()

    clean_courses = []

    for item in courses:

        if isinstance(item, dict):

            course_name = str(
                item.get("name")
                or item.get("course_name")
                or item.get("course")
                or ""
            ).strip()

        else:

            course_name = str(item).strip()

        if course_name:
            clean_courses.append(course_name)

    # -------------------------
    # SEND NEW NOTICE
    # -------------------------
    if request.method == "POST":

        course = request.form.get(
            "course", ""
        ).strip()

        semester = request.form.get(
            "semester", ""
        ).strip()

        section = request.form.get(
            "section", ""
        ).strip().upper()

        title = request.form.get(
            "title", ""
        ).strip()

        message = request.form.get(
            "message", ""
        ).strip()

        # -------------------------
        # Optional PDF
        # -------------------------
        pdf_file = request.files.get("pdf_file")

        # -------------------------
        # Required fields
        # -------------------------
        if (
            not course
            or not semester
            or not section
            or not title
            or not message
        ):
            return "Please fill all fields."

        # -------------------------
        # Semester validation
        # -------------------------
        if semester not in [
            "1", "2", "3", "4",
            "5", "6", "7", "8", "9", "10"
        ]:
            return "Invalid Semester."

        # -------------------------
        # Section validation
        # -------------------------
        if section not in [
            "A", "B", "C", "D"
        ]:
            return "Invalid Section."

        # -------------------------
        # Course validation
        # -------------------------
        valid_course = None

        for item in clean_courses:

            if item.lower() == course.lower():

                valid_course = item
                break

        if not valid_course:
            return "Invalid Course."

        # -------------------------
        # Optional PDF Upload
        # -------------------------
        pdf_filename = ""

        if pdf_file and pdf_file.filename:

            if not pdf_file.filename.lower().endswith(".pdf"):
                return "Only PDF files are allowed."

            import uuid

            original_name = secure_filename(
                pdf_file.filename
            )

            if not original_name:
                return "Invalid PDF filename."

            pdf_filename = (
                str(uuid.uuid4())
                + "_"
                + original_name
            )

            save_uploaded_file(
                pdf_file,
                pdf_filename
            )

        # -------------------------
        # Existing notices
        # -------------------------
        notices = []

        if os.path.exists(NOTICES_FILE):

            try:
                with open(
                    NOTICES_FILE,
                    "r"
                ) as file:

                    notices = json.load(file)

                if not isinstance(
                    notices,
                    list
                ):
                    notices = []

            except:
                notices = []

        # -------------------------
        # New Notice
        # -------------------------
        notice = {

            "faculty_id":
                str(faculty_id).strip(),

            "faculty_name":
                faculty_name,

            "course":
                valid_course,

            "semester":
                semester,

            "section":
                section,

            "title":
                title,

            "message":
                message,

            "pdf_filename":
                pdf_filename
        }

        notices.append(notice)

        # -------------------------
        # Save Notice
        # -------------------------
        with open(
            NOTICES_FILE,
            "w"
        ) as file:

            json.dump(
                notices,
                file,
                indent=4
            )

        return redirect("/send-notice")

    # =========================
    # SHOW FACULTY'S OWN NOTICES
    # =========================

    all_notices = []

    if os.path.exists(NOTICES_FILE):

        try:
            with open(
                NOTICES_FILE,
                "r"
            ) as file:

                all_notices = json.load(file)

            if not isinstance(
                all_notices,
                list
            ):
                all_notices = []

        except:
            all_notices = []

    notices = []

    for index, notice in enumerate(all_notices):

        if str(
            notice.get("faculty_id", "")
        ).strip() == str(faculty_id).strip():

            notice_copy = dict(notice)

            # Original JSON index
            # delete ke liye
            notice_copy["_index"] = index

            notices.append(notice_copy)

    return render_template(
        "send_notice.html",
        notices=notices,
        courses=clean_courses,
        faculty_name=faculty_name
    )

@app.route("/notice-file/<filename>")
def notice_file(filename):

    # =========================
    # LOGIN CHECK
    # Student ya Faculty login
    # hona zaroori hai
    # =========================

    if (
        "student_id" not in session
        and "faculty_id" not in session
    ):
        return redirect(url_for("student_login"))


    stored = get_uploaded_file(filename)

    if not stored:
        return "File not found", 404

    return send_file(
        io.BytesIO(bytes(stored["content"])),
        mimetype="application/pdf",
        as_attachment=False,
        download_name=filename
    )


@app.route("/head-delete-faculty-records", methods=["GET"])
def head_delete_faculty_records():

    # ==============================
    # HEAD LOGIN CHECK
    # ==============================

    if not session.get("head_logged_in"):
        return redirect(url_for("head_login"))

    # ==============================
    # CONFIRMATION PAGE
    # ==============================

    return render_template(
        "head_faculty_records.html"
    )


@app.route("/head-confirm-delete-faculty-data", methods=["POST"])
def head_confirm_delete_faculty_data():

    # ==============================
    # HEAD LOGIN CHECK
    # ==============================

    if not session.get("head_logged_in"):
        return redirect(url_for("head_login"))

    # ==============================
    # DELETE NOTES DATA
    # ==============================

    try:
        json_delete(NOTES_FILE)
    except Exception:
        pass


    # ==============================
    # DELETE SYLLABUS DATA
    # ==============================

    try:
        json_delete(SYLLABUS_FILE)
    except Exception:
        pass


    # ==============================
    # DELETE NOTICE DATA
    # ==============================

    try:
        json_delete(NOTICES_FILE)
    except Exception:
        pass


    # ==============================
    # DELETE HOLIDAY DATA
    # ==============================

    try:
        json_delete(HOLIDAYS_FILE)
    except Exception:
        pass


    # ==============================
    # DELETE ALL UPLOADED FILES
    # FROM POSTGRESQL DATABASE
    # ==============================

    try:

        conn = get_connection()

        try:

            with conn.cursor() as cur:

                cur.execute(
                    "DELETE FROM uploaded_files"
                )

            conn.commit()

        finally:

            conn.close()

    except Exception:
        pass


    # ==============================
    # SUCCESS
    # ==============================

    return """
    <script>

        alert(
            "All faculty uploaded records have been deleted successfully."
        );

        window.location.href = "/head";

    </script>
    """

# =========================
# STUDENT NOTICE BOARD
# =========================

@app.route("/notices")
def student_notices():

    # Student login check
    if "student_enrollment" not in session:
        return redirect(url_for("student_login"))

    student_enrollment = str(
        session.get("student_enrollment", "")
    ).strip()

    # Students load
    students = []

    if os.path.exists(STUDENTS_FILE):
        try:
            with open(STUDENTS_FILE, "r") as file:
                students = json.load(file)
        except:
            students = []

    # Current student find
    student = None

    for item in students:
        if str(item.get("enrollment", "")).strip() == student_enrollment:
            student = item
            break

    if not student:
        return "Student not found", 404

    # Student Course + Semester + Section
    student_course = str(
        student.get("course", "")
    ).strip().lower()

    student_semester = str(
        student.get("semester", "")
    ).strip()

    student_section = str(
        student.get("section", "")
    ).strip().upper()

    # Notices load
    all_notices = []

    if os.path.exists(NOTICES_FILE):
        try:
            with open(NOTICES_FILE, "r") as file:
                all_notices = json.load(file)
        except:
            all_notices = []

    # Matching notices only
    filtered_notices = []

    for notice in all_notices:

        notice_course = str(
            notice.get("course", "")
        ).strip().lower()

        notice_semester = str(
            notice.get("semester", "")
        ).strip()

        notice_section = str(
            notice.get("section", "")
        ).strip().upper()

        if (
            notice_course == student_course
            and notice_semester == student_semester
            and notice_section == student_section
        ):
            filtered_notices.append(notice)

    return render_template(
        "notices.html",
        notices=filtered_notices,
        student=student
    )


# =========================
# DELETE NOTICE
# =========================

@app.route("/delete-notice/<int:notice_index>")
def delete_notice(notice_index):

    # -------------------------
    # Faculty login check
    # -------------------------
    faculty_id = session.get("faculty_id")

    if not faculty_id:
        return redirect("/faculty-login")

    # -------------------------
    # Notices load
    # -------------------------
    notices = []

    if os.path.exists(NOTICES_FILE):

        try:
            with open(
                NOTICES_FILE,
                "r"
            ) as file:

                notices = json.load(file)

        except:
            notices = []

    # -------------------------
    # Index validation
    # -------------------------
    if not (
        0 <= notice_index < len(notices)
    ):
        return redirect("/send-notice")

    notice = notices[notice_index]

    # -------------------------
    # OWNERSHIP CHECK
    # Faculty sirf apna
    # notice delete kar sakta hai
    # -------------------------
    notice_faculty_id = str(
        notice.get("faculty_id", "")
    ).strip()

    if notice_faculty_id != str(
        faculty_id
    ).strip():

        return redirect("/send-notice")

    # -------------------------
    # Delete
    # -------------------------
    notices.pop(notice_index)

    with open(
        NOTICES_FILE,
        "w"
    ) as file:

        json.dump(
            notices,
            file,
            indent=4
        )

    return redirect("/send-notice")

# =========================
# FACULTY HOLIDAY INFORMATION
# =========================

HOLIDAYS_FILE = "holidays.json"


@app.route("/holiday-information", methods=["GET", "POST"])
def holiday_information():

    # =========================
    # FACULTY LOGIN CHECK
    # =========================

    faculty_id = session.get("faculty_id", "").strip()
    faculty_name = session.get("faculty_name", "").strip()

    if not faculty_id:
        return redirect("/faculty-login")


    # =========================
    # FACULTY NAME FALLBACK
    # =========================

    if not faculty_name and os.path.exists(FACULTY_FILE):

        try:
            with open(FACULTY_FILE, "r") as file:
                faculties = json.load(file)

            for faculty in faculties:

                if str(
                    faculty.get("faculty_id", "")
                ).strip() == faculty_id:

                    faculty_name = str(
                        faculty.get("name", "")
                    ).strip()

                    break

        except:
            pass


    # =========================
    # ADMIN COURSES READ
    # =========================

    courses = get_admin_courses()

    course_names = []

    for course in courses:

        if isinstance(course, dict):

            course_name = str(
                course.get("course", "")
                or course.get("name", "")
            ).strip()

        else:

            course_name = str(course).strip()

        if course_name:
            course_names.append(course_name)


    # =========================
    # HOLIDAY DATA READ
    # =========================

    holidays_data = []

    if os.path.exists(HOLIDAYS_FILE):

        try:

            with open(HOLIDAYS_FILE, "r") as file:
                holidays_data = json.load(file)

        except:

            holidays_data = []


    # =========================
    # ADD HOLIDAY
    # =========================

    if request.method == "POST":

        course = request.form.get(
            "course", ""
        ).strip()

        semester = request.form.get(
            "semester", ""
        ).strip()

        section = request.form.get(
            "section", ""
        ).strip()

        holiday_date = request.form.get(
            "holiday_date", ""
        ).strip()

        holiday_name = request.form.get(
            "holiday_name", ""
        ).strip()

        description = request.form.get(
            "description", ""
        ).strip()


        # =========================
        # REQUIRED FIELDS
        # =========================

        if (
            not course
            or not semester
            or not section
            or not holiday_date
            or not holiday_name
        ):

            return "Please fill all required fields."


        # =========================
        # COURSE VALIDATION
        # ADMIN COURSE MANAGEMENT
        # =========================

        valid_course = None

        for admin_course in course_names:

            if admin_course.lower() == course.lower():

                valid_course = admin_course

                break


        if not valid_course:

            return "Invalid Course."


        # =========================
        # SEMESTER VALIDATION
        # =========================

        if semester not in [
            "1", "2", "3", "4",
            "5", "6", "7", "8"
        ]:

            return "Invalid Semester."


        # =========================
        # SECTION VALIDATION
        # =========================

        if section not in [
            "A", "B", "C", "D"
        ]:

            return "Invalid Section."


        # =========================
        # SAVE HOLIDAY
        # =========================

        holiday = {

            "faculty_id": faculty_id,

            "faculty_name": faculty_name,

            "course": valid_course,

            "semester": semester,

            "section": section,

            "date": holiday_date,

            "name": holiday_name,

            "description": description
        }


        holidays_data.append(holiday)


        # =========================
        # SAVE JSON
        # =========================

        with open(
            HOLIDAYS_FILE,
            "w"
        ) as file:

            json.dump(
                holidays_data,
                file,
                indent=4
            )


        return redirect(
            "/holiday-information"
        )


    # =========================
    # SHOW ONLY CURRENT FACULTY
    # =========================

    my_holidays = []

    for index, holiday in enumerate(
        holidays_data
    ):

        if not isinstance(holiday, dict):
            continue

        if str(
            holiday.get("faculty_id", "")
        ).strip() == faculty_id:

            holiday_copy = dict(holiday)

            # Original JSON index
            holiday_copy["_index"] = index

            my_holidays.append(
                holiday_copy
            )


    # =========================
    # SHOW PAGE
    # =========================

    return render_template(
        "holiday_information.html",

        holidays=my_holidays,

        courses=course_names,

        faculty_name=faculty_name
    )




# =========================
# STUDENT HOLIDAY INFORMATION
# =========================

@app.route("/holidays")
def student_holidays():

    # Student login check
    if "student_enrollment" not in session:
        return redirect(url_for("student_login"))

    student_enrollment = str(
        session.get("student_enrollment", "")
    ).strip()

    # =========================
    # STUDENTS LOAD
    # =========================

    students = []

    if os.path.exists(STUDENTS_FILE):
        try:
            with open(STUDENTS_FILE, "r") as file:
                students = json.load(file)
        except:
            students = []

    # =========================
    # CURRENT STUDENT FIND
    # =========================

    student = None

    for item in students:
        if (
            str(item.get("enrollment", "")).strip()
            == student_enrollment
        ):
            student = item
            break

    if not student:
        return "Student not found", 404

    # =========================
    # STUDENT DETAILS
    # =========================

    student_course = str(
        student.get("course", "")
    ).strip().lower()

    student_semester = str(
        student.get("semester", "")
    ).strip()

    student_section = str(
        student.get("section", "")
    ).strip().upper()

    # =========================
    # HOLIDAYS LOAD
    # =========================

    all_holidays = []

    if os.path.exists(HOLIDAYS_FILE):
        try:
            with open(HOLIDAYS_FILE, "r") as file:
                all_holidays = json.load(file)
        except:
            all_holidays = []

    # =========================
    # FILTER HOLIDAYS
    # Course + Semester + Section
    # =========================

    holidays = []

    for holiday in all_holidays:

        holiday_course = str(
            holiday.get("course", "")
        ).strip().lower()

        holiday_semester = str(
            holiday.get("semester", "")
        ).strip()

        holiday_section = str(
            holiday.get("section", "")
        ).strip().upper()

        if (
            holiday_course == student_course
            and holiday_semester == student_semester
            and holiday_section == student_section
        ):
            holidays.append(holiday)

    # =========================
    # SHOW HOLIDAYS
    # =========================

    return render_template(
        "holidays.html",
        holidays=holidays,
        student=student
    )

# =========================
# STUDENT INFORMATION
# =========================

STUDENTS_FILE = "students.json"


@app.route("/student-information")
def student_information():

    students = []

    # students.json se data read karo
    if os.path.exists(STUDENTS_FILE):

        try:
            with open(STUDENTS_FILE, "r") as file:
                students = json.load(file)

        except:
            students = []


    # Enrollment number search
    search_enrollment = request.args.get(
        "enrollment", ""
    ).strip()


    # Student initially None rahega
    student = None


    # Check ki search ki gayi hai ya nahi
    searched = False


    if search_enrollment:

        searched = True

        for item in students:

            if (
                str(item.get("enrollment", "")).strip()
                == search_enrollment
            ):

                student = item
                break


    return render_template(
        "student_information.html",

        student=student,

        search_enrollment=search_enrollment,

        searched=searched
    )

# =========================
# ADMIN LOGIN
# =========================

@app.route("/admin-login", methods=["GET", "POST"])
def admin_login():

    # Departments load karo
    departments = []

    if os.path.exists(DEPARTMENTS_FILE):

        try:
            with open(DEPARTMENTS_FILE, "r") as file:
                departments = json.load(file)

        except:
            departments = []

    if request.method == "POST":

        admin_id = request.form.get(
            "admin_id", ""
        ).strip()

        department = request.form.get(
            "department", ""
        ).strip()

        password = request.form.get(
            "password", ""
        ).strip()

        if not admin_id or not department or not password:
            return "Please enter Admin ID, Department and Password."

        conn = get_connection()

        try:

            with conn.cursor() as cur:

                cur.execute("""
                    SELECT
                        admin_id,
                        department_id,
                        department_name,
                        password,
                        is_closed
                    FROM department_admins
                    WHERE admin_id = %s
                      AND department_name = %s
                """, (
                    admin_id,
                    department
                ))

                admin_data = cur.fetchone()

        finally:
            conn.close()

        # Admin account check
        if not admin_data:
            return "Invalid Admin ID, Password or Department"

        # Temporarily closed check
        if admin_data["is_closed"]:
            return "This Department is temporarily closed."

        # Password check
        if not check_password_hash(
            admin_data["password"],
            password
        ):
            return "Invalid Admin ID, Password or Department"

        # -------------------------
        # ROLE SESSION ISOLATION
        # -------------------------

        session.pop("student_id", None)
        session.pop("student_enrollment", None)
        session.pop("student_course", None)
        session.pop("student_semester", None)
        session.pop("student_section", None)
        session.pop("faculty_id", None)
        session.pop("faculty_name", None)

        # Admin session
        session["admin_id"] = admin_data["admin_id"]
        session["admin_department"] = admin_data["department_name"]
        session["admin_department_id"] = admin_data["department_id"]

        return redirect("/admin-dashboard")

    return render_template(
        "admin_login.html",
        departments=departments
    )


@app.route("/admin-change-password", methods=["GET", "POST"])
def admin_change_password():

    # Logged-in Admin check
    if not session.get("admin_id"):
        return redirect(url_for("admin_login"))

    message = ""

    admin_id = session.get("admin_id")
    admin_department = session.get("admin_department")

    if request.method == "POST":

        current_password = request.form.get(
            "current_password", ""
        ).strip()

        new_password = request.form.get(
            "new_password", ""
        ).strip()

        confirm_password = request.form.get(
            "confirm_password", ""
        ).strip()

        conn = get_connection()

        try:

            with conn.cursor() as cur:

                cur.execute("""
                    SELECT
                        admin_id,
                        department_name,
                        password
                    FROM department_admins
                    WHERE admin_id = %s
                      AND department_name = %s
                """, (
                    admin_id,
                    admin_department
                ))

                admin_data = cur.fetchone()

                if not admin_data:

                    message = "Admin account not found."

                elif not check_password_hash(
                    admin_data["password"],
                    current_password
                ):

                    message = "Current password is incorrect."

                elif new_password != confirm_password:

                    message = "New passwords do not match."

                elif len(new_password) < 6:

                    message = "New password must be at least 6 characters."

                elif check_password_hash(
                    admin_data["password"],
                    new_password
                ):

                    message = "New password must be different from the current password."

                else:

                    new_password_hash = generate_password_hash(
                        new_password
                    )

                    cur.execute("""
                        UPDATE department_admins
                        SET password = %s
                        WHERE admin_id = %s
                          AND department_name = %s
                    """, (
                        new_password_hash,
                        admin_id,
                        admin_department
                    ))

                    conn.commit()

                    message = "Admin password changed successfully."

        finally:
            conn.close()

    return render_template(
        "admin_change_password.html",
        message=message,
        admin_id=admin_id,
        admin_department=admin_department
    )

@app.route("/admin-logout")
def admin_logout():
    # Clear the complete role session so another portal cannot
    # inherit stale authentication data from the Admin portal.
    session.pop("admin_id", None)
    session.pop("student_id", None)
    session.pop("student_enrollment", None)
    session.pop("student_course", None)
    session.pop("student_semester", None)
    session.pop("student_section", None)
    session.pop("faculty_id", None)
    session.pop("faculty_name", None)
    return redirect(url_for("admin_login"))


# =========================
# ADMIN DASHBOARD
# =========================



@app.route("/admin-dashboard")
def admin_dashboard():

    # =========================
    # ADMIN LOGIN CHECK
    # =========================

    security_check = admin_required()

    if security_check:
        return security_check

    # =========================
    # STUDENTS COUNT
    # =========================

    students = []

    try:
        if os.path.exists(STUDENTS_FILE):

            with open(STUDENTS_FILE, "r") as file:
                students = json.load(file)

            if not isinstance(students, list):
                students = []

    except Exception:
        students = []

    total_students = len(students)


    # =========================
    # FACULTY COUNT
    # =========================

    faculty_list = []

    try:
        if os.path.exists(FACULTY_FILE):

            with open(FACULTY_FILE, "r") as file:
                faculty_list = json.load(file)

            if not isinstance(faculty_list, list):
                faculty_list = []

    except Exception:
        faculty_list = []

    total_faculty = len(faculty_list)


    # =========================
    # DASHBOARD
    # =========================

    return render_template(
        "admin_dashboard.html",
        total_students=total_students,
        total_faculty=total_faculty
    )


# =========================
# DCEA STUDENT ADMISSION CARDS
# =========================

@app.route("/student-admission-cards")
def student_admission_cards():

    # =========================
    # ADMIN LOGIN CHECK
    # =========================

    security_check = admin_required()

    if security_check:
        return security_check

    # =========================
    # ONLY DCEA ACCESS
    # =========================

    if session.get("admin_department") != "DCEA":
        return "Access Denied"

    # =========================
    # LOAD STUDENTS
    # =========================

    students = []

    try:
        students = json_load(
            os.path.basename(STUDENTS_FILE)
        )

        if not isinstance(students, list):
            students = []

    except Exception:
        students = []

    # =========================
    # ONLY DCEA STUDENTS
    # =========================

    dcea_students = []

    for student in students:

        if student.get("department") == "DCEA":

            dcea_students.append(student)

    # =========================
    # ADMISSION CARD PAGE
    # =========================

    return render_template(
        "student_admission_cards.html",
        students=dcea_students
    )

# =========================================================
# HEAD - DELETE WHOLE ATTENDANCE RECORD
# =========================================================

@app.route("/head-delete-whole-attendance")
def head_delete_whole_attendance():

    # =========================
    # HEAD LOGIN CHECK
    # =========================

    if not session.get("head_logged_in"):
        return redirect(url_for("head_login"))


    # =========================
    # ATTENDANCE KA POORA RECORD DELETE
    # =========================

    try:

        with open("attendance.json", "w") as file:
            json.dump([], file, indent=4)

    except Exception as e:

        return f"Unable to delete attendance record: {e}"


    return redirect("/head")




@app.route("/admin-search-faculty", methods=["GET"])
def admin_search_faculty():

    # ==============================
    # ADMIN LOGIN CHECK
    # ==============================

    security_check = admin_required()

    if security_check:
        return security_check


    faculty_list = []

    # ==============================
    # Faculty data load
    # ==============================

    if os.path.exists(FACULTY_FILE):

        try:
            with open(FACULTY_FILE, "r") as file:
                faculty_list = json.load(file)

            if not isinstance(faculty_list, list):
                faculty_list = []

        except:
            faculty_list = []


    # ==============================
    # Search value
    # ==============================

    query = request.args.get(
        "q",
        ""
    ).strip().lower()


    # ==============================
    # Search Faculty
    # ==============================

    results = []

    if query:

        for faculty in faculty_list:

            faculty_id = str(
                faculty.get(
                    "faculty_id",
                    ""
                )
            ).strip().lower()

            name = str(
                faculty.get(
                    "name",
                    ""
                )
            ).strip().lower()

            department = str(
                faculty.get(
                    "department",
                    ""
                )
            ).strip().lower()


            if (
                query in faculty_id
                or query in name
                or query in department
            ):

                results.append(faculty)


    # ==============================
    # Search result
    # ==============================

    return render_template(
        "admin_search_faculty.html",
        faculty_list=results,
        query=query
    )

# =========================
# ADMIN STUDENT MANAGEMENT
# =========================

STUDENTS_FILE = "students.json"


# =========================================================
# ADMIN STUDENTS
# =========================================================

@app.route("/admin-students", methods=["GET", "POST"])
def admin_students():

    # =====================================================
    # ADMIN LOGIN CHECK
    # =====================================================

    security_check = admin_required()

    if security_check:
        return security_check

    students = []
    courses = get_admin_courses()

    # =====================================================
    # LOAD DEPARTMENTS
    # =====================================================

    departments = []

    if os.path.exists(DEPARTMENTS_FILE):

        try:

            with open(
                DEPARTMENTS_FILE,
                "r",
                encoding="utf-8"
            ) as file:

                departments = json.load(file)

            if not isinstance(departments, list):
                departments = []

        except Exception:

            departments = []

    # =====================================================
    # LOAD STUDENTS
    # =====================================================

    if os.path.exists(STUDENTS_FILE):

        try:

            with open(
                STUDENTS_FILE,
                "r",
                encoding="utf-8"
            ) as file:

                students = json.load(file)

            if not isinstance(students, list):
                students = []

        except Exception:

            students = []

    # =====================================================
    # COURSE NAME HELPER
    # =====================================================

    def get_course_name(item):

        if isinstance(item, dict):

            return str(
                item.get("name")
                or item.get("course_name")
                or item.get("course")
                or item.get("Course Name")
                or ""
            ).strip()

        return str(item).strip()

    # =====================================================
    # ENROLLMENT SORTING
    # =====================================================

    def enrollment_sort_key(student):

        enrollment = str(
            student.get(
                "enrollment",
                ""
            )
        ).strip()

        try:

            return (
                0,
                int(enrollment)
            )

        except (
            ValueError,
            TypeError
        ):

            return (
                1,
                enrollment.lower()
            )

    # =====================================================
    # ADD NEW STUDENT
    # =====================================================

    if request.method == "POST":

        student_name = request.form.get(
            "student_name",
            ""
        ).strip()

        enrollment = request.form.get(
            "enrollment",
            ""
        ).strip()

        # =================================================
        # NEW: DEPARTMENT
        # =================================================

        department = request.form.get(
            "department",
            ""
        ).strip()

        course = request.form.get(
            "course",
            ""
        ).strip()

        semester = request.form.get(
            "semester",
            ""
        ).strip()

        branch = request.form.get(
            "branch",
            ""
        ).strip().upper()

        section = request.form.get(
            "section",
            ""
        ).strip().upper()

        group = request.form.get(
            "group",
            ""
        ).strip().upper()

        # =================================================
        # REQUIRED FIELDS
        # =================================================

        if (
            not student_name
            or not enrollment
            or not department
            or not course
            or not semester
        ):

            return (
                "Please fill all required fields.",
                400
            )

        # =================================================
        # DEPARTMENT VALIDATION
        # =================================================

        valid_department = None

        for item in departments:

            if isinstance(item, dict):

                department_name = str(
                    item.get("name", "")
                ).strip()

                if (
                    department_name.lower()
                    == department.lower()
                ):

                    valid_department = department_name

                    break

        if not valid_department:

            return (
                "Invalid Department.",
                400
            )

        department = valid_department

        # =================================================
        # SEMESTER VALIDATION
        # =================================================

        if semester not in [
            "1",
            "2",
            "3",
            "4",
            "5",
            "6",
            "7",
            "8",
            "9",
            "10"
        ]:

            return (
                "Invalid Semester.",
                400
            )

        # =================================================
        # COURSE VALIDATION
        # =================================================

        valid_course = None

        for item in courses:

            if isinstance(item, dict):

                name = get_course_name(item)

                course_department = str(
                    item.get(
                        "department",
                        ""
                    )
                ).strip()

                if (
                    name
                    and
                    name.lower()
                    == course.lower()
                    and
                    course_department.lower()
                    == department.lower()
                ):

                    valid_course = name

                    break

            else:

                name = get_course_name(item)

                if (
                    name
                    and
                    name.lower()
                    == course.lower()
                ):

                    valid_course = name

                    break

        if not valid_course:

            return (
                "Invalid Course for selected Department.",
                400
            )

        # Save official course name
        course = valid_course

        # =================================================
        # BRANCH VALIDATION
        # =================================================

        branches = [

            "CSE",

            "MECHANICAL",

            "MECHANICAL WITH AI & ML",

            "ELECTRICAL",

            "CIVIL",

            "ECE",

            "CSE WITH AI & ML",

            "CSE WITH CYBER SECURITY",

            "CSE WITH DATA SCIENCE"

        ]

        is_btech_all = (
            course.strip().lower()
            == "b.tech all"
        )

        # Branch required only for
        # B.Tech All + Semester 1 or 2

        if (
            is_btech_all
            and semester in [
                "1",
                "2"
            ]
        ):

            if not branch:

                return (
                    "Please select Branch.",
                    400
                )

            if branch not in branches:

                return (
                    "Invalid Branch.",
                    400
                )

        else:

            # Other courses/semesters
            # do not save branch.

            branch = ""

        # =================================================
        # SECTION VALIDATION
        # =================================================

        # Section is optional.

        if section and section not in [
            "A",
            "B",
            "C",
            "D"
        ]:

            return (
                "Invalid Section.",
                400
            )

        # =================================================
        # GROUP VALIDATION
        # =================================================

        # Group is optional.

        if group and group not in [
            "G1",
            "G2",
            "G3",
            "G4"
        ]:

            return (
                "Invalid Group.",
                400
            )

        # =================================================
        # DUPLICATE ENROLLMENT
        # =================================================

        for old_student in students:

            if not isinstance(
                old_student,
                dict
            ):
                continue

            old_enrollment = str(
                old_student.get(
                    "enrollment",
                    ""
                )
            ).strip()

            if (
                old_enrollment.lower()
                == enrollment.lower()
            ):

                return (
                    "This Enrollment Number "
                    "already exists.",
                    400
                )

        # =================================================
        # DCEA SEMESTER 1/2 ADMISSION FLOW
        # =================================================

        admission_type = request.form.get(
            "admission_type",
            ""
        ).strip()

        admission_date = request.form.get(
            "admission_date",
            ""
        ).strip()

        application_no = request.form.get(
            "application_no",
            ""
        ).strip()

        admission_card = request.files.get(
            "admission_card"
        )

        is_dcea_sem_1_2 = (
            department.strip().lower() == "dcea"
            and semester in ["1", "2"]
        )

        # -------------------------------------------------
        # FIRST STEP: OPEN ADMISSION PROGRAM PAGE
        # -------------------------------------------------

        if (
            is_dcea_sem_1_2
            and not admission_type
        ):

            return render_template(
                "student_admission_type.html",

                form_action="/admin-students",

                student_name=student_name,

                enrollment=enrollment,

                department=department,

                course=course,

                semester=semester,

                branch=branch,

                section=section,

                group=group
            )

        # -------------------------------------------------
        # DCEA SEMESTER 1/2 FINAL VALIDATION
        # -------------------------------------------------

        if is_dcea_sem_1_2:

            if admission_type not in [
                "orientation",
                "after_orientation"
            ]:

                return (
                    "Invalid admission type.",
                    400
                )

            if not admission_date:

                return (
                    "Please fill Date of Admission.",
                    400
                )

            try:

                from datetime import date

                admission_date_obj = (
                    date.fromisoformat(
                        admission_date
                    )
                )

            except ValueError:

                return (
                    "Invalid admission date.",
                    400
                )

            orientation_start = date(
                2026,
                6,
                1
            )

            orientation_end = date(
                2026,
                8,
                23
            )

            after_orientation_start = date(
                2026,
                8,
                24
            )

            # -------------------------------------------------
            # BEFORE 1 JUNE 2026
            # -------------------------------------------------

            if (
                admission_date_obj
                < orientation_start
            ):

                return (
                    "Admission date cannot be before 1 June 2026.",
                    400
                )

            # -------------------------------------------------
            # ORIENTATION PROGRAM
            # -------------------------------------------------

            if admission_type == "orientation":

                if (
                    admission_date_obj
                    >= after_orientation_start
                ):

                    return (
                        "Please go to After Orientation Program side.",
                        400
                    )

            # -------------------------------------------------
            # AFTER ORIENTATION PROGRAM
            # -------------------------------------------------

            if admission_type == "after_orientation":

                if (
                    admission_date_obj
                    <= orientation_end
                ):

                    return (
                        "Please use Orientation Program side for this date.",
                        400
                    )

                if not application_no:

                    return (
                        "Please enter Application No.",
                        400
                    )

                if not admission_card:

                    return (
                        "Please upload Admission Card PDF.",
                        400
                    )

                if not admission_card.filename:

                    return (
                        "Please upload Admission Card PDF.",
                        400
                    )

                if not admission_card.filename.lower().endswith(
                    ".pdf"
                ):

                    return (
                        "Only PDF Admission Card is allowed.",
                        400
                    )

        else:

            # Existing students:
            # No admission flow.

            admission_type = ""

            admission_date = ""

            application_no = ""

            admission_card = None

        # =================================================
        # CREATE NEXT STUDENT ID
        # =================================================

        numbers = []

        for old_student in students:

            if not isinstance(
                old_student,
                dict
            ):
                continue

            old_id = str(
                old_student.get(
                    "student_id",
                    ""
                )
            ).strip().upper()

            if old_id.startswith("STU"):

                try:

                    number = int(
                        old_id[3:]
                    )

                    numbers.append(
                        number
                    )

                except Exception:

                    pass
            

        # =================================================
        # CREATE NEXT STUDENT ID
        # =================================================

        numbers = []

        for old_student in students:

            if not isinstance(
                old_student,
                dict
            ):

                continue

            old_id = str(
                old_student.get(
                    "student_id",
                    ""
                )
            ).strip().upper()

            if old_id.startswith("STU"):

                try:

                    number = int(
                        old_id[3:]
                    )

                    numbers.append(
                        number
                    )

                except Exception:

                    pass

        if numbers:

            next_number = (
                max(numbers)
                + 1
            )

        else:

            next_number = 1

        student_id = (
            "STU"
            + str(
                next_number
            ).zfill(3)
        )

        # =================================================
        # CREATE STUDENT
        # =================================================

        student = {

            "student_id":
                student_id,

            "name":
                student_name,

            "enrollment":
                enrollment,

            "department":
                department,

            "course":
                course,

            "semester":
                semester,

            "branch":
                branch,

            "section":
                section,

            "group":
                group,

            "admission_type":
                admission_type,

            "admission_date":
                admission_date,

            "application_no":
                application_no,

            "admission_card":
                "",

            "password_hash":
                generate_password_hash(
                    student_id
                )

        }

        # =================================================
        # ADD TO LIST
        # =================================================

        if (
            is_dcea_sem_1_2
            and admission_type == "after_orientation"
        ):
            try:

                card_filename = (
                    "admission_card_" + student_id + ".pdf"
                )

                save_uploaded_file(
                    admission_card,
                    card_filename
                )

                student["admission_card"] = card_filename

            except Exception as e:

                return (
                    f"Admission Card could not be saved: {e}",
                    500
                )

        students.append(
            student
        )

        # =================================================
        # SORT BEFORE SAVE
        # =================================================

        students.sort(
            key=enrollment_sort_key
        )

        # =================================================
        # SAVE STUDENTS
        # =================================================

        with open(
            STUDENTS_FILE,
            "w",
            encoding="utf-8"
        ) as file:

            json.dump(
                students,
                file,
                indent=4,
                ensure_ascii=False
            )

        # =================================================
        # REDIRECT
        # =================================================

        return redirect(
            "/admin-students"
        )

    # =====================================================
    # COURSE → SEMESTER → SECTION → STUDENTS
    # =====================================================

    course_structure = {}

    for student in students:

        if not isinstance(
            student,
            dict
        ):

            continue

        student_course = str(
            student.get(
                "course",
                ""
            )
        ).strip()

        # =================================================
        # OLD DATA SUPPORT
        # =================================================

        if not student_course:

            student_course = str(
                student.get(
                    "department",
                    ""
                )
            ).strip()

        semester = str(
            student.get(
                "semester",
                ""
            )
        ).strip()

        section = str(
            student.get(
                "section",
                ""
            )
        ).strip().upper()

        # =================================================
        # INVALID RECORDS
        # =================================================

        if not student_course:

            continue

        if semester not in [
            "1",
            "2",
            "3",
            "4",
            "5",
            "6",
            "7",
            "8",
            "9",
            "10"
        ]:

            continue

        # =================================================
        # OPTIONAL SECTION
        # =================================================

        if section not in [
            "A",
            "B",
            "C",
            "D"
        ]:

            section = "No Section"

        course_key = (
            student_course
            .lower()
            .strip()
        )

        # =================================================
        # COURSE
        # =================================================

        if course_key not in course_structure:

            course_structure[
                course_key
            ] = {

                "name":
                    student_course,

                "semesters":
                    {}

            }

        # =================================================
        # SEMESTER
        # =================================================

        if semester not in course_structure[
            course_key
        ]["semesters"]:

            course_structure[
                course_key
            ]["semesters"][
                semester
            ] = {}

        # =================================================
        # SECTION
        # =================================================

        if section not in course_structure[
            course_key
        ]["semesters"][
            semester
        ]:

            course_structure[
                course_key
            ]["semesters"][
                semester
            ][section] = []

        course_structure[
            course_key
        ]["semesters"][
            semester
        ][section].append(
            student
        )

    # =====================================================
    # SORT COURSE STRUCTURE
    # =====================================================

    sorted_course_structure = {}

    for course_key in sorted(
        course_structure.keys()
    ):

        course_data = (
            course_structure[
                course_key
            ]
        )

        sorted_semesters = {}

        for semester_number in sorted(
            course_data[
                "semesters"
            ].keys(),

            key=lambda value:
                int(value)
        ):

            semester_data = (
                course_data[
                    "semesters"
                ][semester_number]
            )

            section_order = [

                "A",

                "B",

                "C",

                "D",

                "No Section"

            ]

            sorted_sections = {}

            for section_name in section_order:

                if (
                    section_name
                    in semester_data
                    and semester_data[
                        section_name
                    ]
                ):

                    sorted_sections[
                        section_name
                    ] = sorted(

                        semester_data[
                            section_name
                        ],

                        key=enrollment_sort_key
                    )

            if sorted_sections:

                sorted_semesters[
                    semester_number
                ] = sorted_sections

        if sorted_semesters:

            sorted_course_structure[
                course_key
            ] = {

                "name":
                    course_data[
                        "name"
                    ],

                "semesters":
                    sorted_semesters

            }

    course_structure = (
        sorted_course_structure
    )

    # =====================================================
    # SORT MAIN STUDENT LIST
    # =====================================================

    students.sort(
        key=enrollment_sort_key
    )

    # =====================================================
    # SORT DISPLAYED STUDENTS
    # =====================================================

    for course_data in (
        course_structure.values()
    ):

        for semester_data in (
            course_data.get(
                "semesters",
                {}
            ).values()
        ):

            for section_students in (
                semester_data.values()
            ):

                section_students.sort(
                    key=enrollment_sort_key
                )

    # =====================================================
    # SHOW ADMIN STUDENTS PAGE
    # =====================================================

    return render_template(

        "admin_students.html",

        students=students,

        courses=courses,

        departments=departments,

        course_structure=course_structure

    )

@app.route(
    "/student-change-password",
    methods=["GET", "POST"]
)
def student_change_password():

    # Check student login
    if "student_enrollment" not in session:
        return redirect("/student-login")

    enrollment = session[
        "student_enrollment"
    ]

    students = []

    # students.json read
    if os.path.exists(
        STUDENTS_FILE
    ):

        try:

            with open(
                STUDENTS_FILE,
                "r"
            ) as file:

                students = json.load(
                    file
                )

        except:

            students = []

    student = None

    # Logged-in student find
    for item in students:

        if (
            str(
                item.get(
                    "enrollment",
                    ""
                )
            ).strip()
            ==
            str(
                enrollment
            ).strip()
        ):

            student = item
            break

    if student is None:
        return "Student not found."

    # -----------------------------
    # PASSWORD CHANGE
    # -----------------------------
    if request.method == "POST":

        current_password = request.form.get(
            "current_password",
            ""
        )

        new_password = request.form.get(
            "new_password",
            ""
        )

        confirm_password = request.form.get(
            "confirm_password",
            ""
        )

        # Empty fields
        if (
            not current_password
            or not new_password
            or not confirm_password
        ):
            return "Please fill all fields."

        # Current password check
        if not check_password_hash(
            student.get(
                "password_hash",
                ""
            ),
            current_password
        ):
            return "Current password is incorrect."

        # New password confirmation
        if (
            new_password
            != confirm_password
        ):
            return "New passwords do not match."

        # New password same as old
        if (
            current_password
            == new_password
        ):
            return "New password must be different from current password."

        # New password save as HASH
        student[
            "password_hash"
        ] = generate_password_hash(
            new_password
        )

        # students.json update
        with open(
            STUDENTS_FILE,
            "w"
        ) as file:

            json.dump(
                students,
                file,
                indent=4
            )

        return redirect(
            "/student-dashboard"
        )

    return render_template(
        "student_change_password.html"
    )



@app.route("/edit-student/<student_id>", methods=["GET", "POST"])
def edit_student(student_id):

    security_check = admin_required()

    if security_check:
        return security_check

    students = []

    # =========================
    # LOAD STUDENTS
    # =========================

    if os.path.exists(STUDENTS_FILE):

        try:

            with open(STUDENTS_FILE, "r") as file:
                students = json.load(file)

            if not isinstance(students, list):
                students = []

        except Exception:

            students = []


    # =========================
    # FIND STUDENT
    # =========================

    student = None

    for item in students:

        if str(
            item.get("student_id", "")
        ).strip() == str(student_id).strip():

            student = item
            break


    if student is None:

        return "Student not found."


    # =========================
    # ADMIN COURSES
    # =========================

    courses = get_admin_courses()


    def get_course_name(item):

        if isinstance(item, dict):

            return str(
                item.get("name")
                or item.get("course_name")
                or item.get("course")
                or item.get("Course Name")
                or ""
            ).strip()

        return str(item).strip()


    valid_courses = []

    for item in courses:

        name = get_course_name(item)

        if name:
            valid_courses.append(name)


    # =========================
    # UPDATE STUDENT
    # =========================

    if request.method == "POST":

        student_name = request.form.get(
            "student_name",
            ""
        ).strip()


        enrollment = request.form.get(
            "enrollment",
            ""
        ).strip()


        course = request.form.get(
            "course",
            ""
        ).strip()


        semester = request.form.get(
            "semester",
            ""
        ).strip()


        branch = request.form.get(
            "branch",
            ""
        ).strip().upper()


        section = request.form.get(
            "section",
            ""
        ).strip().upper()


        group = request.form.get(
            "group",
            ""
        ).strip().upper()


        # =========================
        # REQUIRED FIELDS
        # =========================

        if (
            not student_name
            or not enrollment
            or not course
            or not semester
        ):

            return "Please fill all required fields."


        # =========================
        # SEMESTER VALIDATION
        # =========================

        if semester not in [
            "1",
            "2",
            "3",
            "4",
            "5",
            "6",
            "7",
            "8"
        ]:

            return "Invalid Semester."


        # =========================
        # SECTION VALIDATION
        # =========================

        if section and section not in [
            "A",
            "B",
            "C",
            "D"
        ]:

            return "Invalid Section."


        # =========================
        # GROUP VALIDATION
        # =========================

        if group and group not in [
            "G1",
            "G2",
            "G3",
            "G4"
        ]:

            return "Invalid Group."


        # =========================
        # BRANCH VALIDATION
        # =========================

        valid_branches = [

            "CSE",

            "MECHANICAL",

            "MECHANICAL WITH AI & ML",

            "ELECTRICAL",

            "CIVIL",

            "ECE",

            "CSE WITH AI & ML",

            "CSE WITH CYBER SECURITY",

            "CSE WITH DATA SCIENCE"

        ]


        # ==========================================
        # B.TECH ALL BRANCH RULE
        # ==========================================

        if (
            course.strip().upper() == "B.TECH ALL"
            and semester in ["1", "2"]
        ):

            if not branch:

                return "Please select Branch for B.TECH ALL in Semester 1 or 2."


            if branch not in valid_branches:

                return "Invalid Branch."


        else:

            # Semester 3 onwards / other courses
            # branch completely remove

            branch = ""


        # =========================
        # COURSE VALIDATION
        # =========================

        course_found = False


        for valid_course in valid_courses:

            if (
                valid_course.lower()
                == course.lower()
            ):

                course_found = True

                # Admin Course Management ka
                # exact course name save hoga

                course = valid_course

                break


        if not course_found:

            return "Invalid Course."


        # =========================
        # DUPLICATE ENROLLMENT CHECK
        # =========================

        for item in students:

            if (

                str(
                    item.get(
                        "student_id",
                        ""
                    )
                ).strip()
                != str(student_id).strip()

                and

                str(
                    item.get(
                        "enrollment",
                        ""
                    )
                ).strip().lower()
                == enrollment.lower()

            ):

                return (
                    "This Enrollment Number "
                    "already exists."
                )


        # =========================
        # UPDATE STUDENT
        # =========================

        student["name"] = student_name

        student["enrollment"] = enrollment

        student["course"] = course

        student["semester"] = semester


        # =========================
        # BRANCH SAVE / REMOVE
        # =========================

        if (
            course.strip().upper()
            == "B.TECH ALL"
            and semester in ["1", "2"]
            and branch
        ):

            student["branch"] = branch

        else:

            student.pop(
                "branch",
                None
            )


        # =========================
        # OPTIONAL SECTION
        # =========================

        if section:

            student["section"] = section

        else:

            student.pop(
                "section",
                None
            )


        # =========================
        # OPTIONAL GROUP
        # =========================

        if group:

            student["group"] = group

        else:

            student.pop(
                "group",
                None
            )


        # =========================
        # SAVE
        # =========================

        with open(
            STUDENTS_FILE,
            "w"
        ) as file:

            json.dump(
                students,
                file,
                indent=4,
                ensure_ascii=False
            )


        # =========================
        # RETURN
        # =========================

        return redirect(
            "/admin-students"
        )


    # =========================
    # EDIT PAGE
    # =========================

    return render_template(
        "edit_student.html",
        student=student,
        courses=valid_courses
    )

@app.route("/reset-student-password/<student_id>")
@app.route("/admin-reset-student-password/<student_id>")
def reset_student_password(student_id):

    security_check = admin_required()

    if security_check:
        return security_check

    students = []

    # students.json read
    if os.path.exists(STUDENTS_FILE):

        try:
            with open(STUDENTS_FILE, "r") as file:
                students = json.load(file)

        except:
            students = []

    student_found = False

    # Student find
    for student in students:

        if str(student.get("student_id", "")).strip() == str(student_id).strip():

            # Student ID ko new password banakar hash karna
            student["password_hash"] = generate_password_hash(
                student_id
            )

            student_found = True

            break

    if not student_found:
        return "Student not found."

    # Updated data save
    with open(STUDENTS_FILE, "w") as file:

        json.dump(
            students,
            file,
            indent=4
        )

    return redirect("/admin-students")

@app.route("/admin-search-student")
def admin_search_student():

    security_check = admin_required()

    if security_check:
        return security_check

    search = request.args.get("q", "").strip().lower()

    students = []

    if os.path.exists(STUDENTS_FILE):

        try:
            with open(STUDENTS_FILE, "r") as file:
                students = json.load(file)

        except:
            students = []

    results = []

    if search:

        for student in students:

            name = str(student.get("name", "")).lower()
            enrollment = str(student.get("enrollment", "")).lower()
            student_id = str(student.get("student_id", "")).lower()

            if (
                search in name
                or search in enrollment
                or search in student_id
            ):
                results.append(student)

    return render_template(
        "admin_student_search.html",
        results=results,
        search=search
    )

def fix_student_ids():

    if not os.path.exists(STUDENTS_FILE):
        print("students.json file nahi mili.")
        return

    try:
        with open(STUDENTS_FILE, "r") as file:
            students = json.load(file)
    except:
        print("students.json read nahi ho rahi.")
        return

    # Har student ko unique sequential ID do
    for i, student in enumerate(students, start=1):
        student["student_id"] = "STU" + str(i).zfill(3)

    with open(STUDENTS_FILE, "w") as file:
        json.dump(students, file, indent=4)

    print("Student IDs successfully fixed.")
    print("Total students:", len(students))

# =========================
# DELETE STUDENT
# =========================

@app.route("/delete-student/<student_id>")
@app.route("/admin-delete-student/<student_id>")
def delete_student(student_id):

    security_check = admin_required()

    if security_check:
        return security_check

    students = []

    if os.path.exists(STUDENTS_FILE):

        try:
            with open(STUDENTS_FILE, "r") as file:
                students = json.load(file)
        except:
            students = []

    # जिस student की ID match हो उसे छोड़कर बाकी रखेंगे
    students = [
        student
        for student in students
        if str(student.get("student_id", "")).strip() != str(student_id).strip()
    ]

    # Updated data save
    with open(STUDENTS_FILE, "w") as file:
        json.dump(
            students,
            file,
            indent=4
        )

    return redirect("/admin-students")

# =========================
# ADMIN FACULTY MANAGEMENT
# =========================

FACULTY_FILE = "faculty.json"


@app.route("/admin-faculty", methods=["GET", "POST"])
def admin_faculty():

    security_check = admin_required()

    if security_check:
        return security_check

    faculty_list = []
    departments = []

    # ==============================
    # Faculty data load
    # ==============================

    if os.path.exists(FACULTY_FILE):

        try:
            with open(FACULTY_FILE, "r") as file:
                faculty_list = json.load(file)

            if not isinstance(faculty_list, list):
                faculty_list = []

        except:
            faculty_list = []


    # ==============================
    # Department data load
    # ==============================

    if os.path.exists(DEPARTMENTS_FILE):

        try:
            with open(DEPARTMENTS_FILE, "r") as file:
                departments = json.load(file)

            if not isinstance(departments, list):
                departments = []

        except:
            departments = []


    # ==============================
    # ADD FACULTY
    # ==============================

    if request.method == "POST":

        name = request.form.get(
            "name", ""
        ).strip()

        department = request.form.get(
            "department", ""
        ).strip()


        if not name or not department:

            return "Please fill all fields."


        # ==============================
        # Validate Department
        # ==============================

        valid_department = None

        for dept in departments:

            dept_name = str(
                dept.get("name", "")
            ).strip()

            if dept_name.lower() == department.lower():

                valid_department = dept_name
                break


        if not valid_department:

            return "Invalid Department."


        # ==============================
        # Generate Faculty ID
        # ==============================

        existing_ids = []

        for faculty in faculty_list:

            existing_ids.append(
                str(
                    faculty.get(
                        "faculty_id", ""
                    )
                ).strip()
            )


        number = len(faculty_list) + 1

        faculty_id = (
            "FAC"
            + str(number).zfill(3)
        )


        while faculty_id in existing_ids:

            number += 1

            faculty_id = (
                "FAC"
                + str(number).zfill(3)
            )


        # ==============================
        # Initial Password
        # ==============================

        password = faculty_id


        # ==============================
        # Faculty data
        # ==============================

        faculty = {

            "faculty_id": faculty_id,

            "name": name,

            "department": valid_department,

            "password": password

        }


        faculty_list.append(faculty)


        # ==============================
        # Save Faculty
        # ==============================

        with open(
            FACULTY_FILE,
            "w"
        ) as file:

            json.dump(
                faculty_list,
                file,
                indent=4
            )


        return redirect("/admin-faculty")


    # ==============================
    # FACULTY SEARCH
    # ==============================

    search = request.args.get(
        "search",
        ""
    ).strip().lower()


    if search:

        filtered_faculty = []

        for faculty in faculty_list:

            faculty_id = str(
                faculty.get(
                    "faculty_id",
                    ""
                )
            ).lower()

            name = str(
                faculty.get(
                    "name",
                    ""
                )
            ).lower()

            department = str(
                faculty.get(
                    "department",
                    ""
                )
            ).lower()


            if (
                search in faculty_id
                or search in name
                or search in department
            ):

                filtered_faculty.append(
                    faculty
                )


        faculty_list = filtered_faculty


    # ==============================
    # Faculty page
    # ==============================

    return render_template(
        "admin_faculty.html",
        faculty_list=faculty_list,
        departments=departments,
        search=search
    )


# =========================
# DELETE FACULTY
# =========================

@app.route("/delete-faculty/<faculty_id>")
def delete_faculty(faculty_id):

    security_check = admin_required()

    if security_check:
        return security_check

    faculty_list = []

    if os.path.exists(FACULTY_FILE):

        try:
            with open(FACULTY_FILE, "r") as file:
                faculty_list = json.load(file)
        except:
            faculty_list = []

    faculty_list = [
        faculty
        for faculty in faculty_list
        if faculty.get("faculty_id") != faculty_id
    ]

    with open(FACULTY_FILE, "w") as file:

        json.dump(
            faculty_list,
            file,
            indent=4
        )

    return redirect("/admin-faculty")

@app.route("/admin-departments", methods=["GET", "POST"])
def admin_departments():

    security_check = admin_required()

    if security_check:
        return security_check

    departments = []

    if os.path.exists(DEPARTMENTS_FILE):

        try:
            with open(DEPARTMENTS_FILE, "r") as file:
                departments = json.load(file)

        except:
            departments = []

    if request.method == "POST":

        department_name = request.form.get(
            "department_name", ""
        ).strip()

        if not department_name:
            return "Please enter department name."

        # Duplicate check
        for department in departments:

            if str(
                department.get("name", "")
            ).strip().lower() == department_name.lower():

                return "This Department already exists."

        department_id = "DEPT" + str(
            len(departments) + 1
        ).zfill(3)

        department = {
            "department_id": department_id,
            "name": department_name
        }

        departments.append(department)

        with open(DEPARTMENTS_FILE, "w") as file:

            json.dump(
                departments,
                file,
                indent=4
            )

        return redirect("/admin-departments")

    return render_template(
        "admin_departments.html",
        departments=departments
    )

@app.route("/admin-courses", methods=["GET", "POST"])
def admin_courses():

    security_check = admin_required()

    if security_check:
        return security_check

    courses = []
    departments = []

    # =========================
    # LOAD COURSES
    # =========================

    if os.path.exists(COURSES_FILE):

        try:

            with open(COURSES_FILE, "r") as file:
                courses = json.load(file)

            if not isinstance(courses, list):
                courses = []

        except Exception:

            courses = []


    # =========================
    # LOAD DEPARTMENTS
    # =========================

    if os.path.exists(DEPARTMENTS_FILE):

        try:

            with open(DEPARTMENTS_FILE, "r") as file:
                departments = json.load(file)

            if not isinstance(departments, list):
                departments = []

        except Exception:

            departments = []


    # =========================
    # ADD COURSE
    # =========================

    if request.method == "POST":

        course_name = request.form.get(
            "course_name",
            ""
        ).strip()

        department = request.form.get(
            "department",
            ""
        ).strip()

        total_semesters = request.form.get(
            "total_semesters",
            ""
        ).strip()


        # =========================
        # VALIDATION
        # =========================

        if not course_name or not department or not total_semesters:

            return "Please fill all fields."


        try:

            total_semesters = int(total_semesters)

        except ValueError:

            return "Total semesters must be a number."


        if total_semesters < 1:

            return "Total semesters must be at least 1."


        # =========================
        # DUPLICATE COURSE CHECK
        # =========================

        for course in courses:

            if not isinstance(course, dict):
                continue

            if (
                str(course.get("name", "")).strip().lower()
                == course_name.lower()
                and
                str(course.get("department", "")).strip().lower()
                == department.lower()
            ):

                return "This Course already exists in this Department."


        # =========================
        # COURSE ID
        # =========================

        existing_ids = []

        for course in courses:

            existing_ids.append(
                str(
                    course.get(
                        "course_id",
                        ""
                    )
                ).strip()
            )


        number = len(courses) + 1

        course_id = "COURSE" + str(number).zfill(3)


        while course_id in existing_ids:

            number += 1

            course_id = "COURSE" + str(number).zfill(3)


        # =========================
        # CREATE COURSE
        # =========================

        course = {

            "course_id": course_id,

            "name": course_name,

            "department": department,

            "total_semesters": total_semesters

        }


        courses.append(course)


        # =========================
        # SAVE
        # =========================

        with open(
            COURSES_FILE,
            "w"
        ) as file:

            json.dump(
                courses,
                file,
                indent=4,
                ensure_ascii=False
            )


        return redirect("/admin-courses")


    # =========================
    # PAGE
    # =========================

    return render_template(
        "admin_courses.html",
        courses=courses,
        departments=departments
    )

@app.route("/edit-course/<course_id>", methods=["GET", "POST"])
def edit_course(course_id):

    security_check = admin_required()

    if security_check:
        return security_check

    courses = []
    departments = []

    # =========================
    # LOAD COURSES
    # =========================

    if os.path.exists(COURSES_FILE):

        try:

            with open(COURSES_FILE, "r") as file:
                courses = json.load(file)

            if not isinstance(courses, list):
                courses = []

        except Exception:

            courses = []


    # =========================
    # LOAD DEPARTMENTS
    # =========================

    if os.path.exists(DEPARTMENTS_FILE):

        try:

            with open(DEPARTMENTS_FILE, "r") as file:
                departments = json.load(file)

            if not isinstance(departments, list):
                departments = []

        except Exception:

            departments = []


    # =========================
    # FIND COURSE
    # =========================

    course = None

    for item in courses:

        if (
            str(
                item.get(
                    "course_id",
                    ""
                )
            ).strip()
            ==
            str(course_id).strip()
        ):

            course = item

            break


    if course is None:

        return "Course not found."


    # =========================
    # SAVE EDIT
    # =========================

    if request.method == "POST":

        course_name = request.form.get(
            "course_name",
            ""
        ).strip()

        department = request.form.get(
            "department",
            ""
        ).strip()

        total_semesters = request.form.get(
            "total_semesters",
            ""
        ).strip()


        if (
            not course_name
            or
            not department
            or
            not total_semesters
        ):

            return "Please fill all fields."


        try:

            total_semesters = int(
                total_semesters
            )

        except ValueError:

            return "Total semesters must be a number."


        if total_semesters < 1:

            return "Total semesters must be at least 1."


        # =========================
        # DUPLICATE CHECK
        # =========================

        for item in courses:

            if (
                str(
                    item.get(
                        "course_id",
                        ""
                    )
                ).strip()
                !=
                str(course_id).strip()
                and
                str(
                    item.get(
                        "name",
                        ""
                    )
                ).strip().lower()
                ==
                course_name.lower()
                and
                str(
                    item.get(
                        "department",
                        ""
                    )
                ).strip().lower()
                ==
                department.lower()
            ):

                return "This Course already exists in this Department."


        # =========================
        # UPDATE COURSE
        # =========================

        course["name"] = course_name

        course["department"] = department

        course["total_semesters"] = total_semesters


        # =========================
        # SAVE
        # =========================

        with open(
            COURSES_FILE,
            "w"
        ) as file:

            json.dump(
                courses,
                file,
                indent=4,
                ensure_ascii=False
            )


        return redirect("/admin-courses")


    # =========================
    # EDIT PAGE
    # =========================

    return render_template(
        "edit_course.html",
        course=course,
        departments=departments
    )


@app.route("/delete-course/<course_id>")
def delete_course(course_id):

    security_check = admin_required()

    if security_check:
        return security_check

    courses = []

    if os.path.exists(COURSES_FILE):
        try:
            with open(COURSES_FILE, "r") as file:
                courses = json.load(file)
        except:
            courses = []

    courses = [
        course
        for course in courses
        if str(course.get("course_id", "")).strip()
        != str(course_id).strip()
    ]

    with open(COURSES_FILE, "w") as file:
        json.dump(courses, file, indent=4)

    return redirect("/admin-courses")



@app.route("/head-semester-update")
def head_semester_update():

    # =========================
    # HEAD LOGIN CHECK
    # =========================

    if not session.get("head_logged_in"):
        return redirect(url_for("head_login"))


    # =========================
    # LOAD COURSES
    # =========================

    courses = []

    if os.path.exists(COURSES_FILE):

        try:

            with open(COURSES_FILE, "r") as file:
                courses = json.load(file)

            if not isinstance(courses, list):
                courses = []

        except Exception:

            courses = []


    # =========================
    # PAGE
    # =========================

    return render_template(
        "head_semester_update.html",
        courses=courses
    )


@app.route("/head-confirm-semester-update", methods=["POST"])
def head_confirm_semester_update():

    # =========================
    # HEAD LOGIN CHECK
    # =========================

    if not session.get("head_logged_in"):
        return redirect(url_for("head_login"))


    # =========================
    # LOAD COURSES
    # =========================

    try:

        with open(COURSES_FILE, "r") as file:
            courses = json.load(file)

        if not isinstance(courses, list):
            courses = []

    except Exception:

        courses = []


    # =========================
    # COURSE SEMESTER LIMITS
    # =========================

    course_limits = {}

    for course in courses:

        if not isinstance(course, dict):
            continue

        course_name = str(
            course.get("name", "")
        ).strip()

        course_id = str(
            course.get("course_id", "")
        ).strip()

        try:

            total_semesters = int(
                course.get("total_semesters")
            )

        except (TypeError, ValueError):

            continue

        if total_semesters < 1:
            continue

        if course_name:

            course_limits[
                course_name.lower()
            ] = total_semesters

        if course_id:

            course_limits[
                course_id.lower()
            ] = total_semesters


    # =========================
    # LOAD STUDENTS
    # =========================

    try:

        with open(STUDENTS_FILE, "r") as file:
            students = json.load(file)

        if not isinstance(students, list):
            students = []

    except Exception:

        students = []


    # =========================
    # COUNTERS
    # =========================

    updated_count = 0
    final_semester_count = 0
    skipped_count = 0


    # =========================
    # UPDATE EACH STUDENT
    # =========================

    for student in students:

        if not isinstance(student, dict):

            skipped_count += 1
            continue


        # =========================
        # FIND STUDENT COURSE
        # =========================

        student_course = str(
            student.get("course", "")
        ).strip()

        if not student_course:

            student_course = str(
                student.get("course_id", "")
            ).strip()


        if not student_course:

            skipped_count += 1
            continue


        # =========================
        # FIND COURSE LIMIT
        # =========================

        max_semester = course_limits.get(
            student_course.lower()
        )


        if max_semester is None:

            skipped_count += 1
            continue


        # =========================
        # CURRENT SEMESTER
        # =========================

        try:

            current_semester = int(
                str(
                    student.get(
                        "semester",
                        ""
                    )
                ).strip()
            )

        except (TypeError, ValueError):

            skipped_count += 1
            continue


        # =========================
        # FINAL SEMESTER
        # =========================

        if current_semester >= max_semester:

            final_semester_count += 1
            continue


        # =========================
        # ONLY ONE STEP FORWARD
        # =========================

        new_semester = current_semester + 1


        # Extra safety:
        # maximum semester se kabhi upar nahi jayega

        if new_semester > max_semester:

            new_semester = max_semester


        # =========================
        # SAVE ONLY SEMESTER
        # =========================

        student["semester"] = str(
            new_semester
        )

        updated_count += 1


    # =========================
    # SAVE STUDENTS
    # =========================

    try:

        with open(
            STUDENTS_FILE,
            "w"
        ) as file:

            json.dump(
                students,
                file,
                indent=4,
                ensure_ascii=False
            )

    except Exception as e:

        return f"""
        <h2>Semester Update Failed</h2>

        <p>
            {str(e)}
        </p>

        <br>

        <a href="/head">
            Back to Head Portal
        </a>
        """


    # =========================
    # SUCCESS PAGE
    # =========================

    return f"""
    <!DOCTYPE html>

    <html>

    <head>

        <title>Semester Update</title>

        <meta
            name="viewport"
            content="width=device-width, initial-scale=1.0"
        >

        <style>

            body {{
                font-family: Arial, sans-serif;
                background: #f1f5f9;
                display: flex;
                justify-content: center;
                align-items: center;
                min-height: 100vh;
                margin: 0;
                padding: 20px;
            }}

            .box {{
                background: white;
                width: 100%;
                max-width: 500px;
                padding: 30px;
                border-radius: 15px;
                text-align: center;
                box-shadow: 0 5px 20px #cbd5e1;
            }}

            h1 {{
                color: #166534;
            }}

            p {{
                color: #475569;
                line-height: 1.7;
            }}

            .btn {{
                display: inline-block;
                margin-top: 20px;
                padding: 12px 20px;
                background: #2563eb;
                color: white;
                text-decoration: none;
                border-radius: 8px;
                font-weight: bold;
            }}

        </style>

    </head>

    <body>

        <div class="box">

            <h1>
                ✅ Semester Updated
            </h1>

            <p>
                Students updated:
                <b>{updated_count}</b>
            </p>

            <p>
                Final semester students:
                <b>{final_semester_count}</b>
            </p>

            <p>
                Skipped records:
                <b>{skipped_count}</b>
            </p>

            <p>
                Every eligible student was moved
                <b>exactly one semester forward.</b>
            </p>

            <a
                href="/head"
                class="btn"
            >
                Back to Head Portal
            </a>

        </div>

    </body>

    </html>
    """

# =========================================================
# CREATE EXCEL SHEET
# =========================================================

@app.route("/create-excel-sheet", methods=["GET", "POST"])
def create_excel_sheet():

    # =====================================================
    # GET - CREATE EXCEL PAGE
    # =====================================================

    if request.method == "GET":

        courses = get_admin_courses()

        valid_courses = []

        for item in courses:

            if isinstance(item, dict):

                name = str(
                    item.get("name")
                    or item.get("course_name")
                    or item.get("course")
                    or item.get("Course Name")
                    or ""
                ).strip()

            else:

                name = str(item).strip()

            if name:
                valid_courses.append(name)

        return render_template(
            "create_excel_sheet.html",
            courses=valid_courses
        )

    # =====================================================
    # FORM DATA
    # =====================================================

    topic = request.form.get(
        "topic",
        ""
    ).strip()

    sheet_date = request.form.get(
        "sheet_date",
        ""
    ).strip()

    course = request.form.get(
        "course",
        ""
    ).strip()

    semester = request.form.get(
        "semester",
        ""
    ).strip()

    section = request.form.get(
        "section",
        ""
    ).strip().upper()

    # =====================================================
    # MANUAL ROWS
    # =====================================================

    try:

        manual_rows = int(
            request.form.get(
                "manual_rows",
                "10"
            )
        )

    except Exception:

        manual_rows = 10

    if manual_rows < 1:
        manual_rows = 1

    if manual_rows > 500:
        manual_rows = 500

    # =====================================================
    # MANUAL S.NO OPTION
    # =====================================================

    add_manual_sno = (
        request.form.get(
            "add_manual_sno",
            ""
        ) == "yes"
    )

    # =====================================================
    # COLUMN COUNT
    # =====================================================

    try:

        column_count = int(
            request.form.get(
                "column_count",
                "0"
            )
        )

    except Exception:

        column_count = 0

    # =====================================================
    # VALIDATION
    # =====================================================

    if not topic:

        return "Please enter Topic.", 400

    if not sheet_date:

        return "Please select Date.", 400

    if column_count < 1:

        return "Please select at least 1 column.", 400

    if column_count > 50:

        return "Maximum 50 columns allowed.", 400

    # =====================================================
    # GET USER COLUMNS
    # =====================================================

    columns = []

    for i in range(column_count):

        column_name = request.form.get(
            f"column_{i}",
            ""
        ).strip()

        if not column_name:

            return (
                f"Please enter name for Column {i + 1}.",
                400
            )

        columns.append(
            column_name
        )

    # =====================================================
    # FILTER MODE
    # =====================================================

    filters_selected = bool(
        course
        or semester
        or section
    )

    # =====================================================
    # LOAD STUDENTS
    # =====================================================

    all_students = []

    if os.path.exists(STUDENTS_FILE):

        try:

            with open(
                STUDENTS_FILE,
                "r",
                encoding="utf-8"
            ) as file:

                all_students = json.load(file)

            if not isinstance(
                all_students,
                list
            ):

                all_students = []

        except Exception:

            all_students = []

    # =====================================================
    # FILTER STUDENTS
    # =====================================================

    students = []

    if filters_selected:

        for student in all_students:

            if not isinstance(
                student,
                dict
            ):
                continue

            student_course = str(
                student.get("course")
                or student.get("course_name")
                or student.get("Course")
                or ""
            ).strip().lower()

            student_semester = str(
                student.get("semester")
                or student.get("Semester")
                or ""
            ).strip()

            student_section = str(
                student.get("section")
                or student.get("Section")
                or ""
            ).strip().upper()

            course_match = (
                not course
                or student_course == course.lower()
            )

            semester_match = (
                not semester
                or student_semester == semester
            )

            section_match = (
                not section
                or student_section == section
            )

            if (
                course_match
                and semester_match
                and section_match
            ):

                students.append(
                    student
                )

    # =====================================================
    # NUMERICAL ENROLLMENT SORT
    # =====================================================

    def enrollment_sort_key(student):

        enrollment = str(
            student.get("enrollment")
            or student.get("enrollment_no")
            or student.get("enrollment_number")
            or student.get("enrollmentNo")
            or student.get("Enrollment")
            or student.get("Enrollment No")
            or student.get("Enrollment Number")
            or ""
        ).strip()

        try:

            return (
                0,
                int(enrollment)
            )

        except Exception:

            return (
                1,
                enrollment.lower()
            )

    if filters_selected:

        students.sort(
            key=enrollment_sort_key
        )

    # =====================================================
    # IMPORT EXCEL STYLES
    # =====================================================

    from openpyxl.styles import (
        Font,
        Alignment,
        Border,
        Side
    )

    # =====================================================
    # WORKBOOK
    # =====================================================

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Sheet 1"

    # =====================================================
    # PAGE SETUP
    # =====================================================

    worksheet.sheet_view.showGridLines = False

    worksheet.page_setup.orientation = "portrait"

    worksheet.page_setup.paperSize = (
        worksheet.PAPERSIZE_A4
    )

    worksheet.page_setup.fitToWidth = 1

    worksheet.page_setup.fitToHeight = 0

    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    worksheet.page_margins.left = 0.25
    worksheet.page_margins.right = 0.25
    worksheet.page_margins.top = 0.35
    worksheet.page_margins.bottom = 0.35

    # =====================================================
    # BORDERS
    # =====================================================

    thin_side = Side(
        style="thin"
    )

    medium_side = Side(
        style="medium"
    )

    thin_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side
    )

    medium_border = Border(
        left=medium_side,
        right=medium_side,
        top=medium_side,
        bottom=medium_side
    )

    # =====================================================
    # TOP INFORMATION
    #
    # Topic       | Date
    # Course
    # Semester
    # Section
    #
    # IMPORTANT:
    # Blank information gets NO border.
    # =====================================================

    # -----------------------------------------------------
    # TOP BOX POSITIONS
    # -----------------------------------------------------

    top_boxes = []

    # Topic
    worksheet.merge_cells("A1:C1")

    worksheet["A1"] = topic

    top_boxes.append(
        ("A1", "A1:C1")
    )

    # Date
    worksheet.merge_cells("D1:E1")

    worksheet["D1"] = (
        f"DATE: {sheet_date}"
    )

    top_boxes.append(
        ("D1", "D1:E1")
    )

    # Course
    if course:

        worksheet.merge_cells("A2:E2")

        worksheet["A2"] = course

        top_boxes.append(
            ("A2", "A2:E2")
        )

    # Semester
    if semester:

        worksheet.merge_cells("A3:E3")

        semester_names = {
            "1": "1ST SEM",
            "2": "2ND SEM",
            "3": "3RD SEM",
            "4": "4TH SEM",
            "5": "5TH SEM",
            "6": "6TH SEM",
            "7": "7TH SEM",
            "8": "8TH SEM",
            "9": "9TH SEM",
            "10": "10TH SEM"
            
        }

        semester_text = semester_names.get(
            semester,
            f"{semester}TH SEM"
        )

        worksheet["A3"] = semester_text

        top_boxes.append(
            ("A3", "A3:E3")
        )

    # Section
    if section:

        worksheet.merge_cells("A4:E4")

        worksheet["A4"] = (
            f"SECTION {section}"
        )

        top_boxes.append(
            ("A4", "A4:E4")
        )

    # =====================================================
    # APPLY TOP BOX BORDERS
    # =====================================================

    for start_cell, merged_range in top_boxes:

        cells = worksheet[merged_range]

        for row_cells in cells:

            for cell in row_cells:

                cell.border = medium_border

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True
                )

                cell.font = Font(
                    bold=True,
                    size=12
                )

    # =====================================================
    # TOP ROW HEIGHT
    # =====================================================

    worksheet.row_dimensions[1].height = 30

    if course:
        worksheet.row_dimensions[2].height = 26

    if semester:
        worksheet.row_dimensions[3].height = 26

    if section:
        worksheet.row_dimensions[4].height = 26

    # =====================================================
    # PREPARE COLUMNS
    # =====================================================

    final_columns = []

    # -----------------------------------------------------
    # FILTER MODE:
    # S.NO. ALWAYS AUTOMATICALLY FIRST
    # -----------------------------------------------------

    if filters_selected:

        final_columns.append(
            "S.No."
        )

    # -----------------------------------------------------
    # MANUAL MODE:
    # S.NO. ONLY IF USER SELECTS IT
    # -----------------------------------------------------

    if not filters_selected and add_manual_sno:

        final_columns.append(
            "S.No."
        )

    # -----------------------------------------------------
    # ADD USER COLUMNS
    # -----------------------------------------------------

    for column_name in columns:

        # Prevent duplicate S.No.
        column_lower = (
            column_name
            .lower()
            .strip()
        )

        is_sno = column_lower in [
            "s.no.",
            "s.no",
            "s no",
            "sno",
            "serial no",
            "serial number",
            "sr no",
            "sr. no.",
            "sr.no.",
            "sr.no"
        ]

        if is_sno:

            if (
                filters_selected
                or add_manual_sno
            ):
                continue

        final_columns.append(
            column_name
        )

    # =====================================================
    # HEADER ROW
    # =====================================================

    header_row = 6

    for index, column_name in enumerate(
        final_columns,
        start=1
    ):

        cell = worksheet.cell(
            row=header_row,
            column=index,
            value=column_name
        )

        cell.font = Font(
            bold=True,
            size=11
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = medium_border

    worksheet.row_dimensions[
        header_row
    ].height = 30

    # =====================================================
    # NUMBER OF ROWS
    # =====================================================

    if filters_selected:

        total_rows = len(
            students
        )

    else:

        total_rows = manual_rows

    # =====================================================
    # STUDENT / MANUAL ROWS
    # =====================================================

    for row_number in range(
        total_rows
    ):

        excel_row = (
            header_row
            + 1
            + row_number
        )

        student = None

        if filters_selected:

            student = students[
                row_number
            ]

        # -------------------------------------------------
        # STUDENT NAME
        # -------------------------------------------------

        student_name = ""

        if student:

            student_name = str(
                student.get("name")
                or student.get("student_name")
                or student.get("studentName")
                or ""
            ).strip()

        # -------------------------------------------------
        # ENROLLMENT
        # -------------------------------------------------

        enrollment = ""

        if student:

            enrollment = str(
                student.get("enrollment")
                or student.get("enrollment_no")
                or student.get("enrollment_number")
                or student.get("enrollmentNo")
                or student.get("Enrollment")
                or student.get("Enrollment No")
                or student.get("Enrollment Number")
                or ""
            ).strip()

        # -------------------------------------------------
        # STUDENT ID
        # -------------------------------------------------

        student_id = ""

        if student:

            student_id = str(
                student.get("student_id")
                or student.get("studentId")
                or student.get("id")
                or student.get("studentID")
                or ""
            ).strip()

        # -------------------------------------------------
        # WRITE FINAL COLUMNS
        # -------------------------------------------------

        for column_index, column_name in enumerate(
            final_columns,
            start=1
        ):

            column_lower = (
                str(column_name)
                .lower()
                .strip()
            )

            value = ""

            # S.NO.
            if column_lower in [
                "s.no.",
                "s.no",
                "s no",
                "sno",
                "serial no",
                "serial number",
                "sr no",
                "sr. no.",
                "sr.no.",
                "sr.no"
            ]:

                value = (
                    row_number + 1
                )

            # STUDENT NAME
            elif column_lower in [
                "student name",
                "studentname",
                "name"
            ]:

                value = student_name

            # ENROLLMENT
            elif column_lower in [
                "enrollment",
                "enrollment no",
                "enrollment no.",
                "enrollment number",
                "enrollment_number",
                "enrollmentno",
                "enrollment id",
                "enrollment id.",
                "enrollment id no"
            ]:

                value = enrollment

            # STUDENT ID
            elif column_lower in [
                "student id",
                "student_id",
                "studentid",
                "student id no",
                "student id number"
            ]:

                value = student_id

            # OTHER CUSTOM COLUMNS
            else:

                value = ""

            cell = worksheet.cell(
                row=excel_row,
                column=column_index,
                value=value
            )

            # EVERY CELL HAS BORDER
            cell.border = thin_border

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

        worksheet.row_dimensions[
            excel_row
        ].height = 24

    # =====================================================
    # COLUMN WIDTH
    # =====================================================

    for column_index in range(
        1,
        len(final_columns) + 1
    ):

        column_letter = get_column_letter(
            column_index
        )

        column_name = (
            final_columns[
                column_index - 1
            ]
            .lower()
            .strip()
        )

        if column_name in [
            "s.no.",
            "s.no",
            "s no",
            "sno",
            "serial no",
            "serial number",
            "sr no",
            "sr. no.",
            "sr.no.",
            "sr.no"
        ]:

            width = 8

        elif column_name in [
            "student name",
            "studentname",
            "name"
        ]:

            width = 28

        elif column_name in [
            "enrollment",
            "enrollment no",
            "enrollment no.",
            "enrollment number",
            "enrollment_number",
            "enrollmentno",
            "enrollment id",
            "enrollment id.",
            "enrollment id no"
        ]:

            width = 20

        elif column_name in [
            "student id",
            "student_id",
            "studentid"
        ]:

            width = 18

        elif column_name in [
            "signature",
            "sign"
        ]:

            width = 24

        else:

            width = 18

        worksheet.column_dimensions[
            column_letter
        ].width = width

    # =====================================================
    # FREEZE
    # =====================================================

    worksheet.freeze_panes = "A7"

    # =====================================================
    # PRINT AREA
    # =====================================================

    last_row = (
        header_row
        + 1
        + total_rows
    )

    last_column = get_column_letter(
        len(final_columns)
    )

    worksheet.print_area = (
        f"A1:{last_column}{last_row}"
    )

    # =====================================================
    # REPEAT HEADER ON PRINTED PAGES
    # =====================================================

    worksheet.print_title_rows = "1:6"

   # =====================================================
    # AUTOFILTER
    # =====================================================

    if total_rows > 0:

        worksheet.auto_filter.ref = (
            f"A{header_row}:"
            f"{last_column}{last_row}"
        )

    # =====================================================
    # SAVE FOLDER
    # =====================================================

    excel_folder = "generated_excel"

    os.makedirs(
        excel_folder,
        exist_ok=True
    )

    # =====================================================
    # SAFE FILE NAME
    # =====================================================

    safe_topic = "".join(
        char
        if char.isalnum()
        or char in "-_"
        else "_"
        for char in topic
    )

    filename = (
        f"{safe_topic}_{sheet_date}.xlsx"
    )

    file_path = os.path.join(
        excel_folder,
        filename
    )

    # =====================================================
    # SAVE XLSX
    # =====================================================

    workbook.save(
        file_path
    )

    # =====================================================
    # DOWNLOAD
    # =====================================================

    return redirect(
        url_for(
            "download_created_excel",
            filename=filename
        )
    )


# =========================================================
# DOWNLOAD CREATED EXCEL
# =========================================================

@app.route(
    "/download-created-excel/<path:filename>"
)
def download_created_excel(filename):

    excel_folder = "generated_excel"

    file_path = os.path.join(
        excel_folder,
        filename
    )

    if not os.path.isfile(
        file_path
    ):

        return (
            "Excel file not found.",
            404
        )

    return send_file(
        file_path,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        conditional=False
    )


from functools import wraps

def finance_required(view_func):

    @wraps(view_func)
    def wrapped_view(*args, **kwargs):

        if not session.get("finance_logged_in"):
            return redirect("/finance-login")

        return view_func(*args, **kwargs)

    return wrapped_view


@app.route("/finance")
@finance_required
def finance():

    if not session.get("finance_logged_in"):
        return redirect("/finance-login")

    return render_template(
        "finance.html"
    )


@app.route("/finance-login", methods=["GET", "POST"])
def finance_login():

    if request.method == "POST":

        verification_name = request.form.get(
            "verification_name",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        ).strip()

        if not verification_name or not password:
            return render_template(
                "finance_login.html",
                error="Please enter verification name and password."
            )

        conn = get_connection()

        try:

            with conn.cursor() as cur:

                cur.execute("""
                    CREATE TABLE IF NOT EXISTS finance_credentials (
                        id INTEGER PRIMARY KEY,
                        verification_name TEXT NOT NULL,
                        password TEXT NOT NULL
                    )
                """)

                cur.execute("""
                    SELECT
                        verification_name,
                        password
                    FROM finance_credentials
                    WHERE id = 1
                    LIMIT 1
                """)

                credentials = cur.fetchone()

                # FIRST TIME LOGIN
                if not credentials:

                    cur.execute("""
                        INSERT INTO finance_credentials
                        (
                            id,
                            verification_name,
                            password
                        )
                        VALUES (1, %s, %s)
                    """, (
                        verification_name,
                        password
                    ))

                    conn.commit()

                    session["finance_logged_in"] = True

                    return redirect("/finance")

                # NORMAL LOGIN
                if (
    str(verification_name).strip()
    == str(credentials["verification_name"]).strip()
    and
    str(password).strip()
    == str(credentials["password"]).strip()
):

                    session["finance_logged_in"] = True

                    return redirect("/finance")

                return render_template(
                    "finance_login.html",
                    error="Invalid verification name or password."
                )

        finally:

            conn.close()

    return render_template(
        "finance_login.html"
    )




@app.route("/finance-change-credentials", methods=["GET", "POST"])
@finance_required
def finance_change_credentials():

    # Finance login check
    if not session.get("finance_logged_in"):
        return redirect("/finance-login")


    if request.method == "POST":

        current_name = request.form.get(
            "current_name",
            ""
        ).strip()

        current_password = request.form.get(
            "current_password",
            ""
        ).strip()

        new_name = request.form.get(
            "new_name",
            ""
        ).strip()

        new_password = request.form.get(
            "new_password",
            ""
        ).strip()

        confirm_password = request.form.get(
            "confirm_password",
            ""
        ).strip()


        # =========================
        # BASIC CHECK
        # =========================

        if not current_name or not current_password:
            return render_template(
                "finance_change_credentials.html",
                error="Please enter current verification name and password."
            )


        if not new_name or not new_password:
            return render_template(
                "finance_change_credentials.html",
                error="Please enter new verification name and password."
            )


        if new_password != confirm_password:
            return render_template(
                "finance_change_credentials.html",
                error="New passwords do not match."
            )


        conn = get_connection()

        try:

            with conn.cursor() as cur:

                cur.execute("""
                    SELECT
                        verification_name,
                        password
                    FROM finance_credentials
                    WHERE id = 1
                    LIMIT 1
                """)

                credentials = cur.fetchone()


                if not credentials:
                    return render_template(
                        "finance_change_credentials.html",
                        error="Finance credentials not found."
                    )


                # =========================
                # CURRENT CREDENTIAL CHECK
                # =========================

                if (
                    current_name
                    != credentials["verification_name"]
                    or
                    current_password
                    != credentials["password"]
                ):

                    return render_template(
                        "finance_change_credentials.html",
                        error="Current verification name or password is incorrect."
                    )


                # =========================
                # UPDATE CREDENTIALS
                # =========================

                cur.execute("""
                    UPDATE finance_credentials
                    SET
                        verification_name = %s,
                        password = %s
                    WHERE id = 1
                """, (
                    new_name,
                    new_password
                ))

            conn.commit()

        finally:

            conn.close()


        # Logout after changing credentials
        session.pop(
            "finance_logged_in",
            None
        )


        return redirect(
            "/finance-login"
        )


    return render_template(
        "finance_change_credentials.html"
    )



@app.route("/finance-payment", methods=["GET", "POST"])
def finance_payment():

    # =========================
    # STUDENT LOGIN CHECK
    # =========================

    if "student_enrollment" not in session:
        return redirect("/student-login")

    student_enrollment = str(
        session.get("student_enrollment", "")
    ).strip()


    # =========================
    # STUDENTS LOAD
    # =========================

    students = []

    if os.path.exists(STUDENTS_FILE):

        try:
            with open(
                STUDENTS_FILE,
                "r"
            ) as file:

                students = json.load(file)

        except:
            students = []


    # =========================
    # CURRENT STUDENT FIND
    # =========================

    student = None

    for item in students:

        if str(
            item.get("enrollment", "")
        ).strip() == student_enrollment:

            student = item
            break


    if not student:
        return "Student not found.", 404


    # =========================
    # POST - PAYMENT
    # =========================

    if request.method == "POST":

        enrollment = request.form.get(
            "enrollment",
            ""
        ).strip()

        screenshot = request.files.get(
            "screenshot"
        )

        fee_type = request.form.get(
            "fee_type",
            ""
        ).strip()


        if not enrollment:
            return "Please enter Enrollment Number.", 400


        if not screenshot or not screenshot.filename:
            return "Please upload payment screenshot.", 400


        if not fee_type:
            return "Please select Fee Type.", 400


        # Student apne hi enrollment se payment karega
        if enrollment != student_enrollment:
            return "Invalid Enrollment Number.", 400


        # =========================
        # SCREENSHOT SAVE
        # =========================

        original_name = os.path.basename(
            screenshot.filename
        )

        # Same filename dobara upload na ho
        conn = get_connection()

        try:

            with conn.cursor() as cur:

                cur.execute("""
                    SELECT 1
                    FROM payment_records
                    WHERE LOWER(
                        regexp_replace(
                            screenshot,
                            '^payment_[^_]+_',
                            ''
                        )
                    ) = LOWER(%s)
                    LIMIT 1
                """, (original_name,))

                existing_file = cur.fetchone()

        finally:

            conn.close()


        if existing_file:

            return (
                "This screenshot filename has already been uploaded.",
                400
            )


        stored_filename = (
            "payment_"
            + uuid.uuid4().hex
            + "_"
            + original_name
        )

        save_uploaded_file(
            screenshot,
            stored_filename
        )


        # =========================
        # PAYMENT RECORD
        # =========================

        payment_record = {

            "enrollment": enrollment,

            "student_id": str(
                student.get("student_id", "")
            ),

            "student_name": str(
                student.get("name", "")
            ),

            "course": str(
                student.get("course", "")
            ),

            "semester": str(
                student.get("semester", "")
            ),

            "section": str(
                student.get("section", "")
            ),

            "group": str(
                student.get("group", "")
            ),

            "fee_type": fee_type,

            "screenshot": stored_filename,

            "status": "Recent",

            "amount": "",

            "submitted_date": ""

        }


        # =========================
        # SAVE PAYMENT RECORD IN DATABASE
        # =========================

        conn = get_connection()

        try:

            with conn.cursor() as cur:

                cur.execute("""
                    INSERT INTO payment_records (
                        enrollment,
                        student_id,
                        student_name,
                        course,
                        semester,
                        section,
                        "group",
                        fee_type,
                        screenshot,
                        status,
                        amount,
                        submitted_date
                    )
                    VALUES (
                        %s, %s, %s, %s, %s, %s,
                        %s, %s, %s, %s, %s, %s
                    )
                """, (
                    payment_record["enrollment"],
                    payment_record["student_id"],
                    payment_record["student_name"],
                    payment_record["course"],
                    payment_record["semester"],
                    payment_record["section"],
                    payment_record["group"],
                    payment_record["fee_type"],
                    payment_record["screenshot"],
                    payment_record["status"],
                    payment_record["amount"],
                    payment_record["submitted_date"]
                ))

            conn.commit()

        finally:

            conn.close()


        return redirect(
            "/finance-payment"
        )


    # =========================
    # PENDING PAYMENT RECORDS
    # =========================

    pending_records = []

    conn = get_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                SELECT
                    id,
                    enrollment,
                    student_name,
                    course,
                    semester,
                    section,
                    "group",
                    fee_type,
                    screenshot,
                    status,
                    created_at
                FROM payment_records
                WHERE enrollment = %s
                  AND status != 'Submitted'
                ORDER BY created_at DESC
            """, (
                student_enrollment,
            ))

            pending_records = cur.fetchall()

    finally:

        conn.close()


    # =========================
    # PAYMENT PAGE
    # =========================

    return render_template(
        "finance_payment.html",
        student=student,
        pending_records=pending_records
    )


@app.route("/finance-payment-records")
@finance_required
def finance_payment_records():

    return render_template(
        "finance_payment_record.html"
    )


@app.route("/finance-payment-records/recent")
@finance_required
def finance_recent_records():

    recent_records = []

    # =========================
    # LOAD RECENT PAYMENT RECORDS FROM DATABASE
    # =========================

    conn = get_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                SELECT
                    id,
                    enrollment,
                    student_id,
                    student_name,
                    course,
                    semester,
                    section,
                    "group",
                    fee_type,
                    screenshot,
                    status,
                    amount,
                    submitted_date
                FROM payment_records
                WHERE status != 'Submitted'
                ORDER BY created_at DESC
            """)

            recent_records = cur.fetchall()

    finally:

        conn.close()


    return render_template(
        "finance_recent_records.html",
        recent_records=recent_records
    )


@app.route("/finance-payment-records/submitted")
@finance_required
def finance_submitted_records():

    submitted_records = []

    # =========================
    # LOAD SUBMITTED RECORDS FROM DATABASE
    # =========================

    conn = get_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                SELECT
                    id,
                    enrollment,
                    student_id,
                    student_name,
                    course,
                    semester,
                    section,
                    "group",
                    fee_type,
                    receipt_no,
                    screenshot,
                    status,
                    amount,
                    submitted_date
                FROM payment_records
                WHERE status = 'Submitted'
                ORDER BY created_at DESC
            """)

            submitted_records = cur.fetchall()

    finally:

        conn.close()


    return render_template(
        "finance_submitted_records.html",
        submitted_records=submitted_records
    )


@app.route("/payment-screenshot/<filename>")
def payment_screenshot(filename):

    stored = get_uploaded_file(filename)

    if not stored:
        return "Payment screenshot not found.", 404

    download = request.args.get("download")

    return send_file(
        io.BytesIO(bytes(stored["content"])),
        mimetype=stored["content_type"],
        download_name=filename,
        as_attachment=(download == "1")
    )

@app.route("/submit-payment-record/<int:payment_id>", methods=["POST"])
def submit_payment_record(payment_id):

    amount = request.form.get(
        "amount",
        ""
    ).strip()

    source = request.form.get(
        "source",
        "recent"
    ).strip()


    if not amount:
        return "Please enter payment amount.", 400


    conn = get_connection()

    try:

        with conn.cursor() as cur:

            # =========================
            # PAYMENT RECORD FIND
            # =========================

            cur.execute("""
                SELECT enrollment
                FROM payment_records
                WHERE id = %s
            """, (
                payment_id,
            ))

            payment = cur.fetchone()

            if not payment:
                return "Payment record not found.", 404


            enrollment = str(
                payment["enrollment"]
            ).strip()


            # =========================
            # RECEIPT NUMBER GENERATE
            # =========================

            cur.execute("""
                SELECT COUNT(*)
                FROM payment_records
                WHERE status = 'Submitted'
                  AND EXTRACT(
                      YEAR FROM created_at
                  ) = EXTRACT(
                      YEAR FROM CURRENT_DATE
                  )
            """)

            count_result = cur.fetchone()

            current_count = int(
                count_result["count"]
            )

            receipt_no = (
                str(
                    __import__("datetime")
                    .date.today()
                    .year
                )
                + "-"
                + str(current_count + 1)
            )


            # =========================
            # SUBMIT ONLY THIS PAYMENT
            # =========================

            cur.execute("""
                UPDATE payment_records
                SET
                    status = 'Submitted',
                    amount = %s,
                    submitted_date = CURRENT_DATE::text,
                    receipt_no = %s
                WHERE id = %s
                  AND status != 'Submitted'
            """, (
                amount,
                receipt_no,
                payment_id
            ))


            if cur.rowcount == 0:

                return (
                    "Payment record not found or already submitted.",
                    404
                )


        conn.commit()

    finally:

        conn.close()


    # =========================
    # REDIRECT
    # =========================

    if source == "search":

        return redirect(
            "/finance-search-record?enrollment="
            + enrollment
        )


    return redirect(
        "/finance-payment-records/recent"
    )



@app.route("/fees-slip")
def fees_slip():

    if "student_enrollment" not in session:
        return redirect("/student-login")

    student_enrollment = str(
        session.get("student_enrollment", "")
    ).strip()

    conn = get_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                SELECT
                    id,
                    enrollment,
                    student_name,
                    course,
                    semester,
                    section,
                    "group",
                    fee_type,
                    receipt_no,
                    amount,
                    submitted_date,
                    screenshot,
                    status
                FROM payment_records
                WHERE enrollment = %s
                  AND status = 'Submitted'
                ORDER BY created_at DESC
            """, (student_enrollment,))

            payments = cur.fetchall()

    finally:
        conn.close()

    return render_template(
        "fees_slip.html",
        payments=payments
    )


@app.route("/fees-slip/<int:payment_id>")
def open_fees_slip(payment_id):

    if "student_enrollment" not in session:
        return redirect("/student-login")

    student_enrollment = str(
        session.get("student_enrollment", "")
    ).strip()

    conn = get_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                SELECT
                    id,
                    enrollment,
                    student_name,
                    fee_type,
                    receipt_no,
                    course,
                    semester,
                    section,
                    "group",
                    amount,
                    submitted_date,
                    screenshot,
                    status
                FROM payment_records
                WHERE id = %s
                  AND enrollment = %s
                  AND status = 'Submitted'
                LIMIT 1
            """, (
                payment_id,
                student_enrollment
            ))

            payment = cur.fetchone()

    finally:
        conn.close()

    if not payment:
        return "Fees slip not found.", 404


    # =========================
    # DOWNLOAD PDF
    # =========================

    if request.args.get("download") == "1":

        pdf_buffer = io.BytesIO()

        pdf = canvas.Canvas(
            pdf_buffer,
            pagesize=A4
        )

        width, height = A4

        y = height - 60


        # UNIVERSITY NAME

        pdf.setFont(
            "Helvetica-Bold",
            16
        )

        pdf.drawCentredString(
            width / 2,
            y,
            "MANGALAYATAN UNIVERSITY BESWAN, ALIGARH"
        )

        y -= 30


        # TITLE

        pdf.setFont(
            "Helvetica-Bold",
            14
        )

        pdf.drawCentredString(
            width / 2,
            y,
            "FEES PAYMENT RECEIPT"
        )

        y -= 45


        pdf.setFont(
            "Helvetica",
            11
        )


        # STUDENT DETAILS

        details = [

            ("Student Name", payment["student_name"]),

            ("Fee Type", payment["fee_type"]),

            ("Receipt No.", payment["receipt_no"] or "—"),

            ("Enrollment Number", payment["enrollment"]),

            ("Course", payment["course"]),

            ("Semester", payment["semester"]),

            ("Section", payment["section"] or "—"),

            ("Group", payment["group"] or "—"),

            ("Payment Date", payment["submitted_date"]),

            ("Payment Status", "Submitted"),

            ("Payment Amount", "Rs. " + str(payment["amount"]))

        ]


        for label, value in details:

            pdf.setFont(
                "Helvetica-Bold",
                10
            )

            pdf.drawString(
                70,
                y,
                str(label) + ":"
            )

            pdf.setFont(
                "Helvetica",
                10
            )

            pdf.drawString(
                210,
                y,
                str(value)
            )

            y -= 25


        y -= 15


        pdf.setFont(
            "Helvetica-Bold",
            11
        )

        pdf.drawCentredString(
            width / 2,
            y,
            "PAYMENT RECORD SUBMITTED"
        )


        y -= 45


        pdf.setFont(
            "Helvetica",
            9
        )

        pdf.drawString(
            70,
            y,
            "Student Fees Receipt"
        )


        pdf.save()

        pdf_buffer.seek(0)


        return send_file(
            pdf_buffer,
            mimetype="application/pdf",
            as_attachment=True,
            download_name=(
                "fees_slip_"
                + str(payment["enrollment"])
                + "_"
                + str(payment["id"])
                + ".pdf"
            )
        )


    # =========================
    # OPEN FEES SLIP
    # =========================

    return render_template(
        "fees_slip_single.html",
        payment=payment
    )


@app.route("/finance-search-record", methods=["GET", "POST"])
@finance_required
def finance_search_record():

    searched_enrollment = ""
    recent_records = []
    submitted_records = []
    semester_totals = {}


    # =========================
    # GET - AFTER SUBMIT
    # =========================

    if request.method == "GET":

        searched_enrollment = request.args.get(
            "enrollment",
            ""
        ).strip()


    # =========================
    # POST - SEARCH
    # =========================

    if request.method == "POST":

        searched_enrollment = request.form.get(
            "enrollment",
            ""
        ).strip()


    # =========================
    # LOAD PAYMENT RECORDS
    # =========================

    if searched_enrollment:

        conn = get_connection()

        try:

            with conn.cursor() as cur:

                cur.execute("""
                    SELECT
                        id,
                        enrollment,
                        student_name,
                        course,
                        semester,
                        section,
                        "group",
                        fee_type,
                        screenshot,
                        status,
                        amount,
                        submitted_date,
                        created_at
                    FROM payment_records
                    WHERE enrollment = %s
                    ORDER BY created_at ASC
                """, (
                    searched_enrollment,
                ))

                records = cur.fetchall()

        finally:

            conn.close()


        # =========================
        # SEPARATE RECORDS
        # =========================

        for record in records:

            if record["status"] == "Submitted":

                submitted_records.append(
                    record
                )

                # =========================
                # SEMESTER-WISE TOTAL
                # =========================

                semester = record["semester"]

                if semester not in semester_totals:
                    semester_totals[semester] = 0

                semester_totals[semester] += float(
                    record["amount"] or 0
                )

            else:

                recent_records.append(
                    record
                )


    # =========================
    # SORT SEMESTERS 1 TO 10
    # =========================

    def semester_number(semester):

        try:

            return int(
                ''.join(
                    filter(
                        str.isdigit,
                        str(semester)
                    )
                )
            )

        except:

            return 999


    semester_totals = dict(
        sorted(
            semester_totals.items(),
            key=lambda item: semester_number(item[0])
        )
    )


    return render_template(
        "finance_search_record.html",
        searched_enrollment=searched_enrollment,
        recent_records=recent_records,
        submitted_records=submitted_records,
        semester_totals=semester_totals
    )


@app.route("/finance-search-receipt", methods=["GET", "POST"])
@finance_required
def finance_search_receipt():

    receipt_no = ""
    payment = None

    if request.method == "POST":

        receipt_no = request.form.get(
            "receipt_no",
            ""
        ).strip()

        if receipt_no:

            conn = get_connection()

            try:

                with conn.cursor() as cur:

                    cur.execute("""
                        SELECT
                            id,
                            enrollment,
                            student_id,
                            student_name,
                            course,
                            semester,
                            section,
                            "group",
                            fee_type,
                            screenshot,
                            status,
                            amount,
                            submitted_date,
                            receipt_no,
                            created_at
                        FROM payment_records
                        WHERE receipt_no = %s
                          AND status = 'Submitted'
                        LIMIT 1
                    """, (
                        receipt_no,
                    ))

                    payment = cur.fetchone()

            finally:

                conn.close()


    return render_template(
        "finance_search_receipt.html",
        receipt_no=receipt_no,
        payment=payment
    )


# =========================
# OWNER LOGIN
# =========================

@app.route("/owner-login", methods=["GET", "POST"])
def owner_login():

    message = ""

    if request.method == "POST":

        verification_name = request.form.get(
            "verification_name", ""
        ).strip()

        password = request.form.get(
            "password", ""
        ).strip()

        if not verification_name or not password:

            message = "Please enter Verification Name and Password."

        else:

            conn = get_connection()

            try:

                with conn.cursor() as cur:

                    cur.execute("""
                        SELECT
                            verification_name,
                            password
                        FROM owner_credentials
                        WHERE id = 1
                    """)

                    owner_data = cur.fetchone()

                    # First time Owner login
                    if not owner_data:

                        default_name = "OWNER"
                        default_password = "owner123"

                        cur.execute("""
                            INSERT INTO owner_credentials (
                                id,
                                verification_name,
                                password
                            )
                            VALUES (1, %s, %s)
                        """, (
                            default_name,
                            generate_password_hash(
                                default_password
                            )
                        ))

                        conn.commit()

                        owner_data = {
                            "verification_name": default_name,
                            "password": generate_password_hash(
                                default_password
                            )
                        }

            finally:
                conn.close()

            if (
                verification_name
                != owner_data["verification_name"]
            ):

                message = "Invalid Verification Name or Password."

            elif not check_password_hash(
                owner_data["password"],
                password
            ):

                message = "Invalid Verification Name or Password."

            else:

                # Purane roles ki session clear
                session.pop("admin_id", None)
                session.pop("admin_department", None)
                session.pop("admin_department_id", None)

                session.pop("student_id", None)
                session.pop("student_enrollment", None)
                session.pop("student_course", None)
                session.pop("student_semester", None)
                session.pop("student_section", None)

                session.pop("faculty_id", None)
                session.pop("faculty_name", None)

                session["owner_logged_in"] = True

                return redirect("/owner")

    return render_template(
        "owner_login.html",
        message=message
    )


# =========================
# OWNER LOGIN SECURITY
# =========================

def owner_required():

    if not session.get("owner_logged_in"):

        session.pop("owner_logged_in", None)

        return redirect(
            url_for("owner_login")
        )

    return None

@app.route("/owner")
def owner():

    security_check = owner_required()

    if security_check:
        return security_check

    departments = []
    admins = []

    if os.path.exists(DEPARTMENTS_FILE):

        try:
            with open(DEPARTMENTS_FILE, "r") as file:
                departments = json.load(file)

        except:
            departments = []

    conn = get_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                SELECT
                    department_id,
                    department_name,
                    admin_id,
                    is_closed
                FROM department_admins
                ORDER BY id
            """)

            admins = cur.fetchall()

    finally:
        conn.close()

    return render_template(
        "owner.html",
        departments=departments,
        admins=admins
    )


# =========================
# OWNER - CHANGE CREDENTIALS
# =========================

@app.route("/owner-change-password", methods=["GET", "POST"])
def owner_change_password():

    security_check = owner_required()

    if security_check:
        return security_check

    message = ""

    if request.method == "POST":

        current_name = request.form.get(
            "current_name", ""
        ).strip()

        current_password = request.form.get(
            "current_password", ""
        ).strip()

        new_name = request.form.get(
            "new_name", ""
        ).strip()

        new_password = request.form.get(
            "new_password", ""
        ).strip()

        confirm_password = request.form.get(
            "confirm_password", ""
        ).strip()

        conn = get_connection()

        try:

            with conn.cursor() as cur:

                cur.execute("""
                    SELECT
                        verification_name,
                        password
                    FROM owner_credentials
                    WHERE id = 1
                """)

                owner_data = cur.fetchone()

                if not owner_data:

                    message = "Owner account not found."

                elif current_name != owner_data["verification_name"]:

                    message = "Current Verification Name is incorrect."

                elif not check_password_hash(
                    owner_data["password"],
                    current_password
                ):

                    message = "Current Password is incorrect."

                elif not new_name or not new_password:

                    message = "Please enter new Verification Name and Password."

                elif new_password != confirm_password:

                    message = "New passwords do not match."

                elif len(new_password) < 6:

                    message = "New password must be at least 6 characters."

                else:

                    cur.execute("""
                        UPDATE owner_credentials
                        SET
                            verification_name = %s,
                            password = %s
                        WHERE id = 1
                    """, (
                        new_name,
                        generate_password_hash(new_password)
                    ))

                    conn.commit()

                    message = "Owner credentials changed successfully."

        finally:
            conn.close()

    return render_template(
        "owner_change_password.html",
        message=message
    )

# =========================
# OWNER - ADD ADMIN
# =========================

@app.route("/owner-add-admin", methods=["POST"])
def owner_add_admin():

    department_id = request.form.get("department_id", "").strip()

    if not department_id:
        return "Please select a department."

    # Existing departments load karo
    departments = []

    if os.path.exists(DEPARTMENTS_FILE):
        try:
            with open(DEPARTMENTS_FILE, "r") as file:
                departments = json.load(file)
        except:
            departments = []

    # Selected department find karo
    selected_department = None

    for department in departments:

        if department.get("department_id") == department_id:
            selected_department = department
            break

    if not selected_department:
        return "Invalid Department."

    department_name = selected_department.get("name", "").strip()

    if not department_name:
        return "Invalid Department."

    conn = get_connection()

    try:

        with conn.cursor() as cur:

            # Check department already has an Admin
            cur.execute("""
                SELECT id
                FROM department_admins
                WHERE department_id = %s
            """, (department_id,))

            existing_admin = cur.fetchone()

            if existing_admin:
                return "Admin already exists for this Department."

            # Next Admin ID
            cur.execute("""
                SELECT admin_id
                FROM department_admins
                ORDER BY id DESC
                LIMIT 1
            """)

            last_admin = cur.fetchone()

            if last_admin:
                last_id = last_admin["admin_id"]

                try:
                    number = int(
                        last_id.replace("ADMIN", "")
                    ) + 1
                except:
                    number = 1

            else:
                number = 1

            admin_id = "ADMIN" + str(number).zfill(3)

            # Initial password = Admin ID
            password_hash = generate_password_hash(admin_id)

            cur.execute("""
                INSERT INTO department_admins (
                    department_id,
                    department_name,
                    admin_id,
                    password,
                    is_closed
                )
                VALUES (%s, %s, %s, %s, FALSE)
            """, (
                department_id,
                department_name,
                admin_id,
                password_hash
            ))

        conn.commit()

    finally:
        conn.close()

    return redirect("/owner")

# =========================
# OWNER - RESET ADMIN PASSWORD
# =========================

@app.route("/owner-reset-admin-password/<admin_id>", methods=["POST"])
def owner_reset_admin_password(admin_id):

    conn = get_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                UPDATE department_admins
                SET password = %s
                WHERE admin_id = %s
            """, (
                generate_password_hash(admin_id),
                admin_id
            ))

        conn.commit()

    finally:
        conn.close()

    return redirect("/owner")


# =========================
# OWNER - TEMPORARY CLOSE / OPEN ADMIN
# =========================

@app.route("/owner-toggle-admin/<admin_id>", methods=["POST"])
def owner_toggle_admin(admin_id):

    conn = get_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                UPDATE department_admins
                SET is_closed = NOT is_closed
                WHERE admin_id = %s
            """, (admin_id,))

        conn.commit()

    finally:
        conn.close()

    return redirect("/owner")


# =========================
# OWNER - DELETE ADMIN
# =========================

@app.route("/owner-delete-admin/<admin_id>", methods=["POST"])
def owner_delete_admin(admin_id):

    conn = get_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                DELETE FROM department_admins
                WHERE admin_id = %s
            """, (admin_id,))

        conn.commit()

    finally:
        conn.close()

    return redirect("/owner")



def head_required(view_func):

    @wraps(view_func)
    def wrapped_view(*args, **kwargs):

        if not session.get("head_logged_in"):
            return redirect("/head-login")

        return view_func(*args, **kwargs)

    return wrapped_view


@app.route("/head")
@head_required
def head():

    response = make_response(
        render_template("head.html")
    )

    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"

    return response


@app.route("/head-login", methods=["GET", "POST"])
def head_login():

    if request.method == "POST":

        verification_name = request.form.get(
            "verification_name",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        ).strip()

        if not verification_name or not password:

            return render_template(
                "head_login.html",
                error="Please enter verification name and password."
            )

        conn = get_connection()

        try:

            with conn.cursor() as cur:

                cur.execute("""
                    CREATE TABLE IF NOT EXISTS head_credentials (
                        id INTEGER PRIMARY KEY,
                        verification_name TEXT NOT NULL,
                        password TEXT NOT NULL
                    )
                """)

                cur.execute("""
                    SELECT
                        verification_name,
                        password
                    FROM head_credentials
                    WHERE id = 1
                    LIMIT 1
                """)

                credentials = cur.fetchone()


                # FIRST TIME LOGIN

                if not credentials:

                    cur.execute("""
                        INSERT INTO head_credentials
                        (
                            id,
                            verification_name,
                            password
                        )
                        VALUES (1, %s, %s)
                    """, (
                        verification_name,
                        password
                    ))

                    conn.commit()

                    session["head_logged_in"] = True

                    return redirect("/head")


                # NORMAL LOGIN

                if (
                    str(verification_name).strip()
                    == str(credentials["verification_name"]).strip()
                    and
                    str(password).strip()
                    == str(credentials["password"]).strip()
                ):

                    session["head_logged_in"] = True

                    return redirect("/head")


                return render_template(
                    "head_login.html",
                    error="Invalid verification name or password."
                )

        finally:

            conn.close()


    return render_template(
        "head_login.html"
    )


    
@app.route("/head-change-password", methods=["GET", "POST"])
def head_change_password():

    # =========================
    # HEAD LOGIN CHECK
    # =========================

    if not session.get("head_logged_in"):
        return redirect(url_for("head_login"))


    # =========================
    # CHANGE LOGIN DETAILS
    # =========================

    if request.method == "POST":

        current_verification_name = request.form.get(
            "current_verification_name",
            ""
        ).strip()

        current_password = request.form.get(
            "current_password",
            ""
        ).strip()

        new_verification_name = request.form.get(
            "verification_name",
            ""
        ).strip()

        new_password = request.form.get(
            "password",
            ""
        ).strip()

        confirm_password = request.form.get(
            "confirm_password",
            ""
        ).strip()


        # =========================
        # EMPTY FIELD CHECK
        # =========================

        if (
            not current_verification_name
            or not current_password
            or not new_verification_name
            or not new_password
            or not confirm_password
        ):

            return render_template(
                "head_change_password.html",
                error="Please fill all fields."
            )


        # =========================
        # NEW PASSWORD CONFIRM
        # =========================

        if new_password != confirm_password:

            return render_template(
                "head_change_password.html",
                error="New password and confirm password do not match."
            )


        # =========================
        # DATABASE
        # =========================

        conn = get_connection()

        try:

            with conn.cursor() as cur:

                # =========================
                # CREATE TABLE IF NOT EXISTS
                # =========================

                cur.execute("""
                    CREATE TABLE IF NOT EXISTS head_credentials (
                        id INTEGER PRIMARY KEY,
                        verification_name TEXT NOT NULL,
                        password TEXT NOT NULL
                    )
                """)


                # =========================
                # GET CURRENT DETAILS
                # =========================

                cur.execute("""
                    SELECT
                        verification_name,
                        password
                    FROM head_credentials
                    WHERE id = 1
                    LIMIT 1
                """)

                credentials = cur.fetchone()


                # =========================
                # CHECK CURRENT DETAILS
                # =========================

                if not credentials:

                    return render_template(
                        "head_change_password.html",
                        error="Head login credentials were not found."
                    )


                if (
                    str(current_verification_name).strip()
                    != str(
                        credentials["verification_name"]
                    ).strip()
                    or
                    str(current_password).strip()
                    != str(
                        credentials["password"]
                    ).strip()
                ):

                    return render_template(
                        "head_change_password.html",
                        error="Current verification name or password is incorrect."
                    )


                # =========================
                # UPDATE NEW DETAILS
                # =========================

                cur.execute("""
                    UPDATE head_credentials
                    SET
                        verification_name = %s,
                        password = %s
                    WHERE id = 1
                """, (
                    new_verification_name,
                    new_password
                ))


            conn.commit()


        finally:

            conn.close()


        # =========================
        # LOGIN SESSION CONTINUE
        # =========================

        session["head_logged_in"] = True


        # =========================
        # SUCCESS
        # =========================

        return """
        <script>

            alert(
                "Head ID & Password changed successfully."
            );

            window.location.href = "/head";

        </script>
        """


    # =========================
    # GET PAGE
    # =========================

    return render_template(
        "head_change_password.html"
    )


    


# =========================
# RUN APPLICATION
# =========================

if __name__ == "__main__":
    # fix_student_ids()
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)